"""Stage-scoped diagnostics with capability-tested Office application selection."""
import argparse
import importlib.util
import json
import shutil
from canva_bridge import active_identity
from datetime import datetime, timezone
import os
import tempfile
from pathlib import Path
from runtime import (validate_manifest_hashes, load_runtime_state, stage2_artifact_current,
                     source_fingerprint, text_clean_fingerprint, active_attempt_matches)

SCOPES = ('package', 'merge', 'acquisition', 'full', 'image_layer')
TEST_DEPENDENCIES = ('pytest', 'PIL')


def _probe_powerpoint(output_dir):
    app = None
    prs = None
    ok, detail = False, ''
    cleanup = []
    try:
        import win32com.client
        with tempfile.TemporaryDirectory(dir=output_dir) as td:
            source, target = Path(td)/'source.pptx', Path(td)/'target.pptx'
            app = win32com.client.DispatchEx('PowerPoint.Application')
            prs = app.Presentations.Add(-1)
            prs.Slides.Add(1, 12)
            prs.SaveAs(str(source), 24)
            prs.Close(); prs = None
            prs = app.Presentations.Add(-1)
            prs.SaveAs(str(target), 24)
            prs.Close(); prs = None
            prs = app.Presentations.Open(str(target), 0, 0, 0)
            assert prs.Slides.InsertFromFile(str(source), 0) == 1
            prs.Save()
            prs.Close(); prs = None
            prs = app.Presentations.Open(str(target), 0, 0, 0)
            assert prs.Slides.Count == 1
            prs.Close(); prs = None
        ok = True
    except Exception as exc:
        detail = str(exc)
    finally:
        if prs is not None:
            try:
                prs.Close()
            except Exception as exc:
                cleanup.append('document close: ' + str(exc))
        if app is not None:
            try:
                app.Quit()
            except Exception as exc:
                cleanup.append('application quit: ' + str(exc))
    return ok and not cleanup, '; '.join(([detail] if detail else []) + cleanup)


def _probe_wps(output_dir):
    from pptx import Presentation
    from merge_pptx import _native_merge
    with tempfile.TemporaryDirectory(dir=output_dir) as td:
        source, target = Path(td)/'source.pptx', Path(td)/'target.pptx'
        prs = Presentation(); prs.slides.add_slide(prs.slide_layouts[6]); prs.save(source)
        _native_merge([{'slide_id': 'probe', 'path': str(source)}], ['probe'], target, office_host='wps')
        if len(Presentation(target).slides) != 1:
            raise RuntimeError('WPS probe did not reopen one slide')
    return True, ''


def _safe_office_probe(probe, output_dir):
    try:
        return probe(output_dir)
    except Exception as exc:
        return False, str(exc)




def _check_codex_host(skill_dir, host_report, runtime_state, slide_id):
    """Validate local evidence, not remote entitlement or reconstruction quality."""
    if not (Path(skill_dir)/'scripts/canva_pptx_finalize.ps1').is_file():
        raise ValueError('local finalizer missing')
    if not (shutil.which('pwsh') or shutil.which('powershell')):
        raise ValueError('PowerShell unavailable')
    if not host_report or not runtime_state or not slide_id:
        raise ValueError('host report, trusted runtime state and slide_id required')
    report = json.loads(Path(host_report).read_text(encoding='utf-8'))
    state = load_runtime_state(runtime_state)
    attempt = active_identity(state,slide_id)
    for key in ('attempt_id','slide_id','source_fingerprint','text_clean_fingerprint','design_id'):
        if report.get(key) != attempt[key]:
            raise ValueError('host report identity mismatch: '+key)
    if report.get('route') != 'codex-canva' or report.get('schema_version') != 1:
        raise ValueError('invalid host report route/schema')
    checked = datetime.fromisoformat(report['checked_at'].replace('Z','+00:00'))
    if checked.tzinfo is None or not 0 <= (datetime.now(timezone.utc)-checked).total_seconds() <= 1800:
        raise ValueError('host checks must be refreshed within 30 minutes')
    for key in ('connector_readonly', 'browser_session', 'pptx_option'):
        if report.get('checks', {}).get(key) != 'PASS' or not report.get('evidence', {}).get(key, '').strip():
            raise ValueError('host check missing or incomplete: ' + key)
    return 'HOST_EVIDENCE_ACCEPTED_LIVE_UNVERIFIED'

def check(skill_dir, output_dir, scope='package', acquisition_route='external',
          host_report=None, runtime_state=None, slide_id=None, *, office_host='auto'):
    if scope not in SCOPES:
        raise ValueError(f'unknown preflight scope: {scope}')
    if acquisition_route not in ('external', 'codex-canva'):
        raise ValueError('unknown acquisition route')
    if office_host not in ('auto', 'powerpoint', 'wps'):
        raise ValueError('unknown office host: ' + str(office_host))
    root = Path(skill_dir)
    result = {
        'scope': scope, 'protocol_hashes': 'NOT_EXECUTED',
        'runtime_dependencies': {}, 'test_dependencies': {},
        'output_writable': 'FAIL', 'powerpoint_com': 'NOT_REQUIRED',
        'wps_com': 'NOT_REQUIRED', 'office_host': None, 'office_progid': None, 'office_diagnostics': [],
        'image_generation': 'HOST_CAPABILITY_CHECK_REQUIRED',
        'canva_magic_layers': 'EXTERNAL_SESSION_CHECK_REQUIRED',
        'font_environment': 'PASS' if Path(os.environ.get('WINDIR','C:/Windows'),'Fonts').is_dir() else 'FAIL',
        'canva_pptx_acquisition': 'NOT_REQUIRED', 'blockers': [],
    }
    runtime_names = ('yaml',) + (('pptx','lxml','win32com') if scope in ('merge','full') else ())
    if scope == 'acquisition' and acquisition_route == 'codex-canva':
        runtime_names += ('pptx','lxml')
    if scope=='image_layer': runtime_names += ('PIL','numpy')
    for name in runtime_names:
        result['runtime_dependencies'][name] = 'PASS' if importlib.util.find_spec(name) else 'FAIL'
    for name in TEST_DEPENDENCIES:
        result['test_dependencies'][name] = 'PASS' if importlib.util.find_spec(name) else 'FAIL'
    missing = [name for name,status in result['runtime_dependencies'].items() if status == 'FAIL']
    if missing:
        dependency_stage = ('pptx_acquisition' if scope == 'acquisition' and acquisition_route == 'codex-canva'
                            else 'local_merge' if scope in ('merge','full') else 'image_layer' if scope=='image_layer' else 'package')
        result['blockers'].append({'stage':dependency_stage,
                                   'code':('IMAGE_LAYER_GEOMETRY_DEPENDENCY_MISSING' if scope=='image_layer' and any(n in ('PIL','numpy') for n in missing) else 'PYTHON_DEPENDENCY_MISSING'),'environment':True,
                                   'user_action':True,'detail':', '.join(missing)})
    if result['runtime_dependencies']['yaml'] == 'PASS':
        try:
            mismatches = validate_manifest_hashes(root)
            result['protocol_hashes'] = 'FAIL' if mismatches else 'PASS'
            if mismatches:
                result['blockers'].append({'stage':'all','code':'CANONICAL_HASH_MISMATCH',
                                           'environment':False,'user_action':False,'detail':', '.join(mismatches)})
        except Exception as exc:
            result['protocol_hashes'] = 'FAIL'
            result['blockers'].append({'stage':'all','code':'MANIFEST_INVALID',
                                       'environment':False,'user_action':False,'detail':str(exc)})
    try:
        dest = Path(output_dir)
        dest.mkdir(parents=True, exist_ok=True)
        probe = dest/'.skill-write-probe'
        probe.write_text('ok',encoding='utf-8')
        probe.unlink()
        result['output_writable'] = 'PASS'
    except OSError as exc:
        result['blockers'].append({'stage':'delivery','code':'OUTPUT_PATH_UNWRITABLE',
                                   'environment':True,'user_action':True,'detail':str(exc)})
    if scope in ('merge','full'):
        hosts = ('powerpoint', 'wps') if office_host == 'auto' else (office_host,)
        for host in hosts:
            result[host + '_com'] = 'NOT_EXECUTED'
        if result['output_writable'] == 'PASS' and not missing:
            from merge_pptx import OFFICE_PROGIDS
            for host in hosts:
                probe = _probe_powerpoint if host == 'powerpoint' else _probe_wps
                ok, detail = _safe_office_probe(probe, output_dir)
                result[host + '_com'] = 'PASS' if ok else 'FAIL'
                result['office_diagnostics'].append({'host': host, 'progid': OFFICE_PROGIDS[host], 'status': 'PASS' if ok else 'FAIL', 'detail': detail})
                if ok:
                    result.update(office_host=host, office_progid=OFFICE_PROGIDS[host])
                    break
            if result['office_host'] is None:
                for attempt in result['office_diagnostics']:
                    result['blockers'].append({'stage':'local_merge','code':attempt['host'].upper() + '_NATIVE_IMPORT_UNAVAILABLE',
                                               'environment':True,'user_action':True,'detail':attempt['detail']})
    if scope in ('acquisition','full'):
        result['acquisition_route'] = acquisition_route
        if acquisition_route == 'codex-canva':
            try:
                result['canva_pptx_acquisition'] = _check_codex_host(root, host_report, runtime_state, slide_id)
            except (OSError, ValueError, KeyError, TypeError, AttributeError) as exc:
                result['canva_pptx_acquisition'] = 'HOST_CHECK_REQUIRED'
                result['blockers'].append({'stage':'pptx_acquisition','code':'PPTX_ACQUISITION_FAILED',
                    'environment':True,'user_action':False,'detail':str(exc)})
        else:
            adapter = os.environ.get('CANVA_PPTX_ACQUISITION_ADAPTER')
            if adapter and Path(adapter).is_file():
                result['canva_pptx_acquisition'] = 'ADAPTER_PRESENT_UNVERIFIED'
            else:
                result['canva_pptx_acquisition'] = 'IMPLEMENTATION_DEPENDENCY_PENDING'
                result['blockers'].append({'stage':'pptx_acquisition','code':'PPTX_ACQUISITION_FAILED',
                    'environment':True,'user_action':True,'detail':'adapter not provided'})
    return result


def exit_code(report):
    return 1 if report['blockers'] else 0


if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('--skill-dir',required=True)
    p.add_argument('--output-dir',required=True)
    p.add_argument('--scope',choices=SCOPES,default='package')
    p.add_argument('--acquisition-route',choices=('external','codex-canva'),default='external')
    p.add_argument('--host-report')
    p.add_argument('--runtime-state')
    p.add_argument('--slide-id')
    p.add_argument('--office-host',choices=('auto','powerpoint','wps'),default='auto')
    a = p.parse_args()
    report = check(a.skill_dir,a.output_dir,a.scope,a.acquisition_route,a.host_report,a.runtime_state,a.slide_id,office_host=a.office_host)
    print(json.dumps(report,ensure_ascii=False,indent=2))
    raise SystemExit(exit_code(report))
