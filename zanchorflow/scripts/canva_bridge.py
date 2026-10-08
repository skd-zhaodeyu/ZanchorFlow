"""Codex provenance wiring. No connector calls or automatic remote retries."""
import argparse
import copy
from contextlib import contextmanager
import hashlib
import json
import os
from pathlib import Path
import tempfile
import uuid
from urllib.parse import urlsplit
import zipfile
import xml.etree.ElementTree as ET
import runtime as r


def sha(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def resolve(state, value):
    path = Path(value)
    return path if path.is_absolute() else Path(state['_base_dir'])/path


@contextmanager
def transaction(state_path):
    """Helpers operate on a same-directory copy; commit once, never stale rollback."""
    path = Path(state_path).resolve()
    lock = path.with_name(path.name+'.bridge-lock')
    fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    os.close(fd)
    temporary = None
    try:
        before = path.read_bytes()
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix='.json',
                                         prefix='.bridge-state-', delete=False) as stream:
            temporary = Path(stream.name)
            stream.write(before)
        yield temporary
        r.load_runtime_state(temporary)
        if path.read_bytes() != before:
            raise ValueError('runtime state changed during operation; stop and inspect')
        os.replace(temporary, path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        lock.unlink(missing_ok=True)


def init(state_path, route='codex-canva'):
    path = Path(state_path).resolve()
    if route not in ('codex-canva','external'):
        raise ValueError('invalid acquisition route')
    if path.exists():
        r.load_runtime_state(path)
        return {'status':'EXISTING_STATE_UNCHANGED'}
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('x',encoding='utf-8') as stream:
        json.dump({'slides':{},'canva_bridge':{'route':route,'schema_version':1}},stream)
    return {'status':'INITIALIZED','route':route}


def start_run(state_path, title_choice_required=False):
    """Wrap the original start; archive old page state only on a real run change."""
    with transaction(state_path) as temp:
        before = r.load_runtime_state(temp)
        previous = before.get('stage2_run',{})
        if before.get('canva_bridge',{}).get('route') not in ('codex-canva','external'):
            raise ValueError('route missing; do not infer a legacy task route')
        run = r.start_stage2_run(temp, title_choice_required=title_choice_required)
        if previous.get('run_id') != run['run_id']:
            state = r.load_runtime_state(temp)
            bridge = state['canva_bridge']
            if previous:
                bridge.setdefault('run_history',[]).append({
                    'run_id':previous['run_id'],'slides':copy.deepcopy(before['slides']),
                    'bridge_records':{key:copy.deepcopy(before.get('canva_bridge',{}).get(key))
                                      for key in ('active','downloads','validated','merged','text_plan','text_plan_history','text_preparation','text_preparation_history','upload_observations','upload_results','design_observations','transport_requirements','observation_history')},
                    'deck_records':{key:copy.deepcopy(before.get(key))
                                    for key in ('deck_order','merged_deck','validated_deck')}})
            state['slides'] = {sid:{} for sid in run['page_order']}
            for key in ('active','downloads','validated','merged','text_plan','text_plan_history','text_preparation','text_preparation_history','upload_observations','upload_results','design_observations','transport_requirements','observation_history'):
                bridge.pop(key,None)
            for key in ('deck_order','merged_deck','validated_deck'):
                state.pop(key,None)
            r._save_runtime_state(temp,state)
    return run


def _bridge(state):
    bridge = state.get('canva_bridge')
    if not isinstance(bridge,dict) or bridge.get('route') != 'codex-canva':
        raise ValueError('Codex route not registered; do not infer or switch legacy routes')
    return bridge


def _source_fingerprint(state, slide_id):
    slide = state['slides'][slide_id]
    truth = json.loads(resolve(state,slide['final_content_truth']).read_text(encoding='utf-8'))
    canvas = json.loads(resolve(state,slide['canvas']).read_text(encoding='utf-8'))
    return r.source_fingerprint(resolve(state,slide['approved_render']).read_bytes(),truth,canvas)


def _source(state, slide_id):
    return (_source_fingerprint(state,slide_id),
            r.text_clean_fingerprint(resolve(state,state['slides'][slide_id]['text_clean']).read_bytes()))


def _text_plan_source_hashes(state, slide_id):
    slide = state['slides'][slide_id]
    return {
        'approved_render_sha256': sha(resolve(state, slide['approved_render'])),
        'final_content_truth_sha256': sha(resolve(state, slide['final_content_truth'])),
    }


def _check_text_plan(state, slide_id):
    bridge = _bridge(state)
    run = state['stage2_run']
    record = bridge.get('text_plan', {}).get(slide_id)
    if not isinstance(record, dict):
        raise ValueError('STAGE3_TEXT_PLAN_REQUIRED: validated Finalized Text Manifest + Removal Inventory required')
    manifest_reg = run.get('artifacts', {}).get('finalized_manifest:' + slide_id)
    inventory_reg = run.get('artifacts', {}).get('removal_inventory:' + slide_id)
    if not isinstance(manifest_reg, dict) or not isinstance(inventory_reg, dict):
        raise ValueError('STAGE3_TEXT_PLAN_REQUIRED: finalized_manifest/removal_inventory registration missing')
    expected_sources = _text_plan_source_hashes(state, slide_id)
    try:
        manifest = json.loads(resolve(state, manifest_reg['path']).read_text(encoding='utf-8'))
        inventory = json.loads(resolve(state, inventory_reg['path']).read_text(encoding='utf-8'))
        truth = json.loads(resolve(state, state['slides'][slide_id]['final_content_truth']).read_text(encoding='utf-8'))
    except (OSError, ValueError, KeyError, TypeError) as exc:
        raise ValueError('STAGE3_TEXT_PLAN_REQUIRED: current plan files unavailable') from exc
    r.validate_stage3_text_plan(manifest, inventory, truth, slide_id)
    r.titles.validate_manifest(state, slide_id, manifest)
    expected = {
        'slide_id': slide_id,
        'run_id': run['run_id'],
        'approved_outline_fingerprint': run['approved_outline_fingerprint'],
        **expected_sources,
        'manifest_sha256': manifest_reg.get('sha256'),
        'inventory_sha256': inventory_reg.get('sha256'),
        'text_plan_fingerprint': r.stage3_text_plan_fingerprint(manifest, inventory),
    }
    if (record != expected
            or manifest_reg.get('run_id') != run['run_id']
            or inventory_reg.get('run_id') != run['run_id']
            or manifest_reg.get('approved_outline_fingerprint') != run['approved_outline_fingerprint']
            or inventory_reg.get('approved_outline_fingerprint') != run['approved_outline_fingerprint']
            or manifest_reg.get('sha256') != sha(resolve(state, manifest_reg['path']))
            or inventory_reg.get('sha256') != sha(resolve(state, inventory_reg['path']))):
        raise ValueError('STAGE3_TEXT_PLAN_REQUIRED: plan/source binding changed')
    return expected


def _invalidate_after_text_plan_change(state, slide_id):
    run = state['stage2_run']
    for base_kind in ('text_clean','reconstruction_attempt','graphics_first_pptx','font_fallback'):
        key = base_kind + ':' + slide_id
        old = run.setdefault('artifacts', {}).pop(key, None)
        if old is not None:
            run.setdefault('artifact_history', {}).setdefault(key, []).append(old)
        state.setdefault('slides', {}).setdefault(slide_id, {}).pop(base_kind, None)
    slide = state.setdefault('slides', {}).setdefault(slide_id, {})
    old_seal = slide.pop('validated_single_page', None)
    if old_seal is not None:
        slide.setdefault('superseded_validated_single_pages', []).append(old_seal)
    bridge = _bridge(state)
    for key in ('active','downloads','validated','text_preparation','upload_observations','upload_results','design_observations','transport_requirements'):
        section = bridge.get(key)
        if isinstance(section, dict):
            section.pop(slide_id, None)
    bridge.pop('merged', None)
    state.pop('merged_deck', None)
    state.pop('validated_deck', None)


def register_text_plan(state_path, slide_id, manifest_path, inventory_path):
    """Atomically validate and bind the destructive-cleanup plan before Text-Clean creation."""
    manifest_path = Path(manifest_path).resolve()
    inventory_path = Path(inventory_path).resolve()
    with transaction(state_path) as temp:
        r.require_stage2_entry(temp)
        r.require_stage2_visual_approval(temp)
        state = r.load_runtime_state(temp)
        bridge = _bridge(state)
        if (r.current_stage2_approved_render(temp, slide_id) is None
                or not r._stage2_reconciliation_current(temp, slide_id)
                or not r.stage2_artifact_current(temp, 'final_content_truth:' + slide_id)
                or not r.stage2_artifact_current(temp, 'canvas:' + slide_id)):
            raise ValueError('STAGE3_TEXT_PLAN_REQUIRED: current render/truth/reconciliation/canvas required')
        try:
            manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
            inventory = json.loads(inventory_path.read_text(encoding='utf-8'))
            truth = json.loads(resolve(state, state['slides'][slide_id]['final_content_truth']).read_text(encoding='utf-8'))
        except (OSError, ValueError, KeyError, TypeError) as exc:
            raise ValueError('STAGE3_TEXT_PLAN_INVALID: JSON inputs required') from exc
        r.titles.validate_manifest(state, slide_id, manifest)
        summary = r.validate_stage3_text_plan(manifest, inventory, truth, slide_id)
        plan_fp = r.stage3_text_plan_fingerprint(manifest, inventory)
        sources = _text_plan_source_hashes(state, slide_id)
        current = bridge.get('text_plan', {}).get(slide_id)
        if (isinstance(current, dict) and current.get('text_plan_fingerprint') == plan_fp
                and current.get('approved_render_sha256') == sources['approved_render_sha256']
                and current.get('final_content_truth_sha256') == sources['final_content_truth_sha256']
                and r.stage2_artifact_current(temp, 'finalized_manifest:' + slide_id)
                and r.stage2_artifact_current(temp, 'removal_inventory:' + slide_id)):
            return {**summary,**current,'status':'TEXT_PLAN_REGISTERED','slide_id':slide_id}
        if current is not None:
            bridge.setdefault('text_plan_history', {}).setdefault(slide_id, []).append(copy.deepcopy(current))
        _invalidate_after_text_plan_change(state, slide_id)
        r._save_runtime_state(temp, state)
        manifest_reg = r.register_stage2_artifact(temp, 'finalized_manifest:' + slide_id, manifest_path)
        inventory_reg = r.register_stage2_artifact(temp, 'removal_inventory:' + slide_id, inventory_path)
        state = r.load_runtime_state(temp)
        bridge = _bridge(state)
        run = state['stage2_run']
        record = {
            'slide_id': slide_id,
            'run_id': run['run_id'],
            'approved_outline_fingerprint': run['approved_outline_fingerprint'],
            **_text_plan_source_hashes(state, slide_id),
            'manifest_sha256': manifest_reg['sha256'],
            'inventory_sha256': inventory_reg['sha256'],
            'text_plan_fingerprint': plan_fp,
        }
        bridge.setdefault('text_plan', {})[slide_id] = record
        r._save_runtime_state(temp, state)
        _check_text_plan(r.load_runtime_state(temp), slide_id)
    return {**summary,**record,'status':'TEXT_PLAN_REGISTERED','slide_id':slide_id}


def _check_text_preparation(state, slide_id):
    bridge = _bridge(state)
    run = state['stage2_run']
    binding = bridge.get('text_preparation',{}).get(slide_id)
    source_fp,clean_fp = _source(state,slide_id)
    plan = _check_text_plan(state, slide_id)
    record = run.get('artifacts',{}).get('text_clean:'+slide_id)
    expected = {'slide_id':slide_id,'run_id':run['run_id'],
                'approved_outline_fingerprint':run['approved_outline_fingerprint'],
                'source_fingerprint':source_fp,'text_clean_fingerprint':clean_fp,
                'text_plan_fingerprint':plan['text_plan_fingerprint'],
                'registration':record}
    if (slide_id not in run['page_order'] or not isinstance(binding,dict)
            or binding != expected or not isinstance(record,dict)
            or record.get('run_id') != run['run_id']
            or record.get('approved_outline_fingerprint') != run['approved_outline_fingerprint']
            or resolve(state,record.get('path','')) != resolve(state,state['slides'][slide_id]['text_clean'])
            or record.get('sha256') != clean_fp):
        raise ValueError('current text preparation source binding missing or changed: '+slide_id)
    return source_fp,clean_fp


def register_text_clean(state_path, slide_id, artifact):
    """Record locally checked formal text preparation without changing attempts."""
    with transaction(state_path) as temp:
        state = r.load_runtime_state(temp)
        _bridge(state)
        r.require_stage2_entry(temp)
        r.require_stage2_visual_approval(temp)
        if (r.current_stage2_approved_render(temp,slide_id) is None
                or not r._stage2_reconciliation_current(temp,slide_id)
                or not r.stage2_artifact_current(temp,'canvas:'+slide_id)):
            raise ValueError('current page source, canvas and text reconciliation required')
        plan = _check_text_plan(state, slide_id)
        source_fp = _source_fingerprint(state,slide_id)
        record = r.register_stage2_artifact(temp,'text_clean:'+slide_id,artifact)
        state = r.load_runtime_state(temp)
        bridge = _bridge(state)
        old = bridge.setdefault('text_preparation',{}).get(slide_id)
        if old is not None:
            bridge.setdefault('text_preparation_history',{}).setdefault(slide_id,[]).append(copy.deepcopy(old))
        run = state['stage2_run']
        bridge['text_preparation'][slide_id] = {
            'slide_id':slide_id,'run_id':run['run_id'],
            'approved_outline_fingerprint':run['approved_outline_fingerprint'],
            'source_fingerprint':source_fp,
            'text_clean_fingerprint':r.text_clean_fingerprint(resolve(state,state['slides'][slide_id]['text_clean']).read_bytes()),
            'text_plan_fingerprint':plan['text_plan_fingerprint'],
            'registration':copy.deepcopy(record)}
        _check_text_preparation(state,slide_id)
        r._save_runtime_state(temp,state)
    return {'status':'TEXT_CLEAN_REGISTERED','slide_id':slide_id,**record}


def _active(state, slide_id):
    bridge = _bridge(state)
    active = bridge.get('active',{}).get(slide_id)
    if not isinstance(active,dict):
        raise ValueError('current attempt missing: '+slide_id)
    run = state['stage2_run']
    source_fp, clean_fp = _check_text_preparation(state,slide_id)
    if (slide_id not in run['page_order'] or active.get('run_id') != run['run_id']
            or active.get('approved_outline_fingerprint') != run['approved_outline_fingerprint']
            or active.get('source_fingerprint') != source_fp
            or active.get('text_clean_fingerprint') != clean_fp):
        raise ValueError('current attempt lineage changed: '+slide_id)
    return active


def active_identity(state, slide_id):
    active = _active(state,slide_id)
    if active.get('status') != 'ACCEPT' or not active.get('design_id'):
        raise ValueError('attempt pending or not accepted; do not repeat remote call: '+slide_id)
    slide = state['slides'][slide_id]
    path = resolve(state,slide['reconstruction_attempt'])
    record = state['stage2_run']['artifacts'].get('reconstruction_attempt:'+slide_id,{})
    if (record.get('run_id') != active['run_id']
            or record.get('approved_outline_fingerprint') != active['approved_outline_fingerprint']
            or resolve(state,record.get('path','')) != path or record.get('sha256') != sha(path)):
        raise ValueError('current attempt registration missing or changed')
    attempt = json.loads(path.read_text(encoding='utf-8'))
    keys = ('attempt_id','slide_id','run_id','approved_outline_fingerprint',
            'source_fingerprint','text_clean_fingerprint','design_id')
    if attempt.get('assessment') != 'ACCEPT' or any(attempt.get(k) != active.get(k) for k in keys):
        raise ValueError('active design identity mismatch')
    return {k:active[k] for k in keys}


def register_upload(state_path, slide_id, image_path):
    import download_evidence as de
    with transaction(state_path) as temp:
        state=r.load_runtime_state(temp);active=_active(state,slide_id)
        if active['status']!='PENDING':raise ValueError('upload registration requires current PENDING attempt')
        image=Path(image_path).resolve()
        expected=resolve(state,state['slides'][slide_id]['text_clean']).resolve()
        if image!=expected or sha(image)!=active['text_clean_fingerprint']:
            raise ValueError('actual upload must be the current registered Text-Clean bytes')
        identity={k:active[k] for k in ('attempt_id','slide_id','run_id','approved_outline_fingerprint',
                                      'source_fingerprint','text_clean_fingerprint')}
        value={'status':'UPLOAD_REGISTERED','schema_version':2,'identity':identity,'image_path':str(image),'image_sha256':sha(image),
               'requested_title':de.title_for(slide_id,active['attempt_id'],sha(image))}
        existing=state['canva_bridge'].setdefault('upload_observations',{}).get(slide_id)
        if existing and de.sealed(existing)==value:return value
        path=Path(state_path).resolve().parent/('upload-'+active['attempt_id']+'.json')
        ref=de.write_once(path,value)
        state['canva_bridge']['upload_observations'][slide_id]=ref;r._save_runtime_state(temp,state)
    return value


def register_upload_result(state_path,slide_id,capture_path):
    import download_evidence as de
    with transaction(state_path) as temp:
        state=r.load_runtime_state(temp);active=_active(state,slide_id)
        if active['status']!='PENDING':raise ValueError('result capture requires the current PENDING upload')
        upload_ref=state['canva_bridge'].get('upload_observations',{}).get(slide_id)
        upload=de.sealed(upload_ref);capture_path=Path(capture_path).resolve();capture=de.read(capture_path)
        tool=str(capture.get('tool',''))
        if not tool.endswith('image_to_design'):raise ValueError('actual Magic tool capture required')
        args=capture.get('call_args',{})
        if (Path(args.get('image_file','')).resolve()!=Path(upload['image_path']).resolve()
            or args.get('title')!=upload['requested_title'] or sha(upload['image_path'])!=upload['image_sha256']):
            raise ValueError('actual tool arguments differ from registered image/title')
        returned=de.returned_design_id(capture['result'])
        value={'identity':upload['identity'],'upload':upload_ref,'returned_design_id':returned,
               'capture':{'path':str(capture_path),'sha256':de.sha(capture_path)}}
        path=Path(state_path).resolve().parent/('magic-return-'+active['attempt_id']+'.json')
        ref={'path':str(path),'sha256':de.sha(path)} if path.exists() else de.write_once(path,value)
        if de.sealed(ref)!=value:raise ValueError('original Magic result cannot be overwritten')
        state['canva_bridge'].setdefault('upload_results',{})[slide_id]=ref;r._save_runtime_state(temp,state)
    return {'status':'UPLOAD_RESULT_REGISTERED','design_id':returned}


def register_design_observation(state_path, slide_id, lookup_path, title_observation=None):
    import download_evidence as de
    with transaction(state_path) as temp:
        state=r.load_runtime_state(temp);identity=active_identity(state,slide_id)
        lookup_path=Path(lookup_path).resolve();lookup=json.loads(lookup_path.read_text(encoding='utf-8-sig'))
        _lookup_url(lookup,identity)
        design=lookup['response']['design'];title=design.get('title')
        for other_sid,other in state['canva_bridge'].get('active',{}).items():
            if other_sid!=slide_id and other.get('status')=='ACCEPT' and other.get('design_id')==identity['design_id']:
                raise ValueError('same design_id assigned to multiple current pages')
        expected_title=de.title_for(slide_id,identity['attempt_id'],identity['text_clean_fingerprint'])
        host_title_ref=None
        if title is None and title_observation is not None:
            hp=Path(title_observation).resolve();host=de.read(hp);de.check_identity(host['identity'],identity)
            de.design_url(host['observed_url'],identity['design_id'])
            if host.get('source')!='canva_host.readonly' or not isinstance(host.get('evidence'),str) or not host['evidence'].strip():
                raise ValueError('actual readonly Host title evidence required')
            capture=de.sealed(host['capture'])
            if capture.get('observed_url')!=host['observed_url'] or capture.get('actual_title')!=host.get('actual_title'):
                raise ValueError('Host title capture mismatch')
            host_title_ref={'path':str(hp),'sha256':de.sha(hp)};title=host['actual_title']
        if design.get('page_count')!=1:
            raise ValueError('confirm one page on the current design_id; do not reconstruct')
        upload_ref=state['canva_bridge'].get('upload_observations',{}).get(slide_id)
        if upload_ref:
            upload=de.sealed(upload_ref)
            if any(upload['identity'].get(k)!=identity[k] for k in upload['identity']):raise ValueError('upload identity changed')
            if upload['image_sha256']!=identity['text_clean_fingerprint']:
                raise ValueError('upload bytes changed')
        value={'schema_version':2,'identity':identity,'actual_title':title,
               'lookup':{'path':str(lookup_path),'sha256':sha(lookup_path)},
               'updated_at':design.get('updated_at'),'upload':upload_ref,
               'legacy_attempt':upload_ref is None,'host_title':host_title_ref}
        path=Path(state_path).resolve().parent/('design-'+identity['attempt_id']+'-'+sha(lookup_path)[:12]+'.json')
        ref={'path':str(path),'sha256':de.sha(path)} if path.exists() else de.write_once(path,value)
        if de.sealed(ref)!=value:raise ValueError('immutable design observation collision')
        state['canva_bridge'].setdefault('design_observations',{})[slide_id]=ref;r._save_runtime_state(temp,state)
    return {'status':'DESIGN_OBSERVATION_REGISTERED','identity':identity,'actual_title':title}


def validate_completion_evidence(state, identity, payload, pptx, required_version=None):
    import download_evidence as de
    requirement=state.get('canva_bridge',{}).get('transport_requirements',{}).get(identity['slide_id'])
    if requirement and requirement.get('identity')!=identity:requirement=None
    if requirement and requirement.get('terminated'):raise ValueError('transport acquisition explicitly terminated')
    parsed=payload
    if isinstance(payload,str):
        try:parsed=json.loads(payload)
        except (ValueError,TypeError):parsed=None
    if (isinstance(parsed,dict) and 'transport_receipt' in parsed) or requirement:
        if not isinstance(parsed,dict) or not parsed.get('transport_receipt'):
            raise ValueError('TRANSPORT_RECEIPT_REQUIRED: new acquisition cannot downgrade')
        if not requirement or requirement.get('acquisition_id')!=parsed.get('acquisition_id'):
            raise ValueError('transport acquisition not registered')
        if parsed['transport_receipt'] not in requirement.get('receipts',[]):
            raise ValueError('transport receipt not registered')
        snap=de.validate_transport_proof(parsed,identity,pptx)
        current=state['canva_bridge'].get('design_observations',{}).get(identity['slide_id'])
        if current!=snap['context']['design_observation']:raise ValueError('current design observation changed')
        body=dict(parsed);expected=body.pop('evidence_digest',None)
        if expected!=__import__('hashlib').sha256(json.dumps(body,sort_keys=True,separators=(',',':')).encode()).hexdigest():
            raise ValueError('completion evidence digest mismatch')
        return
    if isinstance(payload,str):
        try:payload=json.loads(payload)
        except (ValueError,TypeError):
            if required_version==2:raise ValueError('v2 completion evidence cannot be downgraded')
            return  # Existing v1 opaque evidence stays on the legacy route.
    if isinstance(payload,dict) and payload.get('schema_version') not in (None,1,2):raise ValueError('unsupported completion evidence version')
    if not isinstance(payload,dict) or payload.get('schema_version')!=2:
        if required_version==2:raise ValueError('v2 completion evidence cannot be downgraded')
        return
    snap=de.validate_file_proof(payload,identity,pptx)
    de.source_association(payload,identity,snap['context'],pptx)
    # Observation must still be the one registered for this current attempt.
    current=state.get('canva_bridge',{}).get('design_observations',{}).get(identity['slide_id'])
    if current!=snap['context']['design_observation']:raise ValueError('current design observation changed')
    if not payload.get('page_check'):raise ValueError('PAGE_CHECK_REQUIRED: downloaded file retained; event is optional')
    review=de.sealed(payload['page_check']);de.check_identity(review['identity'],identity)
    if review.get('assessment')!='MATCH' or review.get('file_sha256')!=payload['selected']['sha256'] or review.get('input_sha256')!=identity['text_clean_fingerprint']:
        raise ValueError('page comparison identity mismatch')
    preview=de.sealed(review['preview_receipt']);de.check_identity(preview['identity'],identity)
    if preview['source_sha256']!=payload['selected']['sha256'] or de.sha(preview['preview_path'])!=preview['preview_sha256']:
        raise ValueError('page preview changed')
    input_ref=review['input_render']
    if de.sha(input_ref['path'])!=identity['text_clean_fingerprint'] or input_ref['sha256']!=identity['text_clean_fingerprint']:
        raise ValueError('page comparison source changed')
    expected=payload.get('evidence_digest')
    copy_payload=dict(payload);copy_payload.pop('evidence_digest',None)
    actual=__import__('hashlib').sha256(json.dumps(copy_payload,sort_keys=True,separators=(',',':')).encode()).hexdigest()
    if expected!=actual:raise ValueError('completion evidence digest mismatch')


def begin_attempt(state_path, slide_id, rebuild=False, recovery_evidence=None):
    r.require_stage2_entry(state_path)
    if r.stage2_handoff_status(state_path)['status'] != 'COMPLETE':
        raise ValueError('current Stage 2 handoff incomplete')
    current_status = status(state_path).get('status')
    with transaction(state_path) as temp:
        state = r.load_runtime_state(temp)
        bridge = _bridge(state)
        source_fp,clean_fp = _check_text_preparation(state,slide_id)
        for kind in ('canvas','text_clean'):
            if not r.stage2_artifact_current(temp,kind+':'+slide_id):
                raise ValueError('current registered '+kind+' required')
        old = bridge.setdefault('active',{}).get(slide_id)
        if old and not rebuild:
            raise ValueError('attempt already exists; resume it; new rebuild needs explicit request')
        if old and old.get('status') == 'PENDING':
            raise ValueError('previous completion uncertain; diagnose before any new remote call')
        if (old and rebuild and current_status == 'RESTORE_TEXT'
                and isinstance(bridge.get('downloads', {}).get(slide_id), dict)):
            raise ValueError('current Graphics-first PPTX already bound; resume text restoration from the bound PPTX')
        if old and old.get('status') == 'TOOL_FAILED':
            if not recovery_evidence:
                raise ValueError('confirmed failure needs observed environment repair evidence')
            repair = json.loads(Path(recovery_evidence).read_text(encoding='utf-8'))
            if (repair.get('attempt_id') != old['attempt_id']
                    or repair.get('failure_code') != old.get('failure_code')
                    or repair.get('terminal_confirmed') is not True
                    or not isinstance(repair.get('change_evidence'),str)
                    or not repair['change_evidence'].strip()):
                raise ValueError('recovery evidence must match terminal failure and actual environment change')
        if old:
            bridge.setdefault('history',{}).setdefault(slide_id,[]).append(copy.deepcopy(old))
        run = state['stage2_run']
        retry_state = old.get('retry_state','available') if old and old.get('source_fingerprint') == source_fp and old.get('text_clean_fingerprint') == clean_fp else 'available'
        active = {'retry_state':retry_state,'attempt_id':uuid.uuid4().hex,'slide_id':slide_id,'run_id':run['run_id'],
                  'approved_outline_fingerprint':run['approved_outline_fingerprint'],
                  'source_fingerprint':source_fp,'text_clean_fingerprint':clean_fp,'status':'PENDING'}
        bridge['active'][slide_id] = active
        for section in ('upload_observations','upload_results','design_observations','transport_requirements'):
            previous=bridge.get(section,{}).pop(slide_id,None)
            if previous:bridge.setdefault('observation_history',[]).append({'identity':old,'kind':section,'ref':previous})
        # Old sources and seals remain history; provenance checks stop their delivery.
        r._save_runtime_state(temp,state)
    return active


def accept_attempt(state_path, slide_id, attempt_id, result_path):
    result = json.loads(Path(result_path).read_text(encoding='utf-8'))
    with transaction(state_path) as temp:
        r.require_stage2_entry(temp)
        state = r.load_runtime_state(temp)
        active = _active(state,slide_id)
        if active['attempt_id'] != attempt_id or active['status'] not in ('PENDING','ASSISTED_CLEANUP_REQUIRED'):
            raise ValueError('result must belong to the pending attempt')
        required = ('attempt_id','slide_id','source_fingerprint','text_clean_fingerprint')
        if any(result.get(k) != active[k] for k in required):
            raise ValueError('returned result identity mismatch')
        assessment = result.get('assessment')
        if state['canva_bridge'].get('upload_observations',{}).get(slide_id) and assessment not in ('MAGIC_LAYERS_EXECUTION_FAILED','MAGIC_LAYERS_TOOL_UNAVAILABLE'):
            import download_evidence as de
            returned=de.sealed(state['canva_bridge'].get('upload_results',{}).get(slide_id))
            if returned['upload']!=state['canva_bridge']['upload_observations'][slide_id]:raise ValueError('Magic result belongs to another upload')
            captured=de.sealed(returned['capture'])
            if result.get('design_id')!=de.returned_design_id(captured['result']) or result.get('design_id')!=returned['returned_design_id']:
                raise ValueError('accepted design_id differs from the actual Magic return')
        if assessment in ('MAGIC_LAYERS_EXECUTION_FAILED','MAGIC_LAYERS_TOOL_UNAVAILABLE'):
            if (active['status'] != 'PENDING' or result.get('terminal_confirmed') is not True
                    or result.get('design_id') or not isinstance(result.get('evidence'),str)
                    or not result['evidence'].strip()):
                raise ValueError('confirmed terminal tool failure without design needs actual evidence')
            active.update(status='TOOL_FAILED',failure_code=assessment,
                          terminal_confirmed=True,failure_evidence=result['evidence'])
        else:
            if (assessment not in ('ACCEPT','RETRY_ONCE','ASSISTED_CLEANUP_REQUIRED','RETURN_TO_STAGE_2') or not isinstance(result.get('design_id'),str)
                    or not result['design_id'].strip() or not isinstance(result.get('evidence'),str)
                    or not result['evidence'].strip()):
                raise ValueError('actual structural assessment, design ID and evidence required')
            retry_state = active.get('retry_state','available')
            if active['status'] == 'PENDING' and retry_state == 'pending':
                retry_state = r.retry_transition(retry_state,'assessable_result')
            if assessment == 'RETRY_ONCE':
                retry_state = r.retry_transition(retry_state,'structural_failure')
                if retry_state != 'pending':
                    raise ValueError('structural retry already consumed')
            if active['status'] == 'ASSISTED_CLEANUP_REQUIRED' and assessment == 'ASSISTED_CLEANUP_REQUIRED':
                raise ValueError('ASSISTED_CLEANUP_UNRESOLVED; stop repeated cleanup')
            active.update(status=assessment,design_id=result['design_id'],retry_state=retry_state)
        attempt = {**active,'assessment':assessment,'evidence':result['evidence']}
        metadata = Path(state_path).resolve().parent/('attempt-'+attempt_id+'.json')
        if metadata.exists():
            if json.loads(metadata.read_text(encoding='utf-8')) != attempt:
                metadata = Path(state_path).resolve().parent/('attempt-'+attempt_id+'-'+uuid.uuid4().hex+'.json')
                with metadata.open('x',encoding='utf-8') as stream:
                    json.dump(attempt,stream,ensure_ascii=False,indent=2)
        else:
            with metadata.open('x',encoding='utf-8') as stream:
                json.dump(attempt,stream,ensure_ascii=False,indent=2)
        r._save_runtime_state(temp,state)
        r.register_stage2_artifact(temp,'reconstruction_attempt:'+slide_id,metadata)
    return {'status':active['status'],'attempt_id':attempt_id,'design_id':result.get('design_id')}


def single_page(path):
    try:
        from pptx import Presentation
        from merge_pptx import inspect_ooxml_dependencies
    except ImportError as error:
        raise ValueError('PPTX_VALIDATION_DEPENDENCY_MISSING: '+str(error)) from error
    try:
        with zipfile.ZipFile(path) as archive:
            if archive.testzip():
                raise ValueError('invalid PPTX archive')
            required = {'[Content_Types].xml','_rels/.rels','ppt/presentation.xml',
                        'ppt/_rels/presentation.xml.rels'}
            if not required.issubset(archive.namelist()):
                raise ValueError('required presentation parts or relationships missing')
        errors = inspect_ooxml_dependencies(path)
        if errors:
            raise ValueError('; '.join(errors))
        presentation = Presentation(path)
        if len(presentation.slides) != 1:
            raise ValueError('expected exactly one presentation page')
        # Resolve the slide relationship and parse the actual page, not just its ID.
        list(presentation.slides[0].shapes)
    except Exception as error:
        raise ValueError('GRAPHICS_FIRST_PPTX_INVALID: '+str(error)) from error


def _official_design_url(url, identity):
    """Validate only official HTTPS design paths; alias identity needs a separate binding."""
    if (not isinstance(url, str) or not url.strip() or url != url.strip()
            or any(ord(ch) < 33 or ord(ch) == 127 for ch in url)):
        raise ValueError('observed Canva design URL required')
    parsed = urlsplit(url)
    port = parsed.port  # Validate numeric/range syntax before trusting the URL.
    if (parsed.scheme != 'https' or parsed.hostname not in ('canva.com', 'www.canva.com')
            or parsed.username is not None or parsed.password is not None
            or port not in (None, 443)):
        raise ValueError('download edit URL identity mismatch')
    parts = parsed.path.split('/')
    if len(parts) >= 3 and parts[1] == 'design' and parts[2] == identity['design_id']:
        return 'design'
    if (len(parts) == 3 and parts[1] == 'd' and parts[2]
            and all(ch.isascii() and (ch.isalnum() or ch in '_-') for ch in parts[2])):
        return 'alias'
    raise ValueError('download edit URL identity mismatch')


def _lookup_url(observation, identity):
    """Parse a trusted controller-captured readonly connector response, never download claims."""
    if (not isinstance(observation, dict) or observation.get('schema_version') != 1
            or observation.get('source') != 'canva_connector.get_design'
            or observation.get('identity') != identity
            or not isinstance(observation.get('observation'), str)
            or not observation['observation'].strip()):
        raise ValueError('readonly design lookup identity/evidence mismatch')
    response = observation.get('response')
    design = response.get('design') if isinstance(response, dict) else None
    if not isinstance(design, dict) or design.get('id') != identity['design_id']:
        raise ValueError('readonly connector design identity mismatch')
    urls = design.get('urls')
    if not isinstance(urls, dict):
        raise ValueError('readonly connector URL response missing')
    url = urls.get('edit_url')
    _official_design_url(url, identity)
    return url


def _registered_design_url(state, slide_id, identity):
    record = _bridge(state).get('design_urls', {}).get(slide_id)
    if not isinstance(record, dict) or record.get('identity') != identity:
        raise ValueError('current readonly design URL association missing or changed')
    path = resolve(state, record['lookup_path'])
    if record.get('lookup_sha256') != sha(path):
        raise ValueError('readonly design lookup evidence changed')
    url = _lookup_url(json.loads(path.read_text(encoding='utf-8')), identity)
    if record.get('edit_url') != url:
        raise ValueError('readonly design URL association mismatch')
    return url


def register_design_url(state_path, slide_id, lookup_path):
    """Bind an existing ACCEPT design to independently captured get_design evidence; no remote calls."""
    path = Path(lookup_path).resolve()
    with transaction(state_path) as temp:
        r.require_stage2_entry(temp)
        state = r.load_runtime_state(temp)
        identity = active_identity(state, slide_id)
        observation = json.loads(path.read_text(encoding='utf-8'))
        url = _lookup_url(observation, identity)
        downloads = _bridge(state).get('downloads', {}).get(slide_id)
        if (isinstance(downloads, dict) and downloads.get('identity') == identity
                and downloads.get('evidence', {}).get('edit_url') != url):
            raise ValueError('bound download URL cannot be replaced without actual historic evidence')
        _bridge(state).setdefault('design_urls', {})[slide_id] = {
            'identity': identity, 'edit_url': url,
            'lookup_path': str(path), 'lookup_sha256': sha(path)}
        r._save_runtime_state(temp, state)
        _registered_design_url(state, slide_id, identity)
    return {'status': 'DESIGN_URL_REGISTERED', 'identity': identity, 'edit_url': url}


def _check_download_url(evidence, identity, state=None):
    url = evidence.get('edit_url') if isinstance(evidence, dict) else None
    kind = _official_design_url(url, identity)
    if any(evidence.get(k) != identity[k] for k in ('attempt_id', 'slide_id', 'design_id')):
        raise ValueError('download evidence identity mismatch')
    if kind == 'alias':
        if state is None or _registered_design_url(state, identity['slide_id'], identity) != url:
            raise ValueError('download alias lacks current trusted design URL association')


def download_current(state, slide_id):
    identity = active_identity(state,slide_id)
    record = state.get('canva_bridge',{}).get('downloads',{}).get(slide_id)
    if not isinstance(record,dict) or record.get('identity') != identity:
        raise ValueError('current design download binding missing: '+slide_id)
    _check_download_url(record.get('evidence'),identity,state)
    path = resolve(state,record['path'])
    if record.get('sha256') != sha(path) or resolve(state,state['slides'][slide_id].get('graphics_first_pptx','')) != path:
        raise ValueError('download changed or path mismatch: '+slide_id)
    registered = state['stage2_run']['artifacts'].get('graphics_first_pptx:'+slide_id,{})
    if registered.get('sha256') != record['sha256'] or registered.get('run_id') != identity['run_id']:
        raise ValueError('download registration mismatch')
    single_page(path)
    if record.get('require_transport_receipt'):
        requirement=state.get('canva_bridge',{}).get('transport_requirements',{}).get(slide_id)
        if not requirement or requirement.get('acquisition_id')!=record.get('acquisition_id'):
            raise ValueError('bound transport requirement missing')
    validate_completion_evidence(state,identity,record['evidence'].get('completion_evidence'),path,record.get('completion_schema'))
    return record


def bind_download(state_path, slide_id, pptx, evidence_path):
    path = Path(pptx).resolve()
    single_page(path)
    evidence = json.loads(Path(evidence_path).read_text(encoding='utf-8'))
    with transaction(state_path) as temp:
        r.require_stage2_entry(temp)
        state = r.load_runtime_state(temp)
        identity = active_identity(state,slide_id)
        if any(evidence.get(k) != identity[k] for k in ('attempt_id','slide_id','design_id')):
            raise ValueError('download evidence identity mismatch')
        if any(not isinstance(evidence.get(k),str) or not evidence[k].strip()
               for k in ('edit_url','download_entry','completion_evidence')):
            raise ValueError('observed browser download evidence required')
        _check_download_url(evidence,identity,state)
        validate_completion_evidence(state,identity,evidence.get('completion_evidence'),path)
        if evidence.get('pptx_sha256') != sha(path):
            raise ValueError('download evidence hash mismatch')
        r.register_stage2_artifact(temp,'graphics_first_pptx:'+slide_id,path)
        state = r.load_runtime_state(temp)
        state['canva_bridge'].setdefault('downloads',{})[slide_id] = {
            'identity':identity,'path':str(path),'sha256':sha(path),'evidence':evidence}
        try:completion_version=json.loads(evidence['completion_evidence']).get('schema_version',1)
        except (ValueError,AttributeError):completion_version=1
        state['canva_bridge']['downloads'][slide_id]['completion_schema']=completion_version
        parsed=json.loads(evidence['completion_evidence']) if completion_version==2 else {}
        if parsed.get('transport_receipt'):
            state['canva_bridge']['downloads'][slide_id].update(require_transport_receipt=True,acquisition_id=parsed['acquisition_id'])
        r._save_runtime_state(temp,state)
    return {'status':'DOWNLOAD_BOUND','path':str(path)}


def requires_binding(state, slide_id):
    return (state.get('canva_bridge',{}).get('route') == 'codex-canva'
            or 'reconstruction_attempt' in state['slides'][slide_id])


def page_provenance(state, slide_id):
    sealed = state['slides'][slide_id].get('validated_single_page')
    errors = r.validate_current_artifact(sealed,state,slide_id)
    if errors:
        raise ValueError(slide_id+': '+', '.join(errors))
    if not requires_binding(state,slide_id):
        return {'route':'external','slide_id':slide_id,'validated_pptx_sha256':sealed['validated_pptx_sha256']}
    download = download_current(state,slide_id)
    binding = state['canva_bridge'].get('validated',{}).get(slide_id)
    if (not isinstance(binding,dict) or binding.get('identity') != download['identity']
            or binding.get('download_sha256') != download['sha256']
            or binding.get('seal') != sealed):
        raise ValueError('validated page belongs to a superseded or unbound design: '+slide_id)
    single_page(resolve(state,sealed['path']))
    mapping_review = binding.get('canvas_mapping_review')
    if mapping_review is not None:
        review_path = resolve(state,mapping_review['path'])
        if not review_path.is_file() or sha(review_path) != mapping_review['sha256']:
            _canvas_error('sealed mapping review changed or missing')
    return copy.deepcopy(binding)



# Canvas adaptation is local only. Approved canvas and downloaded bytes are immutable.
def _canvas_error(detail):
    raise ValueError('CANVAS_MAPPING_UNRESOLVED: ' + detail)


def _positive_dimension(value):
    from fractions import Fraction
    if isinstance(value, bool):
        _canvas_error('invalid dimension')
    try:
        result = Fraction(str(value))
    except (ValueError, TypeError, ZeroDivisionError):
        _canvas_error('invalid dimension')
    if result <= 0:
        _canvas_error('nonpositive dimension')
    return result


def plan_canvas_mapping(actual_width, actual_height, canvas, registration=None):
    """Return a deterministic plan; same-aspect arithmetic is NOT framing evidence."""
    aw, ah = map(_positive_dimension, (actual_width, actual_height))
    sw, sh = map(_positive_dimension, (canvas.get('source_width'), canvas.get('source_height')))
    tw, th = map(_positive_dimension, (canvas.get('target_slide_width_emu'), canvas.get('target_slide_height_emu')))
    if any(v.denominator != 1 for v in (aw, ah, tw, th)):
        _canvas_error('slide dimensions must be integer EMU')
    from pptx.oxml.simpletypes import ST_SlideSizeCoordinate
    try:
        for dimension in (aw, ah, tw, th):
            ST_SlideSizeCoordinate.validate(int(dimension))
    except ValueError:
        _canvas_error('slide dimension outside supported PowerPoint range')
    delta = (abs(aw * sw / tw - sw), abs(ah * sh / th - sh))
    direct = all(v <= 5 for v in delta)
    aspect = [abs(aw * sh / ah - sw), abs(ah * sw / aw - sh),
              abs(tw * sh / th - sw), abs(th * sw / tw - sh)]
    if not direct:
        if any(v > 5 for v in aspect):
            raise ValueError('DECK_CANVAS_INVARIANT_FAILED: aspect discrepancy exceeds 5 source pixels')
        keys = ('same_framing', 'no_crop', 'no_expansion', 'no_rotation', 'no_relayout')
        if not isinstance(registration, dict):
            _canvas_error('actual framing evidence required')
        checks, evidence = registration.get('checks'), registration.get('evidence')
        if not isinstance(checks, dict) or not isinstance(evidence, dict):
            _canvas_error('invalid framing evidence')
        if any(checks.get(k) != 'PASS' or not isinstance(evidence.get(k), str) or not evidence[k].strip() for k in keys):
            _canvas_error('incomplete framing evidence')
    scale = min(tw / aw, th / ah) if not direct else _positive_dimension(1)
    dx, dy = ((tw - aw * scale) / 2, (th - ah * scale) / 2) if not direct else (aw*0, ah*0)
    rect = [dx, dy, aw * scale, ah * scale] if not direct else [tw*0, th*0, tw, th]
    ratio = lambda f: [f.numerator, f.denominator]
    return {'mode': 'unchanged_objects' if direct else 'uniform_fit',
            'actual_size_emu': [int(aw), int(ah)], 'target_size_emu': [int(tw), int(th)],
            'source_size_pixels': [float(sw), float(sh)],
            'canvas_sha256': hashlib.sha256(r._canonical(canvas)).hexdigest(),
            'scale_ratio': ratio(scale), 'offset_ratio': [ratio(dx), ratio(dy)],
            'content_rect_emu': [float(v) for v in rect],
            'original_delta_source_pixels': [float(v) for v in delta],
            'aspect_delta_source_pixels': [float(v) for v in aspect],
            'registration': copy.deepcopy(registration) if not direct else None}


def _canvas_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape._element.tag.endswith('}grpSp'):
            yield from _canvas_shapes(shape.shapes)


def _canvas_plain_text(shape):
    """Only a plain text-only box may disappear under the existing neutrality gate."""
    from pptx.oxml.ns import qn
    el = shape._element
    nv = el.find(qn('p:nvSpPr'))
    props = el.find(qn('p:spPr'))
    if any(node.tag in {qn('p:extLst'),qn('a:extLst'),qn('p:ph'),qn('a:videoFile'),qn('a:audioFile'),qn('p:oleObj')} for node in el.iter()):
        return False
    return (el.tag == qn('p:sp') and nv is not None and
            nv.find(qn('p:cNvSpPr')) is not None and nv.find(qn('p:cNvSpPr')).get('txBox') == '1' and props is not None and
            props.find(qn('a:noFill')) is not None and el.find(qn('p:style')) is None and
            any((n.text or '').strip() for n in el.xpath('.//a:t')) and
            all(n.tag in {qn('a:xfrm'), qn('a:prstGeom'), qn('a:noFill'), qn('a:ln')} for n in props) and
            all(ln.find(qn('a:noFill')) is not None for ln in props.findall(qn('a:ln'))))


def _canvas_cleanup(prs, source_sha256, cleanup):
    if not isinstance(cleanup, dict) or cleanup.get('download_sha256') != source_sha256:
        _canvas_error('cleanup must bind original download')
    ids = cleanup.get('removed_text_shape_ids')
    if not isinstance(ids, list) or any(type(v) is not int for v in ids) or len(set(ids)) != len(ids):
        _canvas_error('invalid cleanup IDs')
    if ids and (not isinstance(cleanup.get('evidence'), str) or not cleanup['evidence'].strip()):
        _canvas_error('actual neutrality cleanup evidence required')
    shapes = list(_canvas_shapes(prs.slides[0].shapes))
    by_id = {s.shape_id: s for s in shapes}
    if len(by_id) != len(shapes):
        _canvas_error('duplicate shape ID')
    # Validate ALL removals before mutating anything.
    if any(i not in by_id or not _canvas_plain_text(by_id[i]) for i in ids):
        _canvas_error('cleanup cannot remove non-text graphics')
    for i in ids:
        el = by_id[i]._element
        el.getparent().remove(el)


def _canvas_preflight(prs):
    from pptx.oxml.ns import qn
    for shape in _canvas_shapes(prs.slides[0].shapes):
        el = shape._element
        if el.tag not in {qn('p:sp'), qn('p:pic'), qn('p:grpSp')}:
            _canvas_error('unsupported graphic object')
        if any((n.text or '').strip() for n in el.xpath('.//a:t')):
            raise ValueError('GRAPHICS_FIRST_TEXT_CONTAMINATION')
        props = el.find(qn('p:grpSpPr' if el.tag == qn('p:grpSp') else 'p:spPr'))
        xf = props.find(qn('a:xfrm')) if props is not None else None
        if xf is None or xf.find(qn('a:off')) is None or xf.find(qn('a:ext')) is None:
            _canvas_error('explicit top-level transform required')
        if el.tag == qn('p:grpSp') and (xf.find(qn('a:chOff')) is None or xf.find(qn('a:chExt')) is None):
            _canvas_error('explicit group child coordinates required')
        if el.tag == qn('p:grpSp'):
            try:
                child_offset = [int(xf.find(qn('a:chOff')).get(k)) for k in ('x','y')]
                child_extent = [int(xf.find(qn('a:chExt')).get(k)) for k in ('cx','cy')]
            except (TypeError, ValueError):
                _canvas_error('invalid group child coordinates')
            if any(v <= 0 for v in child_extent):
                _canvas_error('nonpositive group child extent')
        if el.find(qn('p:style')) is not None or any(
                node.tag in {qn('p:extLst'),qn('a:extLst'),qn('p:ph'),qn('a:videoFile'),qn('a:audioFile'),qn('p:oleObj')} for node in el.iter()):
            _canvas_error('inherited style, extensions or special objects require separate visual handling')
        allowed_props = {qn('a:xfrm'),qn('a:custGeom'),qn('a:prstGeom'),qn('a:noFill'),
                         qn('a:solidFill'),qn('a:gradFill'),qn('a:blipFill'),qn('a:pattFill'),qn('a:ln'),qn('a:effectLst')}
        if any(node.tag not in allowed_props for node in props):
            _canvas_error('unsupported graphic properties')
        if any(n.tag in {qn('a:effectDag'), qn('a:scene3d'), qn('a:sp3d')} or
               (n.tag == qn('a:effectLst') and len(n)) or
               (n.tag == qn('a:ln') and n.find(qn('a:noFill')) is None) for n in props):
            _canvas_error('stroke or effect cannot be transformed by geometry only')
        for node, attrs in ((xf.find(qn('a:off')), ('x','y')), (xf.find(qn('a:ext')), ('cx','cy'))):
            try:
                values = [int(node.get(k)) for k in attrs]
            except (TypeError, ValueError):
                _canvas_error('invalid shape transform')
            if attrs == ('cx','cy') and any(v < 0 for v in values):
                _canvas_error('negative shape extent')


def _apply_canvas_plan(prs, plan):
    from fractions import Fraction
    from pptx.oxml.ns import qn
    if plan['mode'] == 'uniform_fit':
        _canvas_preflight(prs)
        s = Fraction(*plan['scale_ratio'])
        dx, dy = (Fraction(*v) for v in plan['offset_ratio'])
        for shape in prs.slides[0].shapes:
            el = shape._element
            props = el.find(qn('p:grpSpPr' if el.tag == qn('p:grpSp') else 'p:spPr'))
            xf = props.find(qn('a:xfrm'))
            off, ext = xf.find(qn('a:off')), xf.find(qn('a:ext'))
            off.set('x', str(round(int(off.get('x')) * s + dx)))
            off.set('y', str(round(int(off.get('y')) * s + dy)))
            ext.set('cx', str(round(int(ext.get('cx')) * s)))
            ext.set('cy', str(round(int(ext.get('cy')) * s)))
    prs.slide_width, prs.slide_height = plan['target_size_emu']


def prepare_canvas_mapping(download_pptx, canvas, registration=None, cleanup=None):
    """Return (in-memory restored base, report). Add native text and save ONCE at a new path."""
    from pptx import Presentation
    single_page(download_pptx)
    digest = sha(download_pptx)
    prs = Presentation(str(download_pptx))
    plan = plan_canvas_mapping(prs.slide_width, prs.slide_height, canvas, registration)
    if plan['mode'] == 'uniform_fit':
        if registration.get('download_sha256') != digest or not registration.get('text_clean_fingerprint'):
            _canvas_error('registration must bind download and current text-clean source')
    cleanup = copy.deepcopy(cleanup) if cleanup is not None else {
        'download_sha256':digest, 'removed_text_shape_ids':[], 'evidence':''}
    _canvas_cleanup(prs, digest, cleanup)
    report = {'download_sha256':digest, 'plan':plan, 'cleanup':cleanup,
              'retained_shape_ids':[s.shape_id for s in prs.slides[0].shapes]}
    _apply_canvas_plan(prs, plan)
    return prs, report


def map_native_text_geometry(element, plan):
    """Map both visible bbox and glyph-height target; do not modify Manifest or font-fit rules."""
    import math
    from fractions import Fraction
    if plan['mode'] == 'uniform_fit':
        s = Fraction(*plan['scale_ratio'])
        dx, dy = (Fraction(*v) for v in plan['offset_ratio'])
        w, h = (Fraction(v)*s for v in plan['actual_size_emu'])
    else:
        dx = dy = Fraction(0)
        w, h = map(Fraction, plan['target_size_emu'])
    bbox = element.get('visual_bbox_normalized')
    if not isinstance(bbox, (list, tuple)) or len(bbox) != 4 or any(
            isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v) for v in bbox):
        _canvas_error('invalid normalized text geometry')
    x,y,bw,bh = (Fraction(str(v)) for v in bbox)
    if bw < 0 or bh < 0:
        _canvas_error('negative text extent')
    result = {'bbox_emu':[round(dx+x*w),round(dy+y*h),round(bw*w),round(bh*h)]}
    if 'font_size_normalized' in element:
        value = _positive_dimension(element['font_size_normalized']) * h
        result.update(visible_glyph_height_emu=float(value), visible_glyph_height_pt=float(value/12700))
    return result


def _canvas_resource_signature(pptx):
    with zipfile.ZipFile(pptx) as z:
        media = {n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('ppt/media/')}
        rel_name = 'ppt/slides/_rels/slide1.xml.rels'
        rels = sorted(tuple(sorted(el.attrib.items())) for el in ET.fromstring(z.read(rel_name))) if rel_name in z.namelist() else []
    return {'media':media, 'relationships':rels}


def verify_canvas_mapping(download_pptx, restored_pptx, canvas, report, expected_clean_fingerprint=None):
    from pptx import Presentation
    from lxml import etree
    if Path(download_pptx).resolve() == Path(restored_pptx).resolve():
        _canvas_error('download cannot be overwritten')
    if not isinstance(report, dict) or report.get('download_sha256') != sha(download_pptx):
        _canvas_error('mapping report download mismatch')
    claimed = report.get('plan', {})
    registration = claimed.get('registration')
    if claimed.get('mode') == 'uniform_fit' and expected_clean_fingerprint is not None:
        if not isinstance(registration, dict) or registration.get('text_clean_fingerprint') != expected_clean_fingerprint:
            _canvas_error('mapping text-clean source changed')
    expected, recomputed = prepare_canvas_mapping(download_pptx, canvas, registration, report.get('cleanup'))
    if any(report.get(k) != recomputed[k] for k in ('plan','cleanup','retained_shape_ids')):
        _canvas_error('mapping parameters or baseline changed')
    single_page(restored_pptx)
    actual = Presentation(str(restored_pptx))
    if [actual.slide_width,actual.slide_height] != recomputed['plan']['target_size_emu']:
        raise ValueError('DECK_CANVAS_INVARIANT_FAILED: restored canvas is not approved target')
    before, after = list(expected.slides[0].shapes), list(actual.slides[0].shapes)
    # Existing retained graphics stay in order; only native text boxes may be appended.
    if len(after) < len(before) or [s.shape_id for s in after[:len(before)]] != [s.shape_id for s in before]:
        _canvas_error('retained object loss or reordering')
    for src, dst in zip(before, after):
        if etree.tostring(src._element,method='c14n') != etree.tostring(dst._element,method='c14n'):
            _canvas_error('retained object differs from deterministic transform')
    if any(not _canvas_plain_text(s) for s in after[len(before):]):
        _canvas_error('only native text may follow retained graphics')
    if _canvas_resource_signature(download_pptx) != _canvas_resource_signature(restored_pptx):
        _canvas_error('download media or slide relationships changed')
    if report.get('restored_pptx_sha256') != sha(restored_pptx):
        _canvas_error('mapping report restored hash mismatch')
    return recomputed


def seal_page(state_path, slide_id, pptx, gates_path):
    review = json.loads(Path(gates_path).read_text(encoding='utf-8'))
    gates = review.get('checks',{})
    if any(gates.get(k) != 'PASS' for k in r.GATE_NAMES):
        raise ValueError('four Hard Gates must all be PASS')
    page = Path(pptx).resolve()
    single_page(page)
    with transaction(state_path) as temp:
        r.require_stage2_entry(temp)
        state = r.load_runtime_state(temp)
        download = download_current(state,slide_id)
        if page == resolve(state,download['path']):
            raise ValueError('restored output must preserve the graphics-first source at a separate path')
        if (any(review.get(k) != download['identity'][k] for k in ('attempt_id','slide_id','design_id'))
                or review.get('download_sha256') != download['sha256']
                or review.get('restored_pptx_sha256') != sha(page)):
            raise ValueError('single-page review identity or file hash mismatch')
        if any(not isinstance(review.get('evidence',{}).get(k),str) or not review['evidence'][k].strip() for k in r.GATE_NAMES):
            raise ValueError('actual four-gate evidence required')
        for kind in ('finalized_manifest','font_fallback'):
            if not r.stage2_artifact_current(temp,kind+':'+slide_id):
                raise ValueError('current registered '+kind+' required')
            if not isinstance(json.loads(resolve(state,state['slides'][slide_id][kind]).read_text(encoding='utf-8')),dict):
                raise ValueError(kind+' must be JSON mapping')
        if review.get('text_restore_fingerprint') != r.expected_lineage_from_state(state,slide_id)['text_restore_fingerprint']:
            raise ValueError('single-page review restoration inputs changed')
        slide = state['slides'][slide_id]
        mapping = review.get('canvas_mapping')
        canvas = json.loads(resolve(state,slide['canvas']).read_text(encoding='utf-8'))
        # Old sealed pages are not migrated. New explicit-EMU canvas pages cannot bypass adaptation proof.
        mapping_required = False
        if all(k in canvas for k in ('source_width','source_height','target_slide_width_emu','target_slide_height_emu')):
            from pptx import Presentation
            source_prs = Presentation(str(resolve(state,download['path'])))
            aw, ah = source_prs.slide_width, source_prs.slide_height
            sw, sh = _positive_dimension(canvas['source_width']), _positive_dimension(canvas['source_height'])
            tw, th = _positive_dimension(canvas['target_slide_width_emu']), _positive_dimension(canvas['target_slide_height_emu'])
            if tw.denominator != 1 or th.denominator != 1:
                _canvas_error('target dimensions must be integer EMU')
            mapping_required = abs(aw*sw/tw-sw)>5 or abs(ah*sh/th-sh)>5
            output_prs = Presentation(str(page))
            if [output_prs.slide_width,output_prs.slide_height] != [int(tw),int(th)]:
                raise ValueError('DECK_CANVAS_INVARIANT_FAILED: restored canvas is not approved target')
        if mapping_required and (not isinstance(mapping,dict) or mapping.get('plan',{}).get('mode')!='uniform_fit'):
            _canvas_error('uniform-fit review required for export size discrepancy')
        if mapping is not None:
            verify_canvas_mapping(resolve(state,download['path']),page,canvas,mapping,download['identity']['text_clean_fingerprint'])
        old = slide.pop('validated_single_page',None)
        if old:
            slide.setdefault('superseded_validated_single_pages',[]).append(old)
        state.pop('merged_deck',None);state.pop('validated_deck',None)
        state['canva_bridge'].pop('merged',None)
        r._save_runtime_state(temp,state)
        seal = r.seal_validated_single_page(temp,slide_id,page,gates)
        state = r.load_runtime_state(temp)
        state['canva_bridge'].setdefault('validated',{})[slide_id] = {
            'identity':download['identity'],'download_sha256':download['sha256'],'seal':seal}
        if mapping is not None:
            state['canva_bridge']['validated'][slide_id]['canvas_mapping_review'] = {
                'path':str(Path(gates_path).resolve()),'sha256':sha(gates_path)}
        r._save_runtime_state(temp,state)
    return {'status':'PAGE_SEALED','slide_id':slide_id,'path':str(page)}


def deck_provenance(state, order):
    return [page_provenance(state,sid) for sid in order]


def status(state_path):
    if Path(state_path).resolve().with_name(Path(state_path).name+'.bridge-lock').exists():
        return {'status':'BLOCKED','detail':'operation running or interrupted; inspect lock, do not auto-unlock/retry'}
    state = r.load_runtime_state(state_path)
    try:
        r.require_stage2_entry(state_path)
    except (KeyError,ValueError,OSError):
        return {'status':'AWAITING_STAGE1_APPROVAL'}
    if not state.get('stage2_run'):
        return {'status':'STAGE2_EXPLORATION'}
    visual = r.stage2_visual_status(state_path)
    if visual['status'] != 'COMPLETE':
        result = {'status':'STAGE2_PRODUCTION','missing':visual.get('missing',[])}
        if state['stage2_run'].get('title_policy') is not None:
            result['title_policy'] = r.titles.context_from_state(state)
        return result
    try:
        r.require_stage2_visual_approval(state_path)
    except (KeyError,ValueError,OSError,TypeError):
        return {'status':'AWAITING_USER_VISUAL_APPROVAL'}
    if r.stage2_handoff_status(state_path)['status'] != 'COMPLETE':
        return {'status':'STAGE2_HANDOFF_INCOMPLETE'}
    run_order = state['stage2_run']['page_order']
    order = state.get('deck_order',run_order)
    if not run_order or not isinstance(order,list) or not r.validate_deck_order(run_order,order) or not r.validate_deck_order(list(state['slides']),order):
        return {'status':'BLOCKED','detail':'trusted page order invalid'}
    for sid in order:
        try:
            page_provenance(state,sid)
            continue
        except (OSError,ValueError,KeyError,TypeError) as error:
            detail = str(error)
        if requires_binding(state,sid):
            try:
                _check_text_preparation(state,sid)
            except (OSError,ValueError,KeyError,TypeError) as error:
                pending = state.get('canva_bridge',{}).get('active',{}).get(sid,{})
                if pending.get('status') == 'PENDING':
                    return {'status':'BLOCKED','slide_id':sid,'detail':'completion unknown; do not repeat call; '+str(error)}
                return {'status':'PREPARE_TEXT','slide_id':sid,'detail':str(error)}
        if not r.stage2_artifact_current(state_path,'text_clean:'+sid):
            return {'status':'PREPARE_TEXT','slide_id':sid,'detail':detail}
        if requires_binding(state,sid):
            if not state.get('canva_bridge',{}).get('active',{}).get(sid) and 'reconstruction_attempt' not in state['slides'][sid]:
                return {'status':'WAIT_RECONSTRUCTION','slide_id':sid,'detail':'no attempt started; begin only for authorized formal reconstruction'}
            try:
                active = _active(state,sid)
            except (OSError,ValueError,KeyError,TypeError) as error:
                return {'status':'BLOCKED','slide_id':sid,'detail':str(error)}
            if active['status'] == 'PENDING':
                return {'status':'WAIT_RECONSTRUCTION','slide_id':sid,'detail':'completion unknown; do not repeat call'}
            if active['status'] == 'TOOL_FAILED':
                return {'status':'TOOL_FAILED','slide_id':sid,'failure_code':active['failure_code'],
                        'detail':'terminal failure recorded; actual environment repair and explicit recovery required'}
            if active['status'] in ('RETRY_ONCE','ASSISTED_CLEANUP_REQUIRED','RETURN_TO_STAGE_2'):
                return {'status':active['status'],'slide_id':sid,'detail':'follow canonical structural recovery; no automatic quota use'}
            try:
                active_identity(state,sid)
            except (OSError,ValueError,KeyError,TypeError) as error:
                return {'status':'BLOCKED','slide_id':sid,'detail':str(error)}
            download = state['canva_bridge'].get('downloads',{}).get(sid)
            if download is None or download.get('identity') != active_identity(state,sid):
                return {'status':'WAIT_DOWNLOAD','slide_id':sid,'detail':'verify download completion before any retry'}
            try:
                download_current(state,sid)
            except (OSError,ValueError,KeyError,TypeError) as error:
                return {'status':'BLOCKED','slide_id':sid,'detail':str(error)}
        return {'status':'RESTORE_TEXT','slide_id':sid,'detail':detail}
    provenance = deck_provenance(state,order)
    merged = state.get('canva_bridge',{}).get('merged')
    if r.validate_current_merged_deck(state) or not isinstance(merged,dict) or merged.get('pages') != provenance:
        return {'status':'PREPARE_DECK','deck_order':order}
    if merged.get('merged_pptx_sha256') != state['merged_deck']['merged_pptx_sha256']:
        return {'status':'BLOCKED','detail':'merged provenance hash mismatch'}
    validated = state.get('validated_deck',{})
    if validated.get('status') == 'PASS' and validated.get('merged_pptx_sha256') == merged['merged_pptx_sha256']:
        return {'status':'COMPLETE','path':state['merged_deck']['merged_pptx_path']}
    return {'status':'DECK_VALIDATION','path':state['merged_deck']['merged_pptx_path']}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('action',choices=('init','status','start-run','register-text-plan','register-text-clean','begin-attempt','accept-attempt','register-design-url','bind-download','seal-page','register-upload','register-design-observation','register-upload-result'))
    parser.add_argument('--state',required=True)
    parser.add_argument('--route',choices=('codex-canva','external'),default='codex-canva')
    parser.add_argument('--slide-id');parser.add_argument('--attempt-id');parser.add_argument('--artifact');parser.add_argument('--manifest');parser.add_argument('--inventory')
    parser.add_argument('--lookup');parser.add_argument('--result');parser.add_argument('--pptx');parser.add_argument('--evidence');parser.add_argument('--gates')
    parser.add_argument('--rebuild',action='store_true')
    parser.add_argument('--recovery-evidence')
    parser.add_argument('--title-choice-required',action='store_true')
    parser.add_argument('--title-observation')
    args = parser.parse_args()
    try:
        if args.action == 'init': result = init(args.state,args.route)
        elif args.action == 'status': result = status(args.state)
        elif args.action == 'start-run': result = {**start_run(args.state,args.title_choice_required), 'status':'RUN_READY'}
        elif args.action == 'register-text-plan': result = register_text_plan(args.state,args.slide_id,args.manifest,args.inventory)
        elif args.action == 'register-text-clean': result = register_text_clean(args.state,args.slide_id,args.artifact)
        elif args.action == 'begin-attempt': result = begin_attempt(args.state,args.slide_id,args.rebuild,args.recovery_evidence)
        elif args.action == 'accept-attempt': result = accept_attempt(args.state,args.slide_id,args.attempt_id,args.result)
        elif args.action == 'register-design-url': result = register_design_url(args.state,args.slide_id,args.lookup)
        elif args.action == 'register-upload-result': result = register_upload_result(args.state,args.slide_id,args.result)
        elif args.action == 'register-upload': result = register_upload(args.state,args.slide_id,args.artifact)
        elif args.action == 'register-design-observation': result = register_design_observation(args.state,args.slide_id,args.evidence,args.title_observation)
        elif args.action == 'bind-download': result = bind_download(args.state,args.slide_id,args.pptx,args.evidence)
        else: result = seal_page(args.state,args.slide_id,args.pptx,args.gates)
    except (OSError,ValueError,KeyError,TypeError,zipfile.BadZipFile,ET.ParseError) as error:
        result = {'status':'BLOCKED','detail':str(error)}
    print(json.dumps(result,ensure_ascii=False,indent=2))
    return 1 if result['status']=='BLOCKED' else 0


def register_transport_requirement(state_path,identity,acquisition_id,receipt_ref=None):
    """Persist new-policy requirement before acquisition; old bound tasks stay untouched."""
    import download_evidence as de
    with transaction(state_path) as temp:
        state=r.load_runtime_state(temp)
        if active_identity(state,identity['slide_id'])!=identity:raise ValueError('transport identity changed')
        registry=state['canva_bridge'].setdefault('transport_requirements',{})
        existing=registry.get(identity['slide_id'])
        if existing and existing['identity']==identity and existing['acquisition_id']!=acquisition_id and not existing.get('terminated'):
            raise ValueError('current transport acquisition differs; resume original')
        if existing and existing.get('terminated') and existing['acquisition_id']!=acquisition_id:
            state['canva_bridge'].setdefault('observation_history',[]).append({'kind':'transport_requirements','record':copy.deepcopy(existing)})
            existing=None
        if existing and existing.get('terminated'):raise ValueError('transport acquisition explicitly terminated')
        requirement=existing if existing and existing['identity']==identity else {
            'identity':identity,'acquisition_id':acquisition_id,'receipts':[]}
        if receipt_ref:
            receipt=de.sealed(receipt_ref);de.check_identity(receipt['identity'],identity)
            if receipt['acquisition_id']!=acquisition_id:raise ValueError('receipt acquisition mismatch')
            if receipt_ref not in requirement['receipts']:requirement['receipts'].append(receipt_ref)
        registry[identity['slide_id']]=requirement;r._save_runtime_state(temp,state)

if __name__ == '__main__':
    raise SystemExit(main())
