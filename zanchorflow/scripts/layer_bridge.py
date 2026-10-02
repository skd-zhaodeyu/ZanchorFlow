"""Image Layer backend runtime state machine. Remote work is explicit and lineage-bound."""
import argparse
import hashlib
import json
import os
from pathlib import Path

from PIL import Image

import runtime as r
import reconstruction_router as router
import layer_policy
import layer_review_boxes
import layer_provider_360
import layer_validate_bundle
from pptx import Presentation


def _sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''): h.update(chunk)
    return h.hexdigest()


def _resolve(state,value):
    p=Path(value)
    return p if p.is_absolute() else Path(state['_base_dir'])/p


def _ensure_bridge(state):
    return state.setdefault('layer_bridge',{
        'schema_version':1,'plans':{},'requests':{},'results':{},'bundles':{},
        'visual_qa':{},'graphics_first':{},'validated':{},'call_ledger':[]})


def init(state_path):
    if router.selected_backend(state_path) != router.IMAGE_LAYER:
        raise ValueError('Image Layer backend not selected')
    with router.transaction(state_path) as temp:
        state=r.load_runtime_state(temp); _ensure_bridge(state); r._save_runtime_state(temp,state)
    return {'status':'INITIALIZED'}


def _current_run(state):
    run=state.get('stage2_run')
    if not isinstance(run,dict) or not run.get('run_id'): raise ValueError('STAGE2_RUN_REQUIRED')
    return run


def _current_text_clean(state_path,state,slide_id):
    run=_current_run(state)
    if slide_id not in run.get('page_order',[]): raise ValueError('STAGE2_PAGE_NOT_CURRENT')
    if not r.stage2_artifact_current(state_path,f'text_clean:{slide_id}'):
        raise ValueError('TEXT_CLEAN_REQUIRED')
    path=_resolve(state,state['slides'][slide_id]['text_clean'])
    return path,_sha(path),run['run_id']


def _plan_current(state,slide_id):
    bridge=state.get('layer_bridge',{}); rec=bridge.get('plans',{}).get(slide_id)
    if not isinstance(rec,dict): return None
    run=state.get('stage2_run',{})
    if rec.get('run_id')!=run.get('run_id'): return None
    slide=state.get('slides',{}).get(slide_id,{})
    value=slide.get('text_clean')
    if not value: return None
    path=_resolve(state,value)
    if not path.is_file() or rec.get('text_clean_sha256')!=_sha(path): return None
    plan=_resolve(state,rec.get('plan_path',''))
    if not plan.is_file() or rec.get('plan_sha256')!=_sha(plan): return None
    return rec


def _clear_downstream(bridge,slide_id):
    for key in ('readiness','requests','results','bundles','visual_qa','graphics_first','validated'):
        bridge.setdefault(key,{}).pop(slide_id,None)


def register_plan(state_path,slide_id,plan_path,overlay_path):
    if router.selected_backend(state_path)!=router.IMAGE_LAYER:
        raise ValueError('Image Layer backend not selected')
    state=r.load_runtime_state(state_path); clean,clean_sha,run_id=_current_text_clean(state_path,state,slide_id)
    plan_path=Path(plan_path).resolve(); overlay_path=Path(overlay_path).resolve()
    if not plan_path.is_file() or not overlay_path.is_file(): raise ValueError('LAYER_PLAN_FAILED: plan/overlay missing')
    plan=json.loads(plan_path.read_text(encoding='utf-8'))
    if plan.get('page_id')!=slide_id: raise ValueError('LAYER_PLAN_FAILED: page_id mismatch')
    with Image.open(clean) as image: source_canvas=list(image.size)
    layer_review_boxes.validate_plan(plan,*source_canvas)
    previous=_plan_current(state,slide_id)
    if previous is not None and previous['plan_sha256']==_sha(plan_path):
        return previous # Re-registering the same input/plan cannot erase a consumed retry or current result.
    record={'status':'LAYER_PLAN_READY','slide_id':slide_id,'run_id':run_id,'text_clean_sha256':clean_sha,
            'source_canvas':source_canvas,'plan_path':str(plan_path),'plan_sha256':_sha(plan_path),
            'overlay_path':str(overlay_path),'overlay_sha256':_sha(overlay_path),'approved':False}
    with router.transaction(state_path) as temp:
        current=r.load_runtime_state(temp); bridge=_ensure_bridge(current); bridge['plans'][slide_id]=record; _clear_downstream(bridge,slide_id); r._save_runtime_state(temp,current)
    return record


def approve_plan(state_path,slide_id):
    state=r.load_runtime_state(state_path); rec=_plan_current(state,slide_id)
    if rec is None: raise ValueError('LAYER_PLAN_FAILED: current plan required')
    with router.transaction(state_path) as temp:
        current=r.load_runtime_state(temp); bridge=_ensure_bridge(current); current_rec=bridge['plans'][slide_id]
        current_rec['approved']=True; current_rec['status']='LAYER_PLAN_APPROVED'; r._save_runtime_state(temp,current)
    return {'status':'LAYER_PLAN_APPROVED','slide_id':slide_id}


def register_readiness(state_path,slide_id,evidence_path):
    state=r.load_runtime_state(state_path); plan=_plan_current(state,slide_id)
    if plan is None: raise ValueError('LAYER_PLAN_FAILED: current plan required')
    path=Path(evidence_path).resolve(); report=json.loads(path.read_text(encoding='utf-8'))
    if (report.get('page_id')!=slide_id or report.get('text_clean_sha256')!=plan['text_clean_sha256']
        or report.get('plan_sha256')!=plan['plan_sha256']): raise ValueError('LAYER_PLAN_READINESS: lineage mismatch')
    checks=report.get('checks',{}); reasons=report.get('reasons',{})
    if any(checks.get(k) not in ('READY','NOT_READY') or not layer_policy.text(reasons.get(k)) for k in layer_policy.READINESS):
        raise ValueError('LAYER_PLAN_READINESS: checks/reasons required')
    payload=json.loads(Path(plan['plan_path']).read_text(encoding='utf-8'))
    layer_review_boxes.validate_plan(payload,*plan['source_canvas'])
    rec={'status':'NOT_READY' if 'NOT_READY' in checks.values() else 'READY',
         'text_clean_sha256':plan['text_clean_sha256'],'plan_sha256':plan['plan_sha256'],
         'path':str(path),'sha256':_sha(path),'checks':checks,'reasons':reasons}
    with router.transaction(state_path) as temp:
        cur=r.load_runtime_state(temp); _ensure_bridge(cur).setdefault('readiness',{})[slide_id]=rec; r._save_runtime_state(temp,cur)
    return rec

def readiness_current(state,slide_id):
    plan=_plan_current(state,slide_id); rec=state.get('layer_bridge',{}).get('readiness',{}).get(slide_id)
    if plan is None or not isinstance(rec,dict): return None
    if any(rec.get(k)!=plan.get(k) for k in ('text_clean_sha256','plan_sha256')): return None
    path=_resolve(state,rec.get('path',''))
    return rec if path.is_file() and _sha(path)==rec.get('sha256') else None

def _provider_identity(provider):
    return {'provider':getattr(provider,'PROVIDER_ID',getattr(provider,'__name__',provider.__class__.__name__)),
        'model':getattr(provider,'MODEL','reveal_layer'),'version':getattr(provider,'VERSION','v2.3.4'),
        **getattr(provider,'MODEL_PARAMETERS',{'pipeline_type':'crop','seed':42,'steps':10,'time_out':3600})}

def _quality_key(plan,identity):
    return hashlib.sha256(json.dumps([plan['text_clean_sha256'],plan['plan_sha256'],identity['provider'],identity['model']],sort_keys=True).encode()).hexdigest()

def authorize_quality_retry(state_path,slide_id,evidence_path,authorization_path):
    state=r.load_runtime_state(state_path); plan=_plan_current(state,slide_id); result=result_current(state,slide_id)
    req=state.get('layer_bridge',{}).get('requests',{}).get(slide_id)
    if plan is None or result is None or not isinstance(req,dict) or req.get('status')!='DONE':
        raise ValueError('SEVERE_QUALITY_RETRY: normally completed result required')
    identity=req.get('provider_identity')
    if not isinstance(identity,dict): raise ValueError('SEVERE_QUALITY_RETRY: original provider parameters required')
    key=_quality_key(plan,identity)
    if state.get('layer_bridge',{}).get('quality_retries',{}).get(key,{}).get('quality_retry_count',0):
        raise ValueError('SEVERE_QUALITY_RETRY_EXHAUSTED')
    ep=Path(evidence_path).resolve(); ap=Path(authorization_path).resolve()
    proof=json.loads(ep.read_text(encoding='utf-8')); auth=json.loads(ap.read_text(encoding='utf-8'))
    binding={'text_clean_sha256':plan['text_clean_sha256'],'plan_sha256':plan['plan_sha256'],
        'result_sha256':result['sha256'],'task_id':result['task_id']}
    if any(proof.get(k)!=v for k,v in binding.items()) or proof.get('page_id')!=slide_id:
        raise ValueError('SEVERE_QUALITY_RETRY: evidence identity')
    if proof.get('classification')!='SEVERE' or proof.get('category') not in layer_policy.SEVERE:
        raise ValueError('SEVERE_QUALITY_RETRY: ordinary limitations ineligible')
    if any(proof.get('excluded_errors',{}).get(k)!='PASS' for k in layer_policy.EXCLUSIONS):
        raise ValueError('SEVERE_QUALITY_RETRY: exclude technical/geometry/assembly errors')
    layer_policy.evidence_assets(proof.get('evidence'),ep.parent)
    if (auth.get('authorized') is not True or type(auth.get('remaining_image_calls')) is not int or auth['remaining_image_calls']<1
        or not layer_policy.text(auth.get('authorization_evidence')) or auth.get('revoked',False)
        or any(auth.get(k)!=v for k,v in {**{k:binding[k] for k in ('text_clean_sha256','plan_sha256')},
            'provider':identity['provider'],'model':identity['model']}.items())):
        raise ValueError('SEVERE_QUALITY_RETRY: current paid authorization required')
    rec={'retry_reason':'severe_quality_failure','quality_retry_count':0,'key':key,**binding,
         'provider_identity':identity,'original_result_id':result['task_id'],
         'severe_quality_evidence_file':str(ep),'severe_quality_evidence_sha':_sha(ep),
         'authorization_file':str(ap),'authorization_sha256':_sha(ap),'authorized':True}
    with router.transaction(state_path) as temp:
        cur=r.load_runtime_state(temp); b=_ensure_bridge(cur)
        current=result_current(cur,slide_id)
        if current is None or current['sha256']!=result['sha256'] or b.setdefault('quality_retries',{}).get(key,{}).get('quality_retry_count',0):
            raise ValueError('SEVERE_QUALITY_RETRY: state changed')
        b['quality_retries'][key]=rec; b['requests'][slide_id]['quality_retry_key']=key
        b['call_ledger'].append({'event':'quality_retry_authorized','slide_id':slide_id,**rec})
        r._save_runtime_state(temp,cur)
    return {'status':'SEVERE_QUALITY_RETRY_AUTHORIZED','slide_id':slide_id,'quality_retry_count':0}

def _request_fp(plan_rec,provider_module):
    provider=getattr(provider_module,'MODEL',getattr(provider_module,'__name__',provider_module.__class__.__name__))
    raw=json.dumps({'plan_sha256':plan_rec['plan_sha256'],'text_clean_sha256':plan_rec['text_clean_sha256'],'provider':provider},sort_keys=True).encode()
    return hashlib.sha256(raw).hexdigest()


def authorize_retry(state_path,slide_id,*,accept_duplicate_charge_risk=False):
    state=r.load_runtime_state(state_path); plan=_plan_current(state,slide_id)
    if plan is None: raise ValueError('LAYER_PLAN_FAILED: current plan required')
    request=state.get('layer_bridge',{}).get('requests',{}).get(slide_id)
    if not isinstance(request,dict) or request.get('status') not in ('LAYER_REQUEST_UNCERTAIN','LAYER_RETRY_APPROVAL_REQUIRED'):
        raise ValueError('retry authorization is not applicable')
    if request['status']=='LAYER_REQUEST_UNCERTAIN' and not accept_duplicate_charge_risk:
        raise ValueError('duplicate charge risk must be explicitly accepted')
    if request.get('retry_reason')=='severe_quality_failure': raise ValueError('SEVERE_QUALITY_RETRY_EXHAUSTED: no technical resubmit of quality attempt')
    auth={'retry_reason':'technical_failure','request_fingerprint':request['request_fingerprint'],'accept_duplicate_charge_risk':bool(accept_duplicate_charge_risk)}
    with router.transaction(state_path) as temp:
        current=r.load_runtime_state(temp); bridge=_ensure_bridge(current); bridge['requests'][slide_id]['retry_authorization']=auth; r._save_runtime_state(temp,current)
    return {'status':'RETRY_AUTHORIZED','slide_id':slide_id}


def submit(state_path,slide_id,work_dir,provider_module=layer_provider_360):
    if router.selected_backend(state_path)!=router.IMAGE_LAYER: raise ValueError('Image Layer backend not selected')
    state=r.load_runtime_state(state_path); plan_rec=_plan_current(state,slide_id)
    if plan_rec is None: raise ValueError('LAYER_PLAN_FAILED: current plan required')
    if not plan_rec.get('approved'): raise ValueError('LAYER_PLAN_NOT_APPROVED')
    api_key=os.environ.get(router.IMAGE_LAYER_API_KEY_ENV)
    if not api_key: return router.image_layer_onboarding()
    readiness=readiness_current(state,slide_id)
    if readiness is None or readiness['status']!='READY': raise ValueError('LAYER_PLAN_READINESS_REQUIRED')
    from layer_geometry_verify import dependencies
    dependencies()
    identity=_provider_identity(provider_module)
    quality=None
    bridge=state.get('layer_bridge',{}); old=bridge.get('requests',{}).get(slide_id); fp=_request_fp(plan_rec,provider_module)
    if isinstance(old,dict):
        if old.get('status') in ('SUBMITTING','LAYER_REQUEST_PENDING'): raise ValueError('layer request already pending')
        if old.get('status') in ('LAYER_REQUEST_UNCERTAIN','LAYER_RETRY_APPROVAL_REQUIRED'):
            auth=old.get('retry_authorization')
            if not isinstance(auth,dict) or auth.get('request_fingerprint')!=old.get('request_fingerprint'):
                raise ValueError(old.get('status') if old.get('status')=='LAYER_REQUEST_UNCERTAIN' else 'retry authorization required')
        elif old.get('status')=='DONE':
            key=old.get('quality_retry_key'); quality=bridge.get('quality_retries',{}).get(key)
            if not isinstance(quality,dict) or not quality.get('authorized'):
                return {'status':'LAYER_RESULT_BOUND','slide_id':slide_id}
            if quality.get('quality_retry_count',0): raise ValueError('SEVERE_QUALITY_RETRY_EXHAUSTED')
            result=result_current(state,slide_id)
            if result is None or result['sha256']!=quality['result_sha256'] or result['task_id']!=quality['original_result_id']:
                raise ValueError('SEVERE_QUALITY_RETRY: original result changed')
            proof=json.loads(Path(quality['severe_quality_evidence_file']).read_text(encoding='utf-8'))
            layer_policy.evidence_assets(proof.get('evidence'),Path(quality['severe_quality_evidence_file']).parent)
            if identity!=quality['provider_identity'] or key!=_quality_key(plan_rec,identity):
                raise ValueError('SEVERE_QUALITY_RETRY: input/plan/provider parameters changed')
            for file_key,hash_key in [('authorization_file','authorization_sha256'),('severe_quality_evidence_file','severe_quality_evidence_sha')]:
                if _sha(quality[file_key])!=quality[hash_key]: raise ValueError('SEVERE_QUALITY_RETRY: authorization/evidence changed')
    router.require_image_target(state_path,state)
    attempt=(old or {}).get('attempt_no',0)+1
    reservation={'status':'SUBMITTING','slide_id':slide_id,'run_id':plan_rec['run_id'],'plan_sha256':plan_rec['plan_sha256'],
                 'text_clean_sha256':plan_rec['text_clean_sha256'],'request_fingerprint':fp,'attempt_no':attempt,'provider_identity':identity,
                 'retry_reason':'severe_quality_failure' if quality else ('technical_failure' if old else 'initial_submit')}
    if quality: reservation['quality_retry_key']=quality['key']
    with router.transaction(state_path) as temp:
        current=r.load_runtime_state(temp); b=_ensure_bridge(current)
        router.require_image_target(temp,current)
        if b.get('requests',{}).get(slide_id)!=old: raise ValueError('layer request changed before reservation')
        current_plan=_plan_current(current,slide_id); current_ready=readiness_current(current,slide_id)
        if current_plan!=plan_rec or not current_ready or current_ready['status']!='READY':
            raise ValueError('layer plan/readiness changed before reservation')
        if quality:
            q=b['quality_retries'][quality['key']]
            if q.get('quality_retry_count',0) or not q.get('authorized'): raise ValueError('SEVERE_QUALITY_RETRY_EXHAUSTED')
            q['quality_retry_count']=1; q['authorized']=False
            b.setdefault('result_history',{}).setdefault(slide_id,[]).append(b['results'][slide_id])
            b.setdefault('request_history',{}).setdefault(slide_id,[]).append(old)
            for name in ('results','bundles','visual_qa','graphics_first','validated'): b.setdefault(name,{}).pop(slide_id,None)
            current.pop('merged_deck',None); current.pop('validated_deck',None); b.pop('merged',None)
        b['requests'][slide_id]=reservation
        b['call_ledger'].append({'slide_id':slide_id,'event':'submit_reserved',**reservation,
            **({'quality_retry_count':1,'original_result_id':quality['original_result_id'],
                'severe_quality_evidence_sha':quality['severe_quality_evidence_sha']} if quality else {})})
        r._save_runtime_state(temp,current)
    current=r.load_runtime_state(state_path); clean=_resolve(current,current['slides'][slide_id]['text_clean'])
    plan=json.loads(_resolve(current,plan_rec['plan_path']).read_text(encoding='utf-8')); boxes=[x['bbox'] for x in plan['targets']]
    try:
        response=provider_module.submit_task(clean,boxes,api_key)
    except Exception as exc:
        explicit_failure=isinstance(exc,ValueError) and str(exc).startswith('LAYER_REQUEST_FAILED:')
        next_status='LAYER_RETRY_APPROVAL_REQUIRED' if explicit_failure else 'LAYER_REQUEST_UNCERTAIN'
        event='submit_terminal_failed' if explicit_failure else 'submit_uncertain'
        with router.transaction(state_path) as temp:
            current=r.load_runtime_state(temp); req=_ensure_bridge(current)['requests'][slide_id]
            if req.get('request_fingerprint')==fp and req.get('attempt_no')==attempt:
                req['status']=next_status; req['error']=type(exc).__name__
                current['layer_bridge']['call_ledger'].append({'slide_id':slide_id,'event':event,'attempt_no':attempt})
                r._save_runtime_state(temp,current)
        return {'status':next_status,'slide_id':slide_id}
    task_id=response.get('task_id')
    if not isinstance(task_id,str) or not task_id:
        with router.transaction(state_path) as temp:
            cur=r.load_runtime_state(temp); _ensure_bridge(cur)['requests'][slide_id]['status']='LAYER_REQUEST_UNCERTAIN'; r._save_runtime_state(temp,cur)
        return {'status':'LAYER_REQUEST_UNCERTAIN','slide_id':slide_id}
    with router.transaction(state_path) as temp:
        current=r.load_runtime_state(temp); b=_ensure_bridge(current); req=b['requests'][slide_id]
        if req.get('request_fingerprint')!=fp or req.get('attempt_no')!=attempt: raise ValueError('layer request lineage changed')
        req.update(status='LAYER_REQUEST_PENDING',task_id=task_id); req.pop('retry_authorization',None)
        b['call_ledger'].append({'slide_id':slide_id,'event':'submit_bound','attempt_no':attempt,'task_id':task_id}); r._save_runtime_state(temp,current)
    return {'status':'LAYER_REQUEST_PENDING','slide_id':slide_id,'task_id':task_id}


def result_current(state,slide_id):
    rec=state.get('layer_bridge',{}).get('results',{}).get(slide_id)
    if not isinstance(rec,dict): return None
    plan=_plan_current(state,slide_id)
    if plan is None or rec.get('run_id')!=plan.get('run_id') or rec.get('plan_sha256')!=plan.get('plan_sha256') or rec.get('text_clean_sha256')!=plan.get('text_clean_sha256'): return None
    path=_resolve(state,rec.get('path',''))
    if not path.is_file() or rec.get('sha256')!=_sha(path): return None
    return rec


def _bundle_current(state,slide_id):
    rec=state.get('layer_bridge',{}).get('bundles',{}).get(slide_id)
    if not isinstance(rec,dict): return None
    result=result_current(state,slide_id)
    if result is None or rec.get('result_sha256')!=result.get('sha256'): return None
    bundle=Path(rec.get('bundle_dir',''))
    manifest=bundle/'manifest.json'
    if not manifest.is_file() or rec.get('manifest_sha256')!=_sha(manifest): return None
    try: layer_validate_bundle.validate_bundle(bundle)
    except ValueError: return None
    return rec


def bind_bundle(state_path,slide_id,bundle_dir):
    state=r.load_runtime_state(state_path); result=result_current(state,slide_id)
    if result is None: raise ValueError('LAYER_RESULT_INVALID: current result required')
    bundle=Path(bundle_dir).resolve(); validated=layer_validate_bundle.validate_bundle(bundle); manifest=validated['manifest']
    plan=_plan_current(state,slide_id)
    if plan is None or manifest.get('page_id')!=slide_id: raise ValueError('LAYER_BUNDLE_VALIDATION_FAILED: lineage')
    if manifest.get('source',{}).get('text_clean_sha256')!=plan.get('text_clean_sha256'):
        raise ValueError('LAYER_BUNDLE_VALIDATION_FAILED: source changed')
    if manifest.get('schema_version')==3 and (manifest['source'].get('layer_plan_sha256')!=plan['plan_sha256']
        or manifest['backend'].get('result_sha256')!=result['sha256']):
        raise ValueError('LAYER_BUNDLE_VALIDATION_FAILED: plan/result identity changed')
    request_id=manifest.get('backend',{}).get('request_id')
    if request_id and result.get('task_id') and request_id!=result.get('task_id'):
        raise ValueError('LAYER_BUNDLE_VALIDATION_FAILED: provider task changed')
    target=router.require_image_target(state_path,state)
    actual=[manifest.get('ppt',{}).get('slide_width_emu'),manifest.get('ppt',{}).get('slide_height_emu')]
    if not isinstance(target,list) or len(target)!=2 or actual!=target:
        raise ValueError('TARGET_DECK_CANVAS_UNRESOLVED: bundle canvas differs from current Image Layer run')
    record={'status':'LAYER_BUNDLE_READY','slide_id':slide_id,'run_id':plan['run_id'],
            'plan_sha256':plan['plan_sha256'],'text_clean_sha256':plan['text_clean_sha256'],
            'result_sha256':result['sha256'],'bundle_dir':str(bundle),
            'manifest_sha256':_sha(bundle/'manifest.json')}
    with router.transaction(state_path) as temp:
        current=r.load_runtime_state(temp); router.require_image_target(temp,current)
        if result_current(current,slide_id)!=result or _plan_current(current,slide_id)!=plan:
            raise ValueError('LAYER_BUNDLE_VALIDATION_FAILED: state changed before binding')
        bridge=_ensure_bridge(current); bridge['bundles'][slide_id]=record; bridge['visual_qa'].pop(slide_id,None); bridge['graphics_first'].pop(slide_id,None); bridge['validated'].pop(slide_id,None); r._save_runtime_state(temp,current)
    return record



_VISUAL_QA_CHECKS=('layer_isolation','background_repair','recomposite_fidelity')


def _visual_qa_current(state,slide_id):
    rec=state.get('layer_bridge',{}).get('visual_qa',{}).get(slide_id)
    if not isinstance(rec,dict): return None
    bundle=_bundle_current(state,slide_id)
    if bundle is None or rec.get('bundle_manifest_sha256')!=bundle.get('manifest_sha256'): return None
    path=_resolve(state,rec.get('path',''))
    if not path.is_file() or rec.get('sha256')!=_sha(path): return None
    if rec.get('status') not in layer_policy.ACCEPTED: return None
    try:
        report=json.loads(path.read_text(encoding='utf-8'))
        manifest=json.loads((Path(bundle['bundle_dir'])/'manifest.json').read_text(encoding='utf-8'))
        status,limitations=layer_policy.validate_qa(report,{x['target_id'] for x in manifest['layers'][1:]},path.parent)
        if status!=rec['status'] or limitations!=rec.get('accepted_limitations',[]): return None
    except (ValueError,OSError,KeyError,TypeError): return None
    return rec


def register_visual_qa(state_path,slide_id,evidence_path):
    state=r.load_runtime_state(state_path); bundle=_bundle_current(state,slide_id)
    if bundle is None: raise ValueError('LAYER_VISUAL_QA_FAILED: current bundle required')
    path=Path(evidence_path).resolve()
    if not path.is_file(): raise ValueError('LAYER_VISUAL_QA_FAILED: evidence missing')
    try: report=json.loads(path.read_text(encoding='utf-8'))
    except (OSError,json.JSONDecodeError) as exc: raise ValueError('LAYER_VISUAL_QA_FAILED: invalid evidence') from exc
    if report.get('page_id')!=slide_id or report.get('bundle_manifest_sha256')!=bundle.get('manifest_sha256'):
        raise ValueError('LAYER_VISUAL_QA_FAILED: evidence lineage mismatch')
    manifest=json.loads((Path(bundle['bundle_dir'])/'manifest.json').read_text(encoding='utf-8'))
    targets={x['target_id'] for x in manifest['layers'][1:]}
    qa_status,limitations=layer_policy.validate_qa(report,targets,path.parent)
    record={'status':qa_status,'slide_id':slide_id,'bundle_manifest_sha256':bundle['manifest_sha256'],
            'path':str(path),'sha256':_sha(path),'checks':report['checks'],'accepted_limitations':limitations}
    with router.transaction(state_path) as temp:
        current=r.load_runtime_state(temp); bridge=_ensure_bridge(current); bridge['visual_qa'][slide_id]=record; bridge['graphics_first'].pop(slide_id,None); bridge['validated'].pop(slide_id,None); current.pop('merged_deck',None); current.pop('validated_deck',None); bridge.pop('merged',None); r._save_runtime_state(temp,current)
    if qa_status=='LAYER_VISUAL_QA_FAILED': raise ValueError('LAYER_VISUAL_QA_FAILED: serious error; no automatic retry')
    return record

def _canvas_residual_source_px(manifest,pptx_path):
    source=manifest.get('source',{}).get('canvas')
    target=manifest.get('ppt',{})
    if not isinstance(source,list) or len(source)!=2: raise ValueError('LAYER_GRAPHICS_FIRST_GATE_FAILED: source canvas missing')
    sw,sh=source
    tw,th=target.get('slide_width_emu'),target.get('slide_height_emu')
    if any(not isinstance(v,int) or v<=0 for v in (tw,th)):
        raise ValueError('LAYER_GRAPHICS_FIRST_GATE_FAILED: target deck canvas missing')
    prs=Presentation(pptx_path)
    if len(prs.slides)!=1 or prs.slide_width<=0 or prs.slide_height<=0:
        raise ValueError('LAYER_GRAPHICS_FIRST_GATE_FAILED: single-page PPTX required')
    sx=tw/sw; sy=th/sh
    dw=abs(prs.slide_width/sx-sw)
    dh=abs(prs.slide_height/sy-sh)
    return float(dw),float(dh)


def bind_graphics_first(state_path,slide_id,pptx_path,bundle_dir):
    state=r.load_runtime_state(state_path); bundle_rec=_bundle_current(state,slide_id)
    bundle=Path(bundle_dir).resolve()
    if bundle_rec is None or Path(bundle_rec.get('bundle_dir','')).resolve()!=bundle:
        raise ValueError('LAYER_BUNDLE_VALIDATION_FAILED: current bundle required')
    validated=layer_validate_bundle.validate_bundle(bundle); manifest=validated['manifest']; pptx=Path(pptx_path).resolve()
    qa=_visual_qa_current(state,slide_id)
    if qa is None: raise ValueError('LAYER_VISUAL_QA_FAILED: current visual QA required')
    if not pptx.is_file(): raise ValueError('LAYER_GRAPHICS_FIRST_GATE_FAILED: PPTX missing')
    dw,dh=_canvas_residual_source_px(manifest,pptx)
    if dw>5 or dh>5: raise ValueError(f'LAYER_GRAPHICS_FIRST_GATE_FAILED: source canvas residual {dw:.3f},{dh:.3f}px exceeds 5px')
    prs=Presentation(pptx)
    if len(prs.slides[0].shapes)!=manifest['counts']['actual_total_layers']:
        raise ValueError('LAYER_GRAPHICS_FIRST_GATE_FAILED: picture count mismatch')
    result=result_current(state,slide_id); digest=_sha(pptx)
    with router.transaction(state_path) as temp:
        r.register_stage2_artifact(temp,f'graphics_first_pptx:{slide_id}',pptx)
        current=r.load_runtime_state(temp); bridge=_ensure_bridge(current)
        record={'status':'LAYER_GRAPHICS_FIRST_READY','slide_id':slide_id,'run_id':bundle_rec['run_id'],
                'result_sha256':result['sha256'],'bundle_manifest_sha256':bundle_rec['manifest_sha256'],
                'visual_qa_sha256':qa['sha256'],'accepted_limitations':qa.get('accepted_limitations',[]),'path':str(pptx),'sha256':digest,'canvas_residual_source_px':[dw,dh]}
        bridge['graphics_first'][slide_id]=record; r._save_runtime_state(temp,current)
    return record


def graphics_first_current(state,slide_id):
    rec=state.get('layer_bridge',{}).get('graphics_first',{}).get(slide_id)
    if not isinstance(rec,dict): return None
    result=result_current(state,slide_id); bundle=_bundle_current(state,slide_id); qa=_visual_qa_current(state,slide_id)
    if (result is None or bundle is None or qa is None or rec.get('result_sha256')!=result.get('sha256')
            or rec.get('bundle_manifest_sha256')!=bundle.get('manifest_sha256')
            or rec.get('visual_qa_sha256')!=qa.get('sha256')): return None
    path=_resolve(state,rec.get('path',''))
    if not path.is_file() or rec.get('sha256')!=_sha(path): return None
    run=state.get('stage2_run',{}); registered=run.get('artifacts',{}).get('graphics_first_pptx:'+slide_id,{})
    slide_path=state.get('slides',{}).get(slide_id,{}).get('graphics_first_pptx')
    if (registered.get('run_id')!=run.get('run_id') or registered.get('sha256')!=rec.get('sha256')
            or not slide_path or _resolve(state,slide_path)!=path): return None
    return rec



def page_provenance(state,slide_id):
    sealed=state.get('slides',{}).get(slide_id,{}).get('validated_single_page')
    errors=r.validate_current_artifact(sealed,state,slide_id)
    if errors: raise ValueError(slide_id+': '+', '.join(errors))
    gf=graphics_first_current(state,slide_id)
    if gf is None: raise ValueError('current Image Layer Graphics-first binding missing: '+slide_id)
    binding=state.get('layer_bridge',{}).get('validated',{}).get(slide_id)
    if (not isinstance(binding,dict) or binding.get('backend')!='image_layer'
            or binding.get('graphics_first_sha256')!=gf.get('sha256')
            or binding.get('bundle_manifest_sha256')!=gf.get('bundle_manifest_sha256')
            or binding.get('visual_qa_sha256')!=gf.get('visual_qa_sha256')
            or binding.get('seal')!=sealed):
        raise ValueError('validated page belongs to superseded Image Layer inputs: '+slide_id)
    return dict(binding)


def seal_page(state_path,slide_id,pptx,gates_path):
    review=json.loads(Path(gates_path).read_text(encoding='utf-8'))
    gates=review.get('checks',{})
    if any(gates.get(k)!='PASS' for k in r.GATE_NAMES):
        raise ValueError('four Hard Gates must all be PASS')
    if any(not isinstance(review.get('evidence',{}).get(k),str) or not review['evidence'][k].strip() for k in r.GATE_NAMES):
        raise ValueError('actual four-gate evidence required')
    page=Path(pptx).resolve()
    prs=Presentation(page)
    if len(prs.slides)!=1: raise ValueError('single-page restored PPTX required')
    with router.transaction(state_path) as temp:
        r.require_stage2_entry(temp)
        state=r.load_runtime_state(temp); gf=graphics_first_current(state,slide_id)
        if gf is None: raise ValueError('current Image Layer Graphics-first binding missing: '+slide_id)
        if page==_resolve(state,gf['path']):
            raise ValueError('restored output must preserve the graphics-first source at a separate path')
        if (review.get('backend')!='image_layer' or review.get('slide_id')!=slide_id
                or review.get('graphics_first_sha256')!=gf['sha256']
                or review.get('restored_pptx_sha256')!=_sha(page)):
            raise ValueError('single-page review identity or file hash mismatch')
        layer_policy.validate_context(review,gf.get('accepted_limitations',[]))
        for kind in ('finalized_manifest','font_fallback'):
            if not r.stage2_artifact_current(temp,kind+':'+slide_id):
                raise ValueError('current registered '+kind+' required')
            payload=json.loads(_resolve(state,state['slides'][slide_id][kind]).read_text(encoding='utf-8'))
            if not isinstance(payload,dict): raise ValueError(kind+' must be JSON mapping')
        expected=r.expected_lineage_from_state(state,slide_id)['text_restore_fingerprint']
        if review.get('text_restore_fingerprint')!=expected:
            raise ValueError('single-page review restoration inputs changed')
        slide=state['slides'][slide_id]; old=slide.pop('validated_single_page',None)
        if old: slide.setdefault('superseded_validated_single_pages',[]).append(old)
        state.pop('merged_deck',None); state.pop('validated_deck',None)
        state.setdefault('layer_bridge',{}).pop('merged',None)
        r._save_runtime_state(temp,state)
        seal=r.seal_validated_single_page(temp,slide_id,page,gates)
        state=r.load_runtime_state(temp); bridge=_ensure_bridge(state)
        bridge['validated'][slide_id]={
            'backend':'image_layer','slide_id':slide_id,
            'graphics_first_sha256':gf['sha256'],
            'bundle_manifest_sha256':gf['bundle_manifest_sha256'],
            'visual_qa_sha256':gf['visual_qa_sha256'],
            'result_sha256':gf['result_sha256'],
            'accepted_limitations':gf.get('accepted_limitations',[]),
            'validated_pptx_sha256':seal['validated_pptx_sha256'],
            'seal':seal}
        r._save_runtime_state(temp,state)
    return {'status':'PAGE_SEALED','slide_id':slide_id,'path':str(page)}

def query(state_path,slide_id,work_dir,provider_module=layer_provider_360):
    state=r.load_runtime_state(state_path); plan_rec=_plan_current(state,slide_id)
    if plan_rec is None or not plan_rec.get('approved'): raise ValueError('LAYER_PLAN_NOT_APPROVED')
    api_key=os.environ.get(router.IMAGE_LAYER_API_KEY_ENV)
    if not api_key: return router.image_layer_onboarding()
    req=state.get('layer_bridge',{}).get('requests',{}).get(slide_id)
    if not isinstance(req,dict) or req.get('status')!='LAYER_REQUEST_PENDING' or not req.get('task_id'):
        raise ValueError('LAYER_REQUEST_PENDING required')
    if req.get('provider_identity') is not None and req['provider_identity']!=_provider_identity(provider_module):
        raise ValueError('LAYER_RESULT_INVALID: query provider/model parameters changed')
    response=provider_module.query_task(req['task_id'],api_key); normalized=response.get('normalized_status')
    if response.get('task_id') not in (None,req['task_id']): raise ValueError('LAYER_RESULT_INVALID: task identity mismatch')
    if normalized=='RUNNING': return {'status':'LAYER_REQUEST_PENDING','slide_id':slide_id,'task_id':req['task_id']}
    if normalized in ('FAILED','NOT_FOUND'):
        with router.transaction(state_path) as temp:
            current=r.load_runtime_state(temp); b=_ensure_bridge(current); rrec=b['requests'][slide_id]; rrec['status']='LAYER_RETRY_APPROVAL_REQUIRED'; rrec['terminal_status']=normalized
            b['call_ledger'].append({'slide_id':slide_id,'event':'query_terminal','task_id':rrec['task_id'],'terminal_status':normalized}); r._save_runtime_state(temp,current)
        return {'status':'LAYER_RETRY_APPROVAL_REQUIRED','slide_id':slide_id}
    if normalized!='DONE': raise ValueError('LAYER_RESULT_INVALID: unknown normalized status')
    current=r.load_runtime_state(state_path); plan=json.loads(_resolve(current,plan_rec['plan_path']).read_text(encoding='utf-8'))
    outdir=Path(work_dir).resolve()/slide_id/('provider-result-attempt-%02d'%req.get('attempt_no',1)); result_path=provider_module.normalize_done_result(response,plan,outdir)
    record={'status':'LAYER_RESULT_BOUND','slide_id':slide_id,'run_id':plan_rec['run_id'],'plan_sha256':plan_rec['plan_sha256'],
            'text_clean_sha256':plan_rec['text_clean_sha256'],'task_id':req['task_id'],'path':str(Path(result_path).resolve()),
            'sha256':_sha(result_path),'usage':response.get('usage',{})}
    with router.transaction(state_path) as temp:
        current=r.load_runtime_state(temp); b=_ensure_bridge(current)
        if b.get('requests',{}).get(slide_id)!=req or _plan_current(current,slide_id)!=plan_rec:
            raise ValueError('LAYER_RESULT_INVALID: state changed during query')
        b['results'][slide_id]=record; b['requests'][slide_id]['status']='DONE'; b['requests'][slide_id]['usage']=response.get('usage',{})
        b['call_ledger'].append({'slide_id':slide_id,'event':'result_bound','task_id':req['task_id']}); r._save_runtime_state(temp,current)
    return {'status':'LAYER_RESULT_BOUND','slide_id':slide_id,'path':record['path']}


def status(state_path):
    state=r.load_runtime_state(state_path); run=_current_run(state)
    run_order=run.get('page_order',[]); order=state.get('deck_order',run_order)
    if (not run_order or not isinstance(order,list) or not r.validate_deck_order(run_order,order)
            or not r.validate_deck_order(list(state.get('slides',{})),order)):
        return {'status':'BLOCKED','detail':'trusted page order invalid'}
    for sid in order:
        try:
            page_provenance(state,sid)
            continue
        except (OSError,ValueError,KeyError,TypeError):
            pass
        gf=graphics_first_current(state,sid)
        if gf is not None:
            return {'status':'RESTORE_TEXT','slide_id':sid}
        plan=_plan_current(state,sid)
        if plan is None: return {'status':'LAYER_PLAN_REQUIRED','slide_id':sid}
        if not plan.get('approved'): return {'status':'LAYER_PLAN_APPROVAL_REQUIRED','slide_id':sid}
        failed_qa=state.get('layer_bridge',{}).get('visual_qa',{}).get(sid,{})
        if failed_qa.get('status')=='LAYER_VISUAL_QA_FAILED': return {'status':'LAYER_VISUAL_QA_FAILED','slide_id':sid}
        result=result_current(state,sid)
        if result is not None:
            bundle=_bundle_current(state,sid)
            if bundle is None: return {'status':'LAYER_BUNDLE_REQUIRED','slide_id':sid}
            if _visual_qa_current(state,sid) is None: return {'status':'LAYER_VISUAL_QA_REQUIRED','slide_id':sid}
            return {'status':'LAYER_GRAPHICS_FIRST_REQUIRED','slide_id':sid}
        req=state.get('layer_bridge',{}).get('requests',{}).get(sid)
        if not isinstance(req,dict):
            ready=readiness_current(state,sid)
            return {'status':'LAYER_REQUEST_READY' if ready and ready['status']=='READY' else 'LAYER_PLAN_READINESS_REQUIRED','slide_id':sid}
        st=req.get('status')
        if st=='SUBMITTING': return {'status':'LAYER_REQUEST_UNCERTAIN','slide_id':sid}
        if st in ('LAYER_REQUEST_UNCERTAIN','LAYER_RETRY_APPROVAL_REQUIRED','LAYER_REQUEST_PENDING'):
            return {'status':st,'slide_id':sid}
        return {'status':'BLOCKED','slide_id':sid,'detail':'unrecognized Image Layer request state'}
    pages=[page_provenance(state,sid) for sid in order]
    merged=state.get('layer_bridge',{}).get('merged')
    if r.validate_current_merged_deck(state) or not isinstance(merged,dict) or merged.get('pages')!=pages:
        return {'status':'PREPARE_DECK','deck_order':order}
    if merged.get('merged_pptx_sha256')!=state['merged_deck']['merged_pptx_sha256']:
        return {'status':'BLOCKED','detail':'merged provenance hash mismatch'}
    validated=state.get('validated_deck',{})
    if validated.get('status')=='PASS' and validated.get('merged_pptx_sha256')==merged['merged_pptx_sha256']:
        return {'status':'COMPLETE','path':state['merged_deck']['merged_pptx_path']}
    return {'status':'DECK_VALIDATION','path':state['merged_deck']['merged_pptx_path']}


def main(argv=None):
    parser=argparse.ArgumentParser(); sub=parser.add_subparsers(dest='action',required=True)
    for action in ('init','approve-plan','submit','query','authorize-retry','status','bind-bundle','register-visual-qa','register-readiness','authorize-quality-retry','bind-graphics-first','seal-page'):
        p=sub.add_parser(action); p.add_argument('--state',required=True)
        if action not in ('init','status'): p.add_argument('--slide-id',required=True)
        if action in ('submit','query'): p.add_argument('--work-dir',required=True)
        if action=='authorize-retry': p.add_argument('--accept-duplicate-charge-risk',action='store_true')
        if action in ('bind-bundle','bind-graphics-first'): p.add_argument('--bundle',required=True)
        if action in ('register-visual-qa','register-readiness','authorize-quality-retry'): p.add_argument('--evidence',required=True)
        if action=='authorize-quality-retry': p.add_argument('--authorization',required=True)
        if action=='bind-graphics-first': p.add_argument('--pptx',required=True)
        if action=='seal-page': p.add_argument('--pptx',required=True); p.add_argument('--gates',required=True)
    p=sub.add_parser('register-plan'); p.add_argument('--state',required=True); p.add_argument('--slide-id',required=True); p.add_argument('--plan',required=True); p.add_argument('--overlay',required=True)
    args=parser.parse_args(argv)
    try:
        if args.action=='init': result=init(args.state)
        elif args.action=='register-plan': result=register_plan(args.state,args.slide_id,args.plan,args.overlay)
        elif args.action=='approve-plan': result=approve_plan(args.state,args.slide_id)
        elif args.action=='submit': result=submit(args.state,args.slide_id,args.work_dir)
        elif args.action=='query': result=query(args.state,args.slide_id,args.work_dir)
        elif args.action=='authorize-retry': result=authorize_retry(args.state,args.slide_id,accept_duplicate_charge_risk=args.accept_duplicate_charge_risk)
        elif args.action=='register-readiness': result=register_readiness(args.state,args.slide_id,args.evidence)
        elif args.action=='authorize-quality-retry': result=authorize_quality_retry(args.state,args.slide_id,args.evidence,args.authorization)
        elif args.action=='bind-bundle': result=bind_bundle(args.state,args.slide_id,args.bundle)
        elif args.action=='register-visual-qa': result=register_visual_qa(args.state,args.slide_id,args.evidence)
        elif args.action=='bind-graphics-first': result=bind_graphics_first(args.state,args.slide_id,args.pptx,args.bundle)
        elif args.action=='seal-page': result=seal_page(args.state,args.slide_id,args.pptx,args.gates)
        else: result=status(args.state)
    except (OSError,ValueError,KeyError,TypeError,json.JSONDecodeError) as exc:
        result={'status':'FAIL','details':[str(exc)]}
    print(json.dumps(result,ensure_ascii=False,indent=2)); return 1 if result.get('status')=='FAIL' else 0

if __name__=='__main__': raise SystemExit(main())
