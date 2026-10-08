"""Office-native merge of validated single-page decks.

Prefer PowerPoint; WPS uses the same import and integrity checks. Application
selection ends before real import; failures never retry another merge engine.
"""
import argparse
import hashlib
import json
import os
import posixpath
import tempfile
import zipfile
from pathlib import Path
from urllib.parse import unquote
from xml.etree import ElementTree as ET

from runtime import (validate_deck_order, load_runtime_state, validate_current_artifact,
                     invalidate_merged_deck, seal_merged_deck)

REL_NS = '{http://schemas.openxmlformats.org/package/2006/relationships}'
GEOMETRY_TOLERANCE_EMU = 0  # measured: 44 coordinates, two source slides, zero native drift


def inspect_ooxml_dependencies(path):
    """Return broken package relationships (empty means intact)."""
    path = Path(path)
    if not path.is_file():
        return [f'missing package: {path}']
    errors = []
    try:
        with zipfile.ZipFile(path) as z:
            members = set(z.namelist())
            if '[Content_Types].xml' not in members or 'ppt/presentation.xml' not in members:
                errors.append('missing OOXML root')
            for rel_file in sorted(n for n in members if n.endswith('.rels')):
                if rel_file == '_rels/.rels':
                    source_dir = ''
                else:
                    parent, name = rel_file.rsplit('/_rels/', 1)
                    source = posixpath.join(parent, name[:-5])
                    source_dir = posixpath.dirname(source)
                try:
                    root = ET.fromstring(z.read(rel_file))
                except ET.ParseError:
                    errors.append(f'invalid relationships XML: {rel_file}')
                    continue
                seen = set()
                for rel in root.findall(f'{REL_NS}Relationship'):
                    rid = rel.get('Id')
                    if rid in seen:
                        errors.append(f'duplicate relationship ID: {rel_file} {rid}')
                    seen.add(rid)
                    if rel.get('TargetMode') == 'External':
                        continue
                    target = unquote(rel.get('Target', '')).split('#', 1)[0]
                    if target.startswith('/'):
                        resolved = target.lstrip('/')
                    else:
                        resolved = posixpath.normpath(posixpath.join(source_dir, target))
                    if resolved not in members:
                        errors.append(f'broken relationship: {rel_file} {rid} -> {resolved}')
    except (zipfile.BadZipFile, OSError) as exc:
        errors.append(f'invalid PPTX: {exc}')
    return errors


def _merge_input_diagnostics(inputs, deck_order, runtime_state_path=None):
    ids = [x.get('slide_id') for x in inputs]
    if not validate_deck_order(ids, deck_order) or any(not x.get('validated') for x in inputs):
        return 'DECK_MERGE_INTEGRITY_FAILED', ['invalid_deck_order_or_unvalidated_page']
    if any(not x.get('path') or not Path(x['path']).is_file() for x in inputs):
        return 'DECK_MERGE_INTEGRITY_FAILED', ['validated_pptx_missing']
    try:
        from pptx import Presentation
        sizes = []
        for x in inputs:
            prs = Presentation(x['path'])
            if len(prs.slides) != 1:
                return 'DECK_MERGE_INTEGRITY_FAILED', ['input_not_single_page']
            sizes.append((prs.slide_width, prs.slide_height))
            if inspect_ooxml_dependencies(x['path']):
                return 'DECK_MERGE_INTEGRITY_FAILED', ['input_ooxml_dependency_invalid']
        if len(set(sizes)) != 1:
            return 'DECK_CANVAS_INVARIANT_FAILED', ['canvas_mismatch']
    except Exception:
        return 'DECK_MERGE_INTEGRITY_FAILED', ['invalid_input_pptx']
    if runtime_state_path is None:
        return 'DECK_MERGE_INTEGRITY_FAILED', ['trusted_runtime_state_unavailable']
    try:
        state = load_runtime_state(runtime_state_path)
    except (OSError, ValueError, TypeError):
        return 'DECK_MERGE_INTEGRITY_FAILED', ['trusted_runtime_state_unavailable']
    if state.get('deck_order') != deck_order:
        return 'DECK_MERGE_INTEGRITY_FAILED', ['deck_order_mismatch']
    reasons = []
    for item in inputs:
        reasons.extend(validate_current_artifact(item, state, item['slide_id']))
    if reasons:
        return 'DECK_MERGE_INTEGRITY_FAILED', list(dict.fromkeys(reasons))
    return None, []


def validate_merge_inputs(inputs, deck_order, runtime_state_path=None):
    return _merge_input_diagnostics(inputs, deck_order, runtime_state_path)[0]

def _shape_signature(prs):
    """Ordered visual structure, excluding volatile OOXML identities."""
    pages = []
    for slide in prs.slides:
        shapes = []
        for shape in slide.shapes:
            is_picture = int(shape.shape_type) == 13
            shapes.append({
                'type': int(shape.shape_type),
                'geometry': (int(shape.left), int(shape.top), int(shape.width), int(shape.height)),
                'text': shape.text if shape.has_text_frame else None,
                'has_chart': bool(shape.has_chart),
                'category': 'picture' if is_picture else ('chart' if shape.has_chart else 'other'),
                'picture_sha256': hashlib.sha256(shape.image.blob).hexdigest() if is_picture else None,
            })
        pages.append(shapes)
    return pages


def _shape_signature_errors(expected, actual, tolerance_emu):
    if len(expected) != len(actual):
        return [f'shape_count_changed {len(expected)} -> {len(actual)}']
    errors = []
    for index, (before, after) in enumerate(zip(expected, actual), 1):
        for field in ('type', 'text', 'has_chart', 'category', 'picture_sha256'):
            if before[field] != after[field]:
                errors.append(f'shape_{index}_{field}_changed')
        if any(abs(a-b) > tolerance_emu for a,b in zip(before['geometry'], after['geometry'])):
            errors.append(f'shape_{index}_geometry_changed')
    return errors

def _design_signature(prs):
    def fill_sig(fill):
        kind = str(fill.type)
        if kind.startswith('SOLID'):
            color = fill.fore_color
            return (kind, str(color.type), str(color.rgb) if color.type and str(color.type).startswith('RGB') else str(color.theme_color))
        return (kind,)
    return [(tuple(sorted((ph.placeholder_format.idx, str(ph.placeholder_format.type)) for ph in slide.slide_layout.placeholders)),
             fill_sig(slide.slide_layout.slide_master.background.fill),
             fill_sig(slide.background.fill)) for slide in prs.slides]

OFFICE_PROGIDS = {'powerpoint': 'PowerPoint.Application', 'wps': 'KWPP.Application'}
OFFICE_HOSTS = ('auto', 'powerpoint', 'wps')


def _create_office_app(office_host, diagnostics):
    if office_host not in OFFICE_HOSTS:
        raise ValueError('unknown office host: ' + str(office_host))
    import win32com.client
    hosts = ('powerpoint', 'wps') if office_host == 'auto' else (office_host,)
    diagnostics.update(office_host=None, office_progid=None, office_attempts=[])
    for host in hosts:
        progid = OFFICE_PROGIDS[host]
        try:
            app = win32com.client.DispatchEx(progid)
        except Exception as exc:
            diagnostics['office_attempts'].append({'host': host, 'progid': progid, 'status': 'FAIL', 'detail': str(exc)})
            if host == hosts[-1]:
                raise RuntimeError('native application unavailable: ' + '; '.join(x['detail'] for x in diagnostics['office_attempts'])) from exc
        else:
            diagnostics['office_attempts'].append({'host': host, 'progid': progid, 'status': 'CREATED'})
            diagnostics.update(office_host=host, office_progid=progid)
            return app


def _native_merge(inputs, order, target, *, office_host='auto', diagnostics=None):
    import win32com.client
    diagnostics = {} if diagnostics is None else diagnostics
    app = _create_office_app(office_host, diagnostics)
    selected = diagnostics['office_host']
    prs = None
    failure = None
    try:
        app.Visible = True
        prs = app.Presentations.Add(-1)
        from pptx import Presentation
        source = Presentation(inputs[0]['path'])
        prs.PageSetup.SlideWidth = source.slide_width / 12700
        prs.PageSetup.SlideHeight = source.slide_height / 12700
        by_id = {x['slide_id']: x for x in inputs}
        for slide_id in order:
            inserted = prs.Slides.InsertFromFile(str(Path(by_id[slide_id]['path']).resolve()), prs.Slides.Count)
            if inserted != 1:
                raise RuntimeError(f'native import produced {inserted} slides for {slide_id}')
            design = prs.Designs.Load(str(Path(by_id[slide_id]['path']).resolve()))
            prs.Slides(prs.Slides.Count).Design = design
        prs.SaveAs(str(target), 24)
        prs.Close(); prs = None
        if selected == 'wps':
            app = None
            app = win32com.client.DispatchEx(OFFICE_PROGIDS[selected])
        prs = app.Presentations.Open(str(target), 0, 0, 0)
        actual = prs.Slides.Count
        prs.Save()
        prs.Close()
        prs = None
        if actual != len(order):
            raise RuntimeError(f'{selected} reopened {actual} slides; expected {len(order)}')
    except BaseException as exc:
        failure = exc
        raise
    finally:
        cleanup = []
        if prs is not None:
            try:
                prs.Close()
            except Exception as exc:
                cleanup.append('document close: ' + str(exc))
        if selected == 'powerpoint' and app is not None:
            try:
                app.Quit()
            except Exception as exc:
                cleanup.append('application quit: ' + str(exc))
        app = None
        if cleanup:
            diagnostics['cleanup_errors'] = cleanup
            if failure is None:
                raise RuntimeError('native cleanup failed: ' + '; '.join(cleanup))
    return diagnostics


def merge(inputs, deck_order, output, runtime_state_path=None, *, office_host='auto'):
    """Merge validated single pages; return report. Output appears only on success."""
    if office_host not in OFFICE_HOSTS:
        raise ValueError('unknown office host: ' + str(office_host))
    code, reasons = _merge_input_diagnostics(inputs, deck_order, runtime_state_path)
    if code:
        return {'status': 'FAIL', 'code': code, 'details': reasons}
    from pptx import Presentation
    by_id = {x['slide_id']: x for x in inputs}
    ordered = [by_id[i] for i in deck_order]
    expected = [_shape_signature(Presentation(x['path']))[0] for x in ordered]
    expected_design = [_design_signature(Presentation(x['path']))[0] for x in ordered]
    size = (Presentation(ordered[0]['path']).slide_width, Presentation(ordered[0]['path']).slide_height)
    output = Path(output).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent) as tmp:
        candidate = Path(tmp)/'merged.pptx'
        native = {}
        try:
            _native_merge(inputs, deck_order, candidate, office_host=office_host, diagnostics=native)
            actual_prs = Presentation(candidate)
            errors = inspect_ooxml_dependencies(candidate)
            if len(actual_prs.slides) != len(deck_order):
                errors.append('slide count changed')
            if (actual_prs.slide_width, actual_prs.slide_height) != size:
                errors.append('deck canvas changed')
            actual = _shape_signature(actual_prs)
            actual_design = _design_signature(actual_prs)
            if actual_design != expected_design:
                errors.append(f'master/layout/background changed: {expected_design} -> {actual_design}')
            for i, (before, after) in enumerate(zip(expected, actual)):
                errors.extend(f'slide_{i+1}_{reason}' for reason in _shape_signature_errors(before, after, GEOMETRY_TOLERANCE_EMU))
            if errors:
                return {'status':'FAIL', 'code':'DECK_MERGE_INTEGRITY_FAILED', 'details':errors, **native}
            invalidate_merged_deck(runtime_state_path, deck_order, ordered)
            os.replace(candidate, output)
            seal_merged_deck(runtime_state_path, deck_order, ordered, output)
            return {'status':'PASS', 'path':str(output), 'deck_order':deck_order, 'slides':len(deck_order), 'ooxml_relationships':'PASS', 'native_reopen':'PASS', **native}
        except Exception as exc:
            return {'status':'FAIL', 'code':'DECK_MERGE_INTEGRITY_FAILED', 'details':[str(exc)], **native}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('job_json')
    parser.add_argument('--runtime-state', required=True)
    parser.add_argument('--office-host', choices=OFFICE_HOSTS, default='auto')
    args = parser.parse_args()
    job = json.loads(Path(args.job_json).read_text(encoding='utf-8'))
    result = merge(job['inputs'], job['deck_order'], job['output'], args.runtime_state, office_host=args.office_host)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result['status'] == 'PASS' else 1)

if __name__ == '__main__':
    main()
