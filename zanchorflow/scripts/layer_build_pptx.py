"""Build one canonical textless PowerPoint slide from a validated Layer Bundle v2."""
import argparse
import json
from pathlib import Path
from pptx import Presentation
import layer_validate_bundle


def build_page_pptx(bundle_dir, output_path):
    bundle=Path(bundle_dir).resolve(); output=Path(output_path).resolve()
    validated=layer_validate_bundle.validate_bundle(bundle); manifest=validated['manifest']
    prs=Presentation(); prs.slide_width=int(manifest['ppt']['slide_width_emu']); prs.slide_height=int(manifest['ppt']['slide_height_emu'])
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    layers=sorted(manifest['layers'],key=lambda x:x['z_index'])
    for idx,spec in enumerate(layers):
        x,y,w,h=spec['canonical_bbox_normalized']; path=bundle/spec['file']
        pic=slide.shapes.add_picture(str(path),round(prs.slide_width*x),round(prs.slide_height*y),
                                     width=round(prs.slide_width*w),height=round(prs.slide_height*h))
        if idx==0: pic.name='Background'
        else: pic.name=f"Foreground {idx:02d} - {spec['target_id']}"
    output.parent.mkdir(parents=True,exist_ok=True); prs.save(output)
    check=Presentation(output)
    if len(check.slides)!=1 or len(check.slides[0].shapes)!=len(layers): raise ValueError('LAYER_PPTX_BUILD_FAILED')
    return {'status':'PASS','path':str(output),'slide_width_emu':check.slide_width,'slide_height_emu':check.slide_height,'picture_count':len(layers)}


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__); parser.add_argument('--bundle',required=True); parser.add_argument('--output',required=True); args=parser.parse_args(argv)
    try: result=build_page_pptx(args.bundle,args.output)
    except (OSError,ValueError,KeyError,TypeError,json.JSONDecodeError) as exc: result={'status':'FAIL','details':[str(exc)]}
    print(json.dumps(result,ensure_ascii=False,indent=2)); return 1 if result.get('status')=='FAIL' else 0

if __name__=='__main__': raise SystemExit(main())
