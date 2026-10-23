"""Local native-object builder. Does not grant provenance or visual acceptance."""
import copy,hashlib,json,math,os,tempfile,re
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls,qn
from editable_geometry import validate_shape,custom_path_geometry_xml

KINDS={'text','table','shape','path','image'}
def numbers(values,count,positive=False):
 if not isinstance(values,(list,tuple)) or len(values)!=count or any(isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v) for v in values):raise ValueError('finite numeric coordinates required')
 if positive and any(v<=0 for v in values):raise ValueError('positive dimensions required')
 return values

def validate(plan):
 if plan.get('schema_version')!=1:raise ValueError('unsupported editable plan version')
 sw,sh=numbers(plan.get('source_size_px'),2,True);tw,th=numbers(plan.get('slide_size_emu'),2,True)
 if abs(tw*sh/th-sw)>5 or abs(th*sw/tw-sh)>5:raise ValueError('canvas aspect mismatch')
 objects=plan.get('objects');ids=set();occ=set();logical_lines=set()
 if not isinstance(objects,list) or not objects:raise ValueError('objects required')
 def occurrences(item):
  for ref in item.get('occurrence_ids',[]):
   if not isinstance(ref,str) or not ref or ref in occ:raise ValueError('duplicate or invalid occurrence')
   occ.add(ref)
 for item in objects:
  name=item.get('id')
  if not isinstance(name,str) or not name or name in ids:raise ValueError('duplicate or invalid object id')
  ids.add(name)
  if item.get('type') not in KINDS:raise ValueError('unsupported object type')
  x,y,w,h=numbers(item.get('box_px'),4)
  if min(x,y)<0 or min(w,h)<=0 or x+w>sw or y+h>sh:raise ValueError('invalid source box')
  if item['type']=='path':validate_shape(item)
  if item['type'] in ('shape','path') and item.get('semantic_line_id') is not None:
   line_id=item['semantic_line_id']
   if not isinstance(line_id,str) or not line_id.strip() or line_id in logical_lines:raise ValueError('logical line split across multiple objects')
   logical_lines.add(line_id)
  if 'source_corner_radius_px' in item:
   radius=numbers([item['source_corner_radius_px']],1)[0]
   if item['type']!='shape' or item.get('preset')!='roundRect' or not 0<=radius<=min(w,h)/2:raise ValueError('invalid source corner radius')
  if item['type']=='text':
   if not isinstance(item.get('text'),str):raise ValueError('text required')
   occurrences(item)
  if item['type'] in ('text','shape','path'):
   validate_style(item)
  if item['type']=='table':
   cells=item.get('cells')
   if not isinstance(cells,list) or not cells or not cells[0]:raise ValueError('table cells required')
   nr,nc=len(cells),len(cells[0]);covered=set()
   if any(not isinstance(row,list) or len(row)!=nc for row in cells):raise ValueError('rectangular table required')
   for key,n in [('row_heights',nr),('column_widths',nc)]:numbers(item.get(key,[1]*n),n,True)
   for ri,row in enumerate(cells):
    for ci,cell in enumerate(row):
     if not isinstance(cell,dict) or not isinstance(cell.get('text',''),str):raise ValueError('cell must be a text object')
     if (ri,ci) in covered:
      if cell not in ({},{'text':''}):raise ValueError('covered cell contains hidden content')
      continue
     rs,cs=cell.get('row_span',1),cell.get('col_span',1)
     if type(rs)!=int or type(cs)!=int or min(rs,cs)<1 or ri+rs>nr or ci+cs>nc:raise ValueError('invalid merge')
     for rr in range(ri,ri+rs):
      for cc in range(ci,ci+cs):
       if (rr,cc) in covered:raise ValueError('overlapping merge')
       if (rr,cc)!=(ri,ci):covered.add((rr,cc))
     occurrences(cell);validate_style({**item.get('style',{}),**cell.get('style',{}),**cell})
  if item['type']=='image':
   occurrences(item)
   if item.get('fit','contain') not in ('contain','cover','stretch'):raise ValueError('invalid image fit')
   file=Path(item.get('file',''))
   if not file.is_file() or hashlib.sha256(file.read_bytes()).hexdigest()!=item.get('sha256'):raise ValueError('image identity mismatch')
 for formula in plan.get('formula_inventory',[]):
  if not isinstance(formula,dict) or formula.get('object_id') not in ids or formula.get('decision') not in ('native_text','latex_rendered','source_preserved'):raise ValueError('unresolved formula')
  obj=next(v for v in objects if v['id']==formula['object_id'])
  if formula['decision']=='native_text' and obj['type']!='text':raise ValueError('formula must remain native text')
  if formula['decision']!='native_text' and obj['type']!='image':raise ValueError('formula image missing')
  if formula['decision']=='latex_rendered':
   tex=Path(formula.get('tex_source',''))
   if not tex.is_file() or hashlib.sha256(tex.read_bytes()).hexdigest()!=formula.get('tex_sha256'):raise ValueError('formula source identity mismatch')
 return occ

def validate_style(item):
 for key in ('fill','stroke','color'):
  if key in item and item[key] not in (None,'none') and (not isinstance(item[key],str) or not re.fullmatch(r'#?[0-9A-Fa-f]{6}',item[key])):raise ValueError('invalid color')
 for key in ('font_size','stroke_width'):
  if key in item:numbers([item[key]],1,True)
 for key in ('rotation','z_index','baseline','letter_spacing','space_before','space_after'):
  if key in item:numbers([item[key]],1)
 if item.get('align','left') not in ('left','center','right'):raise ValueError('invalid alignment')
 if item.get('valign','top') not in ('top','middle','bottom'):raise ValueError('invalid vertical alignment')
 if 'paragraphs' in item:
  if not isinstance(item['paragraphs'],list) or not item['paragraphs']:raise ValueError('paragraph list required')
  observed=[]
  for para in item['paragraphs']:
   if not isinstance(para,dict) or not isinstance(para.get('runs'),list):raise ValueError('run list required')
   for run in para['runs']:
    if not isinstance(run,dict) or not isinstance(run.get('text'),str):raise ValueError('run text required')
    validate_style({k:v for k,v in run.items() if k!='paragraphs'})
   observed.append(''.join(run['text'] for run in para['runs']))
  if '\n'.join(observed)!=item.get('text',''):raise ValueError('formatted text mismatch')


def fill(shape,value):
 if value in (None,'none'):shape.fill.background()
 else:shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(value.lstrip('#'))

def text(frame,item):
 frame.clear();frame.word_wrap=item.get('wrap',True)
 for side in ('left','right','top','bottom'):setattr(frame,'margin_'+side,Pt(item.get('margin_'+side,0)))
 frame.vertical_anchor={'top':MSO_ANCHOR.TOP,'middle':MSO_ANCHOR.MIDDLE,'bottom':MSO_ANCHOR.BOTTOM}[item.get('valign','top')]
 paragraphs=item.get('paragraphs') or [{'runs':[{'text':v}]} for v in item.get('text','').split('\n')]
 for index,spec in enumerate(paragraphs):
  paragraph_style={**item,**spec}
  p=frame.paragraphs[0] if index==0 else frame.add_paragraph();p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[paragraph_style.get('align','left')]
  for key in ('space_before','space_after'):
   if key in paragraph_style:setattr(p,key,Pt(paragraph_style[key]))
  if 'line_spacing' in paragraph_style:p.line_spacing=paragraph_style['line_spacing']
  for run in spec.get('runs',[]):
   rr=p.add_run();rr.text=run['text'];style={**item,**run};rr.font.name=style.get('font','Microsoft YaHei');rr.font.size=Pt(style.get('font_size',18));rr.font.bold=bool(style.get('bold',False));rr.font.italic=bool(style.get('italic',False));rr.font.underline=bool(style.get('underline',False));rr.font.color.rgb=RGBColor.from_string(style.get('color','111111').lstrip('#'))
   rp=rr._r.get_or_add_rPr()
   for tag in ('ea','cs'):
    child=rp.find(qn('a:'+tag))
    if child is None:child=parse_xml('<a:'+tag+' '+nsdecls('a')+'/>');rp.append(child)
    child.set('typeface',rr.font.name)
   if 'baseline' in style:rp.set('baseline',str(int(style['baseline'])))
   if 'letter_spacing' in style:rp.set('spc',str(round(style['letter_spacing']*100)))

def cell_borders(cell,borders):
 prop=cell._tc.get_or_add_tcPr()
 for side,spec in borders.items():
  if side not in ('left','right','top','bottom'):raise ValueError('invalid cell border')
  tag={'left':'lnL','right':'lnR','top':'lnT','bottom':'lnB'}[side]
  old=prop.find(qn('a:'+tag))
  if old is not None:prop.remove(old)
  line=parse_xml('<a:'+tag+' '+nsdecls('a')+' w="'+str(round(spec.get('width',1)*12700))+'"/>')
  if spec.get('color') in (None,'none'):line.append(parse_xml('<a:noFill '+nsdecls('a')+'/>'))
  else:line.append(parse_xml('<a:solidFill '+nsdecls('a')+'><a:srgbClr val="'+spec['color'].lstrip('#')+'"/></a:solidFill>'))
  insert_at=next((i for i,child in enumerate(prop) if child.tag not in {qn('a:'+tag) for tag in ('lnL','lnR','lnT','lnB','lnTlToBr','lnBlToTr')}),len(prop))
  prop.insert(insert_at,line)


def build(plan,output):
 plan=copy.deepcopy(plan);expected=validate(plan);output=Path(output)
 manifest_path=output.with_suffix('.manifest.json')
 if output.exists() or manifest_path.exists():raise ValueError('output exists')
 sw,sh=plan['source_size_px'];tw,th=plan['slide_size_emu'];prs=Presentation();prs.slide_width=round(tw);prs.slide_height=round(th);slide=prs.slides.add_slide(prs.slide_layouts[6]);occ={};mapping={};resolved=[]
 def box(item):
  x,y,w,h=item['box_px'];return [round(x*tw/sw),round(y*th/sh),round(w*tw/sw),round(h*th/sh)]
 def bind(item,shape,cell=None):
  for ref in item.get('occurrence_ids',[]):occ[ref]={'shape_id':shape.shape_id,**({'cell':cell} if cell is not None else {})}
 for item in sorted(plan['objects'],key=lambda v:v.get('z_index',0)):
  kind=item['type'];coords=box(item)
  if kind=='text':shape=slide.shapes.add_textbox(*coords);text(shape.text_frame,item);bind(item,shape)
  elif kind=='table':
   cells=item['cells'];shape=slide.shapes.add_table(len(cells),len(cells[0]),*coords);table=shape.table
   for key,seq,extent in [('column_widths',table.columns,coords[2]),('row_heights',table.rows,coords[3])]:
    weights=item.get(key,[1]*len(seq));parts=[round(extent*v/sum(weights)) for v in weights];parts[-1]+=extent-sum(parts)
    for obj,value in zip(seq,parts):setattr(obj,'width' if key=='column_widths' else 'height',value)
   for ri,row in enumerate(cells):
    for ci,data in enumerate(row):
     c=table.cell(ri,ci)
     if c.is_spanned:continue
     rs,cs=data.get('row_span',1),data.get('col_span',1)
     if rs>1 or cs>1:c.merge(table.cell(ri+rs-1,ci+cs-1))
     style={**item.get('style',{}),**data.get('style',{}),'text':data.get('text','')}
     if 'paragraphs' in data:style['paragraphs']=data['paragraphs']
     text(c.text_frame,style)
     c.vertical_anchor={'top':MSO_ANCHOR.TOP,'middle':MSO_ANCHOR.MIDDLE,'bottom':MSO_ANCHOR.BOTTOM}[style.get('valign','top')]
     for side in ('left','right','top','bottom'):setattr(c,'margin_'+side,Pt(style.get('margin_'+side,0)))
     fill(c,style.get('fill','FFFFFF'))
     cell_borders(c,style.get('borders',{}))
     bind(data,shape,[ri,ci])
  elif kind=='image':
   from PIL import Image
   with Image.open(item['file']) as im:iw,ih=im.size
   x,y,w,h=coords;mode=item.get('fit','contain')
   if mode=='contain':
    ratio=min(w/iw,h/ih);nw,nh=round(iw*ratio),round(ih*ratio);shape=slide.shapes.add_picture(item['file'],x+(w-nw)//2,y+(h-nh)//2,nw,nh)
   else:
    shape=slide.shapes.add_picture(item['file'],*coords)
    if mode=='cover':
     ratio=max(w/iw,h/ih);cropx=(iw-w/ratio)/iw/2;cropy=(ih-h/ratio)/ih/2
     shape.crop_left=shape.crop_right=cropx;shape.crop_top=shape.crop_bottom=cropy
   bind(item,shape)
  else:
   preset={'rect':MSO_SHAPE.RECTANGLE,'roundRect':MSO_SHAPE.ROUNDED_RECTANGLE,'ellipse':MSO_SHAPE.OVAL}.get(item.get('preset','rect'))
   if preset is None:raise ValueError('unsupported shape preset')
   shape=slide.shapes.add_shape(preset,*coords);fill(shape,item.get('fill','none'))
   if 'source_corner_radius_px' in item:shape.adjustments[0]=item['source_corner_radius_px']/min(item['box_px'][2:])
   if kind=='path':
    geom=parse_xml(custom_path_geometry_xml(item).replace('<a:custGeom>','<a:custGeom '+nsdecls('a')+'>',1));shape._element.spPr.remove(shape._element.spPr.find(qn('a:prstGeom')));shape._element.spPr.insert(1,geom)
   if item.get('stroke','000000') in (None,'none'):shape.line.fill.background()
   else:shape.line.color.rgb=RGBColor.from_string(item.get('stroke','000000').lstrip('#'))
   shape.line.width=Pt(item.get('stroke_width',1));line=shape._element.spPr.find(qn('a:ln'))
   effect=shape._element.find('.//'+qn('a:effectRef'))
   if effect is not None:effect.set('idx','0')
   if item.get('dash'):line.append(parse_xml('<a:prstDash '+nsdecls('a')+' val="'+item['dash']+'"/>'))
   for field,tag in [('start_arrow','headEnd'),('end_arrow','tailEnd')]:
    if item.get(field)=='triangle':line.append(parse_xml('<a:'+tag+' '+nsdecls('a')+' type="triangle"/>'))
  shape.name='editable:'+item['id'];shape.rotation=item.get('rotation',0);mapping[item['id']]=shape.shape_id
  resolved.append({'id':item['id'],'shape_id':shape.shape_id,'type':kind,'box_emu':[shape.left,shape.top,shape.width,shape.height],'rotation':shape.rotation})
 assert set(occ)==expected
 output.parent.mkdir(parents=True,exist_ok=True)
 fd,tmp=tempfile.mkstemp(suffix='.pptx',dir=output.parent);os.close(fd)
 try:
  prs.save(tmp);check=Presentation(tmp)
  if len(check.slides)!=1:raise ValueError('invalid page build')
  with output.open('xb') as target:target.write(Path(tmp).read_bytes())
 finally:Path(tmp).unlink(missing_ok=True)
 manifest={'schema_version':1,'plan':plan,'pptx_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'objects':resolved,'occurrences':occ}
 raw=json.dumps(manifest,ensure_ascii=False,sort_keys=True,indent=2).encode('utf-8')
 with manifest_path.open('xb') as stream:stream.write(raw)
 return {'resolved_manifest':str(manifest_path),'resolved_manifest_sha256':hashlib.sha256(raw).hexdigest(),'status':'BUILT_NOT_ASSESSED','path':str(output),'sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'objects':mapping,'occurrences':occ}
