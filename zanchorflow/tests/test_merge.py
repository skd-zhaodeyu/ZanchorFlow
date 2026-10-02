import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from merge_pptx import validate_merge_inputs, inspect_ooxml_dependencies

def test_merge_rejects_duplicate_id_before_import(tmp_path):
    inputs = [{'slide_id':'S1','path':str(tmp_path/'a.pptx'),'validated':True}, {'slide_id':'S1','path':str(tmp_path/'b.pptx'),'validated':True}]
    assert validate_merge_inputs(inputs, ['S1','S1']) == 'DECK_MERGE_INTEGRITY_FAILED'

def test_ooxml_inspector_reports_missing_file(tmp_path):
    assert inspect_ooxml_dependencies(tmp_path/'missing.pptx')

def test_merge_rejects_multislide_and_canvas_mismatch(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    a, b = tmp_path/'a.pptx', tmp_path/'b.pptx'
    pa = Presentation()
    pa.slides.add_slide(pa.slide_layouts[6])
    pa.slides.add_slide(pa.slide_layouts[6])
    pa.save(a)
    pb = Presentation()
    pb.slides.add_slide(pb.slide_layouts[6])
    pb.save(b)
    items = [{'slide_id':'S1','path':str(a),'validated':True}, {'slide_id':'S2','path':str(b),'validated':True}]
    assert validate_merge_inputs(items, ['S1','S2']) == 'DECK_MERGE_INTEGRITY_FAILED'
    pa = Presentation()
    pa.slide_width = Inches(12)
    pa.slides.add_slide(pa.slide_layouts[6])
    pa.save(a)
    assert validate_merge_inputs(items, ['S1','S2']) == 'DECK_CANVAS_INVARIANT_FAILED'


# Assembly entry tests use fake native merge and local bytes only: no COM or Canva.
def _assembly_fixture(tmp_path, monkeypatch):
    import json
    import assemble_deck as assembly
    import runtime
    state = {'stage2_run':{'page_order':['S001','S002']},'slides':{}}
    for slide_id in ('S001','S002'):
        slide = {}
        for key, value in (
            ('approved_render',b'render'),('text_clean',b'clean'),
            ('final_content_truth',{'title':slide_id}),('canvas',{'width':1600,'height':900}),
            ('finalized_manifest',{'title':slide_id}),('font_fallback',{}),
        ):
            path = tmp_path/(slide_id+'-'+key)
            if isinstance(value,bytes): path.write_bytes(value)
            else: path.write_text(json.dumps(value),encoding='utf-8')
            slide[key]=str(path)
        state['slides'][slide_id]=slide
    state_path=tmp_path/'state.json'
    state_path.write_text(json.dumps(state),encoding='utf-8')
    gates={key:'PASS' for key in runtime.GATE_NAMES}
    for slide_id in state['slides']:
        page=tmp_path/(slide_id+'.pptx'); page.write_bytes(slide_id.encode())
        runtime.seal_validated_single_page(state_path,slide_id,page,gates)
    monkeypatch.setattr(assembly,'require_stage2_entry',lambda _: 'synthetic-approval')
    monkeypatch.setattr(assembly,'stage2_handoff_status',lambda _: {'status':'COMPLETE'})
    monkeypatch.setattr(assembly.preflight,'check',lambda *_args,**_kwargs: {'blockers':[]})
    def fake_merge(inputs, order, output, state_path):
        output.parent.mkdir(parents=True,exist_ok=True)
        output.write_bytes('|'.join(order).encode())
        runtime.seal_merged_deck(state_path,order,inputs,output)
        return {'status':'PASS','path':str(output),'deck_order':order}
    monkeypatch.setattr(assembly,'merge',fake_merge)
    return assembly, runtime, state_path


def test_assembly_defaults_to_trusted_page_order_then_requires_review(tmp_path, monkeypatch):
    import json, pytest
    assembly,runtime,state_path=_assembly_fixture(tmp_path,monkeypatch)
    result=assembly.prepare(state_path,tmp_path/'work')
    assert result['status']=='AWAITING_DECK_VALIDATION'
    assert runtime.load_runtime_state(state_path)['deck_order']==['S001','S002']
    review=Path(result['validation_report'])
    with pytest.raises(ValueError,match='incomplete'):
        assembly.publish(state_path,review,tmp_path/'outputs'/'deck.pptx')
    assert not (tmp_path/'outputs'/'deck.pptx').exists()
    data=json.loads(review.read_text())
    data['checks']={k:'PASS' for k in assembly.DECK_CHECKS}
    data['evidence']={k:'synthetic local review fixture' for k in assembly.DECK_CHECKS}
    review.write_text(json.dumps(data))
    final=tmp_path/'outputs'/'deck.pptx'
    assert assembly.publish(state_path,review,final)['status']=='PASS'
    assert final.read_bytes()==b'S001|S002'
    assert runtime.validate_current_merged_deck(runtime.load_runtime_state(state_path))==[]
    assert runtime.load_runtime_state(state_path)['validated_deck']['status']=='PASS'
    assert (tmp_path/'S001.pptx').read_bytes()==b'S001'
    with pytest.raises(ValueError,match='no overwrite'):
        assembly.publish(state_path,review,final)


def test_assembly_preserves_explicit_order_and_blocks_stale_pages(tmp_path, monkeypatch):
    import pytest
    assembly,runtime,state_path=_assembly_fixture(tmp_path,monkeypatch)
    state=runtime.load_runtime_state(state_path); state['deck_order']=['S002','S001']
    runtime._save_runtime_state(state_path,state)
    result=assembly.prepare(state_path,tmp_path/'work')
    assert Path(result['path']).read_bytes()==b'S002|S001'
    (tmp_path/'S001.pptx').write_bytes(b'changed')
    with pytest.raises(ValueError,match='single_page_artifact_changed'):
        assembly.publish(state_path,result['validation_report'],tmp_path/'out.pptx')


def test_assembly_blocks_missing_duplicate_unknown_and_empty_orders(tmp_path, monkeypatch):
    import pytest
    assembly,runtime,state_path=_assembly_fixture(tmp_path,monkeypatch)
    state=runtime.load_runtime_state(state_path)
    for order in ([],['S001'],['S001','S001'],['S001','other']):
        state['deck_order']=order
        with pytest.raises(ValueError): assembly.current_inputs(state)
    state['deck_order']=['S001','S002']
    state['slides']['S002'].pop('validated_single_page')
    with pytest.raises(ValueError,match='unsealed'):
        assembly.current_inputs(state)


def test_assembly_publish_rejects_wrong_review_hash(tmp_path, monkeypatch):
    import json, pytest
    assembly,runtime,state_path=_assembly_fixture(tmp_path,monkeypatch)
    result=assembly.prepare(state_path,tmp_path/'work')
    review=Path(result['validation_report']); data=json.loads(review.read_text())
    data['merged_pptx_sha256']='wrong'
    review.write_text(json.dumps(data))
    with pytest.raises(ValueError,match='not bound'):
        assembly.publish(state_path,review,tmp_path/'outputs'/'deck.pptx')


def test_download_finalizer_one_page_hash_and_refuses_overwrite(tmp_path):
    import json, shutil, subprocess, pytest
    from pptx import Presentation
    shell=shutil.which('pwsh') or shutil.which('powershell')
    if not shell: pytest.skip('PowerShell not available')
    script=Path(__file__).resolve().parents[1]/'scripts'/'canva_pptx_finalize.ps1'
    src=tmp_path/'download.pptx'; dest=tmp_path/'graphics-first'
    prs=Presentation(); prs.slides.add_slide(prs.slide_layouts[6]); prs.save(src)
    import hashlib
    expected=hashlib.sha256(src.read_bytes()).hexdigest()
    args=[shell,'-NoProfile','-File',str(script),'-DownloadedFile',str(src),'-OutputDirectory',str(dest),'-TargetFileName','S001.pptx','-StableSamples','1','-PollIntervalMilliseconds','100']
    result=subprocess.run(args,capture_output=True,text=True)
    assert result.returncode==0,result.stderr
    assert hashlib.sha256((dest/'S001.pptx').read_bytes()).hexdigest()==expected
    assert not src.exists()
    prs.save(src)
    assert subprocess.run(args,capture_output=True).returncode!=0
    assert src.exists()
    prs.slides.add_slide(prs.slide_layouts[6]); prs.save(src)
    args[args.index('S001.pptx')]='S002.pptx'
    assert subprocess.run(args,capture_output=True).returncode!=0
    assert not (dest/'S002.pptx').exists()


def test_assembly_publish_requires_outputs_and_rechecks_copied_bytes(tmp_path, monkeypatch):
    import json, pytest
    assembly,runtime,state_path=_assembly_fixture(tmp_path,monkeypatch)
    result=assembly.prepare(state_path,tmp_path/'work')
    review=Path(result['validation_report']); data=json.loads(review.read_text())
    data['checks']={k:'PASS' for k in assembly.DECK_CHECKS}
    data['evidence']={k:'synthetic local review fixture' for k in assembly.DECK_CHECKS}
    review.write_text(json.dumps(data))
    before=state_path.read_bytes()
    with pytest.raises(ValueError,match='under outputs'):
        assembly.publish(state_path,review,tmp_path/'work'/'delivery.pptx')
    def changed_copy(src,dest):
        dest.write(b'changed during copy')
    monkeypatch.setattr(assembly.shutil,'copyfileobj',changed_copy)
    final=tmp_path/'outputs'/'deck.pptx'
    with pytest.raises(ValueError,match='reviewed bytes'):
        assembly.publish(state_path,review,final)
    assert not final.exists()
    assert state_path.read_bytes()==before
    assert runtime.validate_current_merged_deck(runtime.load_runtime_state(state_path))==[]


# rc7 user flow: real internal approval/registration/merge diagnostics; external edges only simulated.
def _download_and_review(state_path, sid, tmp_path, seed):
    import json, canva_bridge as bridge, runtime
    state=runtime.load_runtime_state(state_path)
    identity=bridge.active_identity(state,sid)
    graphics=tmp_path/(identity['attempt_id']+'-graphics.pptx')
    graphics.write_bytes(seed.read_bytes())
    evidence={**identity,'pptx_sha256':bridge.sha(graphics),'edit_url':'https://www.canva.com/design/'+identity['design_id']+'/edit',
              'download_entry':'synthetic browser entry','completion_evidence':'synthetic complete event'}
    evidence_path=tmp_path/(identity['attempt_id']+'-download.json');evidence_path.write_text(json.dumps(evidence))
    bridge.bind_download(state_path,sid,graphics,evidence_path)
    restored=tmp_path/(identity['attempt_id']+'-restored.pptx');restored.write_bytes(seed.read_bytes())
    state=runtime.load_runtime_state(state_path)
    review={**identity,'download_sha256':bridge.sha(graphics),'restored_pptx_sha256':bridge.sha(restored),
            'text_restore_fingerprint':runtime.expected_lineage_from_state(state,sid)['text_restore_fingerprint'],
            'checks':{k:'PASS' for k in runtime.GATE_NAMES},'evidence':{k:'synthetic gate observation' for k in runtime.GATE_NAMES}}
    review_path=tmp_path/(identity['attempt_id']+'-review.json');review_path.write_text(json.dumps(review))
    return graphics,restored,review_path


def _ready_codex_deck(tmp_path):
    from test_preflight_scope import _bridge_fixture,_accepted_attempt
    from pptx import Presentation
    import canva_bridge as bridge
    state_path=_bridge_fixture(tmp_path)
    seed=tmp_path/'seed.pptx';prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(seed)
    pages={}
    for sid in ('S001','S002','S003'):
        _accepted_attempt(state_path,sid,tmp_path,design='design-A-'+sid)
        graphics,restored,review=_download_and_review(state_path,sid,tmp_path,seed)
        bridge.seal_page(state_path,sid,restored,review)
        pages[sid]=(graphics,restored,review)
    return state_path,seed,pages


def _local_native_only(monkeypatch):
    import preflight,merge_pptx
    from pptx import Presentation
    original=preflight.importlib.util.find_spec
    monkeypatch.setattr(preflight.importlib.util,'find_spec',lambda name: object() if name=='win32com' else original(name))
    monkeypatch.setattr(preflight,'_probe_powerpoint',lambda _: (True,'synthetic environment only'))
    def native(inputs,order,path):
        prs=Presentation()
        for sid in order: prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(path)
    monkeypatch.setattr(merge_pptx,'_native_merge',native)


def _complete_deck_review(path):
    import json,assemble_deck
    data=json.loads(path.read_text())
    data['checks']={k:'PASS' for k in assemble_deck.DECK_CHECKS}
    data['evidence']={k:'synthetic deck observation' for k in assemble_deck.DECK_CHECKS}
    path.write_text(json.dumps(data))


def test_rc7_three_page_real_internal_flow_and_readonly_resume(tmp_path,monkeypatch):
    import canva_bridge as bridge,assemble_deck, runtime
    state_path,seed,pages=_ready_codex_deck(tmp_path)
    before=state_path.read_bytes()
    assert 'deck_order' not in runtime.load_runtime_state(state_path)
    assert bridge.status(state_path)['status']=='PREPARE_DECK'
    assert state_path.read_bytes()==before
    _local_native_only(monkeypatch)
    result=assemble_deck.prepare(state_path,tmp_path/'work')
    assert result['status']=='AWAITING_DECK_VALIDATION'
    assert bridge.status(state_path)['status']=='DECK_VALIDATION'
    assert runtime.load_runtime_state(state_path)['deck_order']==['S001','S002','S003']
    review=Path(result['validation_report']);_complete_deck_review(review)
    output=tmp_path/'outputs'/'deck.pptx'
    assert assemble_deck.publish(state_path,review,output)['status']=='PASS'
    assert bridge.status(state_path)['status']=='COMPLETE'
    assert all(graphics.exists() and restored.exists() for graphics,restored,_ in pages.values())


def test_rc7_design_switch_blocks_old_page_report_even_with_identical_bytes(tmp_path,monkeypatch):
    import pytest,json,canva_bridge as bridge,assemble_deck,runtime
    from test_preflight_scope import _accepted_attempt
    state_path,seed,pages=_ready_codex_deck(tmp_path)
    _local_native_only(monkeypatch)
    result=assemble_deck.prepare(state_path,tmp_path/'work-A')
    old_review=Path(result['validation_report']);_complete_deck_review(old_review)
    old_bytes=Path(result['path']).read_bytes()
    old_single=runtime.load_runtime_state(state_path)['slides']['S001']['validated_single_page']
    identity=_accepted_attempt(state_path,'S001',tmp_path,design='design-B',rebuild=True)
    assert bridge.status(state_path)['status']=='WAIT_DOWNLOAD'
    with pytest.raises(ValueError): assemble_deck.current_inputs(runtime.load_runtime_state(state_path))
    with pytest.raises(ValueError): assemble_deck.publish(state_path,old_review,tmp_path/'outputs'/'old.pptx')
    graphics,restored,new_review=_download_and_review(state_path,'S001',tmp_path,seed)
    assert restored.read_bytes()==pages['S001'][1].read_bytes()
    with pytest.raises(ValueError,match='identity'):
        bridge.seal_page(state_path,'S001',restored,pages['S001'][2])
    assert runtime.load_runtime_state(state_path)['slides']['S001']['validated_single_page']==old_single
    bridge.seal_page(state_path,'S001',restored,new_review)
    assert bridge.status(state_path)['status']=='PREPARE_DECK'
    assert runtime.load_runtime_state(state_path)['slides']['S001']['superseded_validated_single_pages'][-1]==old_single
    result=assemble_deck.prepare(state_path,tmp_path/'work-B')
    # Force byte equality of the merged fixtures while keeping real seal checks.
    Path(result['path']).write_bytes(old_bytes)
    state=runtime.load_runtime_state(state_path)
    runtime.seal_merged_deck(state_path,state['deck_order'],assemble_deck.current_inputs(state)[1],result['path'])
    state=runtime.load_runtime_state(state_path)
    state['canva_bridge']['merged']['merged_pptx_sha256']=bridge.sha(result['path']);runtime._save_runtime_state(state_path,state)
    new_report=json.loads(Path(result['validation_report']).read_text());new_report['merged_pptx_sha256']=bridge.sha(result['path'])
    Path(result['validation_report']).write_text(json.dumps(new_report))
    with pytest.raises(ValueError,match='old design'):
        assemble_deck.publish(state_path,old_review,tmp_path/'outputs'/'wrong.pptx')
    _complete_deck_review(Path(result['validation_report']))
    assert assemble_deck.publish(state_path,result['validation_report'],tmp_path/'outputs'/'new.pptx')['status']=='PASS'
    assert pages['S001'][0].exists() and pages['S001'][1].exists()


def test_rc7_pending_missing_fields_and_changed_download_stop_without_retry(tmp_path):
    import pytest,json,canva_bridge as bridge,runtime
    from test_preflight_scope import _bridge_fixture
    from pptx import Presentation
    state_path=_bridge_fixture(tmp_path,n=1)
    assert bridge.status(state_path)['status']=='WAIT_RECONSTRUCTION'
    active=bridge.begin_attempt(state_path,'S001')
    assert bridge.status(state_path)['status']=='WAIT_RECONSTRUCTION'
    before=state_path.read_bytes()
    with pytest.raises(ValueError): bridge.begin_attempt(state_path,'S001',True)
    assert state_path.read_bytes()==before
    result={**active,'design_id':'design-A','assessment':'ACCEPT','evidence':'synthetic observed result'}
    result_path=tmp_path/'result.json';result_path.write_text(json.dumps(result))
    bridge.accept_attempt(state_path,'S001',active['attempt_id'],result_path)
    assert bridge.status(state_path)['status']=='WAIT_DOWNLOAD'
    seed=tmp_path/'seed.pptx';prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(seed)
    graphics,restored,review=_download_and_review(state_path,'S001',tmp_path,seed)
    assert bridge.status(state_path)['status']=='RESTORE_TEXT'
    state=runtime.load_runtime_state(state_path)
    state['stage2_run']['artifacts'].pop('finalized_manifest:S001');runtime._save_runtime_state(state_path,state)
    before=state_path.read_bytes()
    with pytest.raises(ValueError,match='finalized_manifest'): bridge.seal_page(state_path,'S001',restored,review)
    assert state_path.read_bytes()==before
    assert 'validated_single_page' not in runtime.load_runtime_state(state_path)['slides']['S001']
    graphics.write_bytes(b'changed downloaded bytes')
    # The missing validated text plan is now the earliest unresolved state; later download corruption remains hidden until text prep is repaired.
    assert bridge.status(state_path)['status']=='PREPARE_TEXT'


def test_rc7_failed_seal_and_publication_commit_do_not_restore_old_design(tmp_path,monkeypatch):
    import pytest,canva_bridge as bridge,runtime,assemble_deck
    from test_preflight_scope import _accepted_attempt
    state_path,seed,pages=_ready_codex_deck(tmp_path)
    _accepted_attempt(state_path,'S001',tmp_path,design='design-B',rebuild=True)
    graphics,restored,review=_download_and_review(state_path,'S001',tmp_path,seed)
    before=state_path.read_bytes()
    original=bridge.os.replace
    def fail_commit(src,dest):
        if Path(dest).resolve()==state_path.resolve(): raise OSError('synthetic commit failure')
        return original(src,dest)
    with monkeypatch.context() as patch:
        patch.setattr(bridge.os,'replace',fail_commit)
        with pytest.raises(OSError): bridge.seal_page(state_path,'S001',restored,review)
    assert state_path.read_bytes()==before
    assert bridge.active_identity(runtime.load_runtime_state(state_path),'S001')['design_id']=='design-B'
    with pytest.raises(ValueError): assemble_deck.current_inputs(runtime.load_runtime_state(state_path))
    bridge.seal_page(state_path,'S001',restored,review)
    _local_native_only(monkeypatch)
    result=assemble_deck.prepare(state_path,tmp_path/'work');deck_review=Path(result['validation_report']);_complete_deck_review(deck_review)
    before=state_path.read_bytes();output=tmp_path/'outputs'/'deck.pptx'
    with monkeypatch.context() as patch:
        patch.setattr(bridge.os,'replace',fail_commit)
        with pytest.raises(OSError): assemble_deck.publish(state_path,deck_review,output)
    assert not output.exists() and state_path.read_bytes()==before
    assert bridge.status(state_path)['status']=='DECK_VALIDATION'


def test_rc7_transaction_detects_other_writer_without_overwriting_it(tmp_path):
    import pytest,json,canva_bridge as bridge
    state=tmp_path/'state.json';bridge.init(state)
    with pytest.raises(ValueError,match='state changed'):
        with bridge.transaction(state) as temp:
            data=json.loads(state.read_text());data['other_writer']='preserve me';state.write_text(json.dumps(data))
    assert json.loads(state.read_text())['other_writer']=='preserve me'


def test_rc7_missing_binding_cannot_be_relabeled_external(tmp_path):
    import pytest,canva_bridge as bridge,runtime,assemble_deck
    state_path,seed,pages=_ready_codex_deck(tmp_path)
    state=runtime.load_runtime_state(state_path);state['canva_bridge']['route']='external';runtime._save_runtime_state(state_path,state)
    with pytest.raises(ValueError,match='route'): assemble_deck.current_inputs(runtime.load_runtime_state(state_path))


def test_rc7_accept_commit_failure_resumes_actual_result_without_new_call(tmp_path,monkeypatch):
    import pytest,json,canva_bridge as bridge,runtime
    from test_preflight_scope import _bridge_fixture
    state=_bridge_fixture(tmp_path,n=1);active=bridge.begin_attempt(state,'S001')
    result={**active,'design_id':'design-A','assessment':'ACCEPT','evidence':'synthetic observed actual result'}
    result_path=tmp_path/'result.json';result_path.write_text(json.dumps(result))
    before=state.read_bytes();original=bridge.os.replace
    def fail(src,dest):
        if Path(dest).resolve()==state.resolve(): raise OSError('synthetic final commit failure')
        return original(src,dest)
    with monkeypatch.context() as patch:
        patch.setattr(bridge.os,'replace',fail)
        with pytest.raises(OSError):bridge.accept_attempt(state,'S001',active['attempt_id'],result_path)
    assert state.read_bytes()==before
    assert bridge.status(state)['status']=='WAIT_RECONSTRUCTION'
    bridge.accept_attempt(state,'S001',active['attempt_id'],result_path)
    assert bridge.status(state)['status']=='WAIT_DOWNLOAD'
    before=state.read_bytes()
    with pytest.raises(ValueError):bridge.accept_attempt(state,'S001',active['attempt_id'],result_path)
    assert state.read_bytes()==before


def test_rc7_failed_native_merge_and_interrupted_lock_do_not_claim_completion(tmp_path,monkeypatch):
    import pytest,canva_bridge as bridge,assemble_deck,merge_pptx
    state,seed,pages=_ready_codex_deck(tmp_path);_local_native_only(monkeypatch)
    before=state.read_bytes()
    def failure(*args): raise OSError('synthetic native failure')
    monkeypatch.setattr(merge_pptx,'_native_merge',failure)
    with pytest.raises(ValueError,match='native merge failed'):
        assemble_deck.prepare(state,tmp_path/'work')
    assert state.read_bytes()==before
    assert not (tmp_path/'work'/'deck-candidate.pptx').exists()
    assert bridge.status(state)['status']=='PREPARE_DECK'
    lock=state.with_name(state.name+'.bridge-lock');lock.write_text('interrupted fixture')
    assert bridge.status(state)['status']=='BLOCKED'
    with pytest.raises(FileExistsError):bridge.begin_attempt(state,'S001',True)
    assert lock.exists() and state.read_bytes()==before


def test_rc7_structural_retry_keeps_original_one_retry_limit(tmp_path):
    import pytest,json,canva_bridge as bridge
    from test_preflight_scope import _bridge_fixture
    state=_bridge_fixture(tmp_path,n=1)
    first=bridge.begin_attempt(state,'S001')
    result={**first,'design_id':'design-A','assessment':'RETRY_ONCE','evidence':'synthetic non-text structural anomaly'}
    result_path=tmp_path/'result.json';result_path.write_text(json.dumps(result))
    assert bridge.accept_attempt(state,'S001',first['attempt_id'],result_path)['status']=='RETRY_ONCE'
    assert bridge.status(state)['status']=='RETRY_ONCE'
    second=bridge.begin_attempt(state,'S001',True)
    result={**second,'design_id':'design-B','assessment':'RETRY_ONCE','evidence':'synthetic second structural anomaly'}
    result_path.write_text(json.dumps(result))
    with pytest.raises(ValueError,match='consumed'):
        bridge.accept_attempt(state,'S001',second['attempt_id'],result_path)
    result['assessment']='ASSISTED_CLEANUP_REQUIRED';result_path.write_text(json.dumps(result))
    bridge.accept_attempt(state,'S001',second['attempt_id'],result_path)
    assert bridge.status(state)['status']=='ASSISTED_CLEANUP_REQUIRED'
    with pytest.raises(ValueError,match='UNRESOLVED'):
        bridge.accept_attempt(state,'S001',second['attempt_id'],result_path)
    result['assessment']='ACCEPT';result['evidence']='synthetic observed one-time cleanup result';result_path.write_text(json.dumps(result))
    bridge.accept_attempt(state,'S001',second['attempt_id'],result_path)
    assert bridge.status(state)['status']=='WAIT_DOWNLOAD'


def test_rc7_download_identity_and_stale_restoration_review_are_rejected(tmp_path):
    import pytest,json,canva_bridge as bridge,runtime
    state,seed,pages=_ready_codex_deck(tmp_path)
    graphics,restored,review=pages['S001']
    identity=bridge.active_identity(runtime.load_runtime_state(state),'S001')
    evidence={**identity,'design_id':'wrong-design','pptx_sha256':bridge.sha(graphics),
              'edit_url':'synthetic observed URL','download_entry':'synthetic entry','completion_evidence':'synthetic completion'}
    evidence_path=tmp_path/'bad-download.json';evidence_path.write_text(json.dumps(evidence))
    before=state.read_bytes()
    with pytest.raises(ValueError,match='identity'):bridge.bind_download(state,'S001',graphics,evidence_path)
    assert state.read_bytes()==before
    fallback=tmp_path/'S001-font_fallback.json';fallback.write_text('{"missing-font":"actual-fallback"}')
    runtime.register_stage2_artifact(state,'font_fallback:S001',fallback)
    before=state.read_bytes()
    with pytest.raises(ValueError,match='restoration inputs'):bridge.seal_page(state,'S001',restored,review)
    assert state.read_bytes()==before
    data=json.loads(review.read_text());data['checks']['content_truth']='FAIL';review.write_text(json.dumps(data))
    with pytest.raises(ValueError,match='four Hard Gates'):bridge.seal_page(state,'S001',restored,review)
    assert state.read_bytes()==before


def test_rc7_deck_requires_background_fixed_elements_and_preserves_order_checks(tmp_path,monkeypatch):
    import pytest,json,canva_bridge as bridge,assemble_deck,runtime
    state,seed,pages=_ready_codex_deck(tmp_path);_local_native_only(monkeypatch)
    result=assemble_deck.prepare(state,tmp_path/'work');review=Path(result['validation_report']);_complete_deck_review(review)
    data=json.loads(review.read_text());data['checks'].pop('background_family');review.write_text(json.dumps(data))
    with pytest.raises(ValueError,match='background_family'):assemble_deck.publish(state,review,tmp_path/'outputs'/'bad.pptx')
    _complete_deck_review(review);data=json.loads(review.read_text());data['evidence']['page_numbers_fixed_elements']='';review.write_text(json.dumps(data))
    with pytest.raises(ValueError,match='page_numbers_fixed_elements'):assemble_deck.publish(state,review,tmp_path/'outputs'/'bad.pptx')
    snapshot=runtime.load_runtime_state(state)
    for order in ([],['S001'],['S001','S001','S003'],['S001','S002','other']):
        changed=dict(snapshot);changed['deck_order']=order
        with pytest.raises(ValueError):assemble_deck.current_inputs(changed)
    assert not (tmp_path/'outputs'/'bad.pptx').exists()


# rc8: regressions use real internal state; only external/native edges are simulated.
def test_rc8_confirmed_tool_failure_can_resume_without_consuming_structural_retry(tmp_path):
    import json,pytest,canva_bridge as bridge
    from test_preflight_scope import _bridge_fixture
    state=_bridge_fixture(tmp_path,n=1);active=bridge.begin_attempt(state,'S001')
    result={**active,'assessment':'MAGIC_LAYERS_EXECUTION_FAILED','terminal_confirmed':True,
            'evidence':'synthetic observed terminal failure, no result'}
    path=tmp_path/'failure.json';path.write_text(json.dumps(result))
    outcome=bridge.accept_attempt(state,'S001',active['attempt_id'],path)
    assert outcome['status']=='TOOL_FAILED'
    assert bridge.status(state)['status']=='TOOL_FAILED'
    before=state.read_bytes()
    with pytest.raises(ValueError):bridge.begin_attempt(state,'S001',True)
    assert state.read_bytes()==before
    recovery={'attempt_id':active['attempt_id'],'failure_code':'MAGIC_LAYERS_EXECUTION_FAILED',
              'change_evidence':'synthetic observed connection repair','terminal_confirmed':True}
    evidence=tmp_path/'recovery.json';evidence.write_text(json.dumps(recovery))
    resumed=bridge.begin_attempt(state,'S001',True,recovery_evidence=evidence)
    assert resumed['attempt_id']!=active['attempt_id']
    assert resumed['retry_state']=='available'
    assert bridge.status(state)['status']=='WAIT_RECONSTRUCTION'


def test_rc8_page_count_change_archives_previous_run_without_resetting_resume(tmp_path):
    import json,canva_bridge as bridge,runtime
    from test_preflight_scope import _accepted_attempt
    from PIL import Image
    state,seed,pages=_ready_codex_deck(tmp_path)
    original=state.read_bytes();old_run=runtime.load_runtime_state(state)['stage2_run']['run_id']
    same=bridge.start_run(state)
    assert same['run_id']==old_run and state.read_bytes()==original
    for count in (2,4):
        prior=runtime.load_runtime_state(state)
        outline=tmp_path/('outline-'+str(count)+'.json')
        outline.write_text(json.dumps({'slides':[{k:k+' replacement '+str(i) for k in runtime.OUTLINE_FIELDS} for i in range(count)]}),encoding='utf-8')
        runtime.record_stage1_draft(state,outline);runtime.present_stage1_review(state);runtime.handle_stage1_reply(state,'同意，按这个继续。', decision='approve')
        run=bridge.start_run(state)
        current=runtime.load_runtime_state(state)
        assert list(current['slides'])==['S'+str(i).zfill(3) for i in range(1,count+1)]
        assert 'merged_deck' not in current and 'deck_order' not in current
        assert current['canva_bridge']['run_history'][-1]['slides']==prior['slides']
        for sid in run['page_order']:
            p=tmp_path/(sid+'-new.png');Image.new('RGB',(16,9),'blue').save(p)
            candidate=runtime.register_stage2_candidate(state,sid,p,'new approved outline fixture')
            runtime.review_stage2_candidate(state,candidate['candidate_id'],'PASS')
            for kind,value in [('final_content_truth',{'slide_id':sid,'meaningful_text':['new'],'exact_facts':[]}),('canvas',{'width':1600,'height':900})]:
                p=tmp_path/(sid+'-new-'+kind+'.json');p.write_text(json.dumps(value));runtime.register_stage2_artifact(state,kind+':'+sid,p)
            runtime.mark_stage2_text_reconciled(state,sid)
        runtime.present_stage2_formal_display(state)
        runtime.seal_stage2_visual_approval(state,'确认，按这组页面进入下一阶段。')
        for sid in run['page_order']:
            from stage2_test_helpers import write_stage3_text_plan
            manifest,inventory=write_stage3_text_plan(tmp_path/(sid+'-new-manifest.json'),tmp_path/(sid+'-new-inventory.json'),sid,'new')
            bridge.register_text_plan(state,sid,manifest,inventory)
            p=tmp_path/(sid+'-new-font_fallback.json');p.write_text('{}');runtime.register_stage2_artifact(state,'font_fallback:'+sid,p)
            p=tmp_path/(sid+'-clean-new.png');Image.new('RGB',(16,9),'blue').save(p);bridge.register_text_clean(state,sid,p)
        for sid in run['page_order']:
            _accepted_attempt(state,sid,tmp_path,design='design-new-'+str(count)+'-'+sid)
            graphics,restored,review=_download_and_review(state,sid,tmp_path,seed)
            bridge.seal_page(state,sid,restored,review)
        assert bridge.status(state)['status']=='PREPARE_DECK'
        assert all(source.exists() and restored.exists() for source,restored,_ in pages.values())


def test_rc8_review_write_failure_leaves_no_registered_candidate(tmp_path,monkeypatch):
    import pytest,canva_bridge as bridge,assemble_deck
    state,seed,pages=_ready_codex_deck(tmp_path);_local_native_only(monkeypatch)
    original=Path.write_text
    def fail(path,*args,**kwargs):
        if path.name=='deck-review.json':raise OSError('synthetic review write failure')
        return original(path,*args,**kwargs)
    before=state.read_bytes()
    with monkeypatch.context() as patch:
        patch.setattr(Path,'write_text',fail)
        with pytest.raises(OSError):assemble_deck.prepare(state,tmp_path/'work')
    assert state.read_bytes()==before
    assert not (tmp_path/'work/deck-candidate.pptx').exists()
    assert not (tmp_path/'work/deck-review.json').exists()
    assert bridge.status(state)['status']=='PREPARE_DECK'
    assert assemble_deck.prepare(state,tmp_path/'work')['status']=='AWAITING_DECK_VALIDATION'


def test_rc8_missing_review_recovers_existing_candidate_without_com_or_overwrite(tmp_path,monkeypatch):
    import pytest,assemble_deck,preflight,merge_pptx
    state,seed,pages=_ready_codex_deck(tmp_path);_local_native_only(monkeypatch)
    result=assemble_deck.prepare(state,tmp_path/'work');review=Path(result['validation_report']);review.unlink()
    before=state.read_bytes();candidate=Path(result['path']);saved=candidate.read_bytes()
    def forbidden(*args,**kwargs):raise AssertionError('recovery must not invoke COM or merge')
    monkeypatch.setattr(preflight,'_probe_powerpoint',forbidden);monkeypatch.setattr(merge_pptx,'_native_merge',forbidden)
    recovered=assemble_deck.recover_review(state)
    assert recovered['status']=='AWAITING_DECK_VALIDATION'
    assert review.exists() and candidate.read_bytes()==saved and state.read_bytes()==before
    _complete_deck_review(review);filled=review.read_bytes()
    with pytest.raises(ValueError):assemble_deck.recover_review(state)
    assert review.read_bytes()==filled
    assert assemble_deck.publish(state,review,tmp_path/'outputs/deck.pptx')['status']=='PASS'


def test_rc8_unknown_failure_and_mismatched_repair_never_authorize_new_call(tmp_path):
    import json,pytest,canva_bridge as bridge
    from test_preflight_scope import _bridge_fixture
    state=_bridge_fixture(tmp_path,n=1);active=bridge.begin_attempt(state,'S001')
    result={**active,'assessment':'MAGIC_LAYERS_EXECUTION_FAILED','evidence':'synthetic timeout, terminal status not known'}
    path=tmp_path/'result.json';path.write_text(json.dumps(result));before=state.read_bytes()
    with pytest.raises(ValueError):bridge.accept_attempt(state,'S001',active['attempt_id'],path)
    assert state.read_bytes()==before
    with pytest.raises(ValueError):bridge.begin_attempt(state,'S001',True)
    result['terminal_confirmed']=True;result['evidence']='synthetic confirmed terminal failure';path.write_text(json.dumps(result))
    bridge.accept_attempt(state,'S001',active['attempt_id'],path)
    recovery={'attempt_id':'wrong-attempt','failure_code':'MAGIC_LAYERS_EXECUTION_FAILED',
              'terminal_confirmed':True,'change_evidence':'synthetic observed environment change'}
    repair=tmp_path/'repair.json';repair.write_text(json.dumps(recovery));before=state.read_bytes()
    with pytest.raises(ValueError):bridge.begin_attempt(state,'S001',True,recovery_evidence=repair)
    assert state.read_bytes()==before and bridge.status(state)['status']=='TOOL_FAILED'


def test_rc8_recovery_rejects_changed_candidate_without_remerge(tmp_path,monkeypatch):
    import pytest,assemble_deck
    state,seed,pages=_ready_codex_deck(tmp_path);_local_native_only(monkeypatch)
    result=assemble_deck.prepare(state,tmp_path/'work');review=Path(result['validation_report']);review.unlink()
    candidate=Path(result['path']);candidate.write_bytes(candidate.read_bytes()+b'changed')
    before=state.read_bytes()
    with pytest.raises(ValueError):assemble_deck.recover_review(state)
    assert not review.exists() and state.read_bytes()==before


def test_rc8_prepare_commit_failure_removes_new_template_and_preserves_state(tmp_path,monkeypatch):
    import pytest,canva_bridge as bridge,assemble_deck
    state,seed,pages=_ready_codex_deck(tmp_path);_local_native_only(monkeypatch)
    before=state.read_bytes();original=bridge.os.replace
    def fail(src,dest):
        if Path(dest).resolve()==state.resolve():raise OSError('synthetic final prepare commit failure')
        return original(src,dest)
    monkeypatch.setattr(bridge.os,'replace',fail)
    with pytest.raises(OSError):assemble_deck.prepare(state,tmp_path/'work')
    assert state.read_bytes()==before
    assert not (tmp_path/'work/deck-review.json').exists()
    assert not (tmp_path/'work/deck-candidate.pptx').exists()


# rc9: test public CLI output and source/structure checks, not only internal returns.
def _rc9_cli(script, *args):
    import subprocess,json
    root=Path(__file__).resolve().parents[1]
    cp=subprocess.run([sys.executable,'-X','utf8',str(root/'scripts'/script),*map(str,args)],capture_output=True,text=True,encoding='utf-8',cwd=root)
    assert 'Traceback' not in cp.stderr,cp.stderr
    return cp,json.loads(cp.stdout)


def test_rc9_cli_first_start_and_same_run_resume(tmp_path):
    import json,runtime
    state=tmp_path/'state.json'
    assert _rc9_cli('canva_bridge.py','init','--state',state)[0].returncode==0
    outline=tmp_path/'outline.json'
    outline.write_text(json.dumps({'slides':[{k:k+str(i) for k in runtime.OUTLINE_FIELDS} for i in range(3)]}))
    assert _rc9_cli('runtime.py','stage1-draft','--state',state,'--outline',outline)[0].returncode==0
    before=state.read_bytes()
    cp,blocked=_rc9_cli('canva_bridge.py','start-run','--state',state)
    assert cp.returncode==1 and blocked['status']=='BLOCKED' and state.read_bytes()==before
    assert _rc9_cli('runtime.py','stage1-review','--state',state)[1]['review_view']
    assert _rc9_cli('runtime.py','stage1-reply','--state',state,'--message','同意，按这个继续。','--decision','approve')[0].returncode==0
    assert _rc9_cli('runtime.py','stage2-entry','--state',state)[0].returncode==0
    cp,run=_rc9_cli('canva_bridge.py','start-run','--state',state)
    assert cp.returncode==0 and cp.stderr=='' and run['status']=='RUN_READY'
    assert run['page_order']==['S001','S002','S003']
    before=state.read_bytes()
    cp,resumed=_rc9_cli('canva_bridge.py','start-run','--state',state)
    assert cp.returncode==0 and resumed['run_id']==run['run_id'] and state.read_bytes()==before


def test_rc9_cli_outline_rollover_commits_once_and_archives(tmp_path):
    import json,runtime
    state,_,pages=_ready_codex_deck(tmp_path)
    previous=runtime.load_runtime_state(state)['stage2_run']['run_id']
    outline=tmp_path/'two.json'
    outline.write_text(json.dumps({'slides':[{k:k+str(i) for k in runtime.OUTLINE_FIELDS} for i in range(2)]}))
    _rc9_cli('runtime.py','stage1-draft','--state',state,'--outline',outline)
    _rc9_cli('runtime.py','stage1-review','--state',state)
    _rc9_cli('runtime.py','stage1-reply','--state',state,'--message','同意，按这个继续。','--decision','approve')
    cp,run=_rc9_cli('canva_bridge.py','start-run','--state',state)
    assert cp.returncode==0 and cp.stderr=='' and run['status']=='RUN_READY'
    current=runtime.load_runtime_state(state)
    assert list(current['slides'])==['S001','S002']
    assert current['canva_bridge']['run_history'][-1]['run_id']==previous
    assert all(p.exists() for triple in pages.values() for p in triple)
    before=state.read_bytes();assert _rc9_cli('canva_bridge.py','start-run','--state',state)[0].returncode==0
    assert state.read_bytes()==before


def test_rc9_download_url_checks_new_and_historical_evidence(tmp_path):
    import json,pytest,canva_bridge as bridge,runtime
    state,_,pages=_ready_codex_deck(tmp_path)
    sid='S001';graphics,_,_=pages[sid]
    identity=bridge.active_identity(runtime.load_runtime_state(state),sid)
    base={'attempt_id':identity['attempt_id'],'slide_id':sid,'design_id':identity['design_id'],
          'pptx_sha256':bridge.sha(graphics),'download_entry':'synthetic observed entry','completion_evidence':'synthetic completion'}
    evidence=tmp_path/'url.json'
    for url in ['https://www.canva.com/design/OTHER/edit','https://www.canva.com.evil.example/design/'+identity['design_id']+'/edit',
                'https://user@www.canva.com/design/'+identity['design_id']+'/edit','http://www.canva.com/design/'+identity['design_id']+'/edit','not a URL','https://www.canva.com/no-design','https://www.canva.com:not-a-port/design/'+identity['design_id']+'/edit','https://www.canva.com:70000/design/'+identity['design_id']+'/edit']:
        evidence.write_text(json.dumps({**base,'edit_url':url}))
        before=state.read_bytes()
        with pytest.raises(ValueError):bridge.bind_download(state,sid,graphics,evidence)
        assert state.read_bytes()==before
    for url in ['https://canva.com/design/'+identity['design_id']+'/edit',
                'https://www.canva.com/design/'+identity['design_id']+'/token/edit?ui=editor#page1']:
        evidence.write_text(json.dumps({**base,'edit_url':url}));bridge.bind_download(state,sid,graphics,evidence)
        assert bridge.download_current(runtime.load_runtime_state(state),sid)['evidence']['edit_url']==url
    for historic_url in ['https://www.canva.com/design/OTHER/edit','https://www.canva.com:not-a-port/design/'+identity['design_id']+'/edit','https://www.canva.com:70000/design/'+identity['design_id']+'/edit']:
        current=runtime.load_runtime_state(state)
        current['canva_bridge']['downloads'][sid]['evidence']['edit_url']=historic_url
        runtime._save_runtime_state(state,current);before=state.read_bytes()
        with pytest.raises(ValueError):bridge.download_current(runtime.load_runtime_state(state),sid)
        assert bridge.status(state)['status']=='BLOCKED' and state.read_bytes()==before


def _rc9_package_cases(tmp_path):
    import zipfile
    from pptx import Presentation
    valid=tmp_path/'valid.pptx';prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(valid)
    multi=tmp_path/'multi.pptx';prs.slides.add_slide(prs.slide_layouts[6]);prs.save(multi)
    cases={'valid':valid,'multi':multi}
    with zipfile.ZipFile(valid) as z: parts={n:z.read(n) for n in z.namelist()}
    edits={'missing-root-rel':{'_rels/.rels':None},'missing-slide-rel':{'ppt/_rels/presentation.xml.rels':None},
           'missing-slide':{'ppt/slides/slide1.xml':None},
           'broken-rel':{'ppt/_rels/presentation.xml.rels':parts['ppt/_rels/presentation.xml.rels'].replace(b'slides/slide1.xml',b'slides/absent.xml')}}
    for name,changes in edits.items():
        path=tmp_path/(name+'.pptx')
        with zipfile.ZipFile(path,'w') as z:
            for key,data in {**parts,**changes}.items():
                if data is not None:z.writestr(key,data)
        cases[name]=path
    minimal=tmp_path/'minimal.pptx'
    with zipfile.ZipFile(minimal,'w') as z:
        z.writestr('[Content_Types].xml','<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        z.writestr('_rels/.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')
        z.writestr('ppt/presentation.xml','<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:sldIdLst><p:sldId id="256"/></p:sldIdLst></p:presentation>')
        z.writestr('ppt/slides/slide1.xml','<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>')
    cases['minimal']=minimal
    corrupt=tmp_path/'corrupt.pptx';corrupt.write_bytes(b'not a zip');cases['corrupt']=corrupt
    return cases


def test_rc9_python_and_powershell_reject_same_invalid_packages(tmp_path):
    import pytest,subprocess,shutil,canva_bridge as bridge
    shell=shutil.which('pwsh') or shutil.which('powershell')
    if not shell: pytest.skip('PowerShell not available')
    root=Path(__file__).resolve().parents[1]
    cases=_rc9_package_cases(tmp_path)
    for name,path in cases.items():
        before=path.read_bytes()
        if name=='valid':assert bridge.single_page(path) is None
        else:
            with pytest.raises(ValueError,match='GRAPHICS_FIRST_PPTX_INVALID'):bridge.single_page(path)
        # Identical bytes in an independent observed download, preserving the source fixture.
        observed=tmp_path/(name+' observed.pptx');observed.write_bytes(before)
        dest=tmp_path/('final '+name)
        cp=subprocess.run([shell,'-NoProfile','-File',str(root/'scripts/canva_pptx_finalize.ps1'),'-DownloadedFile',str(observed),
                           '-OutputDirectory',str(dest),'-TargetFileName','S001.pptx','-StableSamples','1','-PythonExecutable',sys.executable],capture_output=True,text=True)
        if name=='valid':assert cp.returncode==0,cp.stderr;assert (dest/'S001.pptx').read_bytes()==before
        else:assert cp.returncode!=0;assert observed.read_bytes()==before and not (dest/'S001.pptx').exists()


def test_rc9_validator_dependency_and_interpreter_failure_are_distinct(tmp_path,monkeypatch):
    import builtins,pytest,subprocess,shutil,canva_bridge as bridge
    valid=_rc9_package_cases(tmp_path)['valid'];original=builtins.__import__
    def missing(name,*args,**kwargs):
        if name=='pptx':raise ModuleNotFoundError('offline missing pptx',name='pptx')
        return original(name,*args,**kwargs)
    with monkeypatch.context() as m:
        m.setattr(builtins,'__import__',missing)
        with pytest.raises(ValueError,match='PPTX_VALIDATION_DEPENDENCY_MISSING'):bridge.single_page(valid)
    shell=shutil.which('pwsh') or shutil.which('powershell')
    if not shell: pytest.skip('PowerShell not available')
    script=Path(__file__).resolve().parents[1]/'scripts/canva_pptx_finalize.ps1'
    cp=subprocess.run([shell,'-NoProfile','-File',str(script),'-DownloadedFile',str(valid),'-OutputDirectory',str(tmp_path/'out'),
                       '-TargetFileName','S001.pptx','-StableSamples','1','-PythonExecutable',str(tmp_path/'missing python.exe')],capture_output=True,text=True)
    assert cp.returncode!=0 and 'PPTX_VALIDATION_INTERPRETER_UNAVAILABLE' in cp.stderr
    assert valid.exists() and not (tmp_path/'out/S001.pptx').exists()



def test_rc9_minimal_unopenable_package_is_not_single_page(tmp_path):
    import pytest,canva_bridge as bridge
    path=_rc9_package_cases(tmp_path)['minimal']
    with pytest.raises(ValueError,match='GRAPHICS_FIRST_PPTX_INVALID'):bridge.single_page(path)


def test_rc9_default_finalizer_rejects_minimal_unopenable_package(tmp_path):
    import subprocess,shutil,pytest
    path=_rc9_package_cases(tmp_path)['minimal'];before=path.read_bytes()
    shell=shutil.which('pwsh') or shutil.which('powershell')
    if not shell: pytest.skip('PowerShell not available')
    script=Path(__file__).resolve().parents[1]/'scripts/canva_pptx_finalize.ps1'
    cp=subprocess.run([shell,'-NoProfile','-File',str(script),'-DownloadedFile',str(path),'-OutputDirectory',str(tmp_path/'out'),
                       '-TargetFileName','S001.pptx','-StableSamples','1'],capture_output=True,text=True)
    assert cp.returncode!=0 and 'GRAPHICS_FIRST_PPTX_INVALID' in cp.stderr
    assert path.read_bytes()==before and not (tmp_path/'out/S001.pptx').exists()



def _rc9_assembly_cli(*args):
    import subprocess,json
    root=Path(__file__).resolve().parents[1]
    # Only native environment/import is substituted; run the real assembly CLI and merge checks.
    code="""import sys,runpy,copy
from pathlib import Path
root=Path(sys.argv[1]);sys.path.insert(0,str(root/'scripts'))
import preflight,merge_pptx
from pptx import Presentation
original=preflight.importlib.util.find_spec
preflight.importlib.util.find_spec=lambda n:object() if n=='win32com' else original(n)
preflight._probe_powerpoint=lambda _: (True,'OFFLINE native capability substitute')
def native(inputs,order,target):
    by_id={x['slide_id']:x for x in inputs};prs=Presentation()
    first=Presentation(by_id[order[0]]['path']);prs.slide_width=first.slide_width;prs.slide_height=first.slide_height
    for sid in order:
        src=Presentation(by_id[sid]['path']).slides[0];dst=prs.slides.add_slide(prs.slide_layouts[6])
        for shape in src.shapes:dst.shapes._spTree.insert_element_before(copy.deepcopy(shape.element),'p:extLst')
    prs.save(target)
merge_pptx._native_merge=native
sys.argv=[str(root/'scripts/assemble_deck.py'),*sys.argv[2:]]
runpy.run_path(sys.argv[0],run_name='__main__')
"""
    cp=subprocess.run([sys.executable,'-X','utf8','-c',code,str(root),*map(str,args)],capture_output=True,text=True,encoding='utf-8',cwd=root)
    assert 'Traceback' not in cp.stderr,cp.stderr
    return cp,json.loads(cp.stdout)


def test_rc9_three_page_documented_cli_path_with_real_native_text(tmp_path):
    import json,hashlib,runtime,canva_bridge as bridge
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    state=tmp_path/'state.json'
    def run(script,*args):
        cp,data=_rc9_cli(script,*args);assert cp.returncode==0 and cp.stderr=='',data;return data
    def write(name,data):
        path=tmp_path/name;path.write_text(json.dumps(data),encoding='utf-8');return path
    run('canva_bridge.py','init','--state',state)
    outline=write('outline.json',{'slides':[{k:k+str(i) for k in runtime.OUTLINE_FIELDS} for i in range(3)]})
    run('runtime.py','stage1-draft','--state',state,'--outline',outline)
    assert run('runtime.py','stage1-review','--state',state)['review_view']
    run('runtime.py','stage1-reply','--state',state,'--message','同意，按这个继续。','--decision','approve')
    run('runtime.py','stage2-entry','--state',state)
    assert run('canva_bridge.py','start-run','--state',state)['status']=='RUN_READY'
    for kind in ('anchor','style_dna'):
        run('runtime.py','stage2-artifact-register','--state',state,'--kind',kind,'--artifact',write(kind+'.json',{'fixture':'OFFLINE'}))
    for sid in ('S001','S002','S003'):
        render=tmp_path/(sid+'.png');Image.new('RGB',(1600,900),'white').save(render)
        candidate=run('runtime.py','stage2-candidate-register','--state',state,'--slide-id',sid,'--artifact',render,'--generation-intent','OFFLINE qualification fixture')
        run('runtime.py','stage2-candidate-review','--state',state,'--candidate-id',candidate['candidate_id'],'--verdict','PASS')
        truth=write(sid+'-truth.json',{'slide_id':sid,'meaningful_text':['OFFLINE '+sid],'exact_facts':[]})
        run('runtime.py','stage2-artifact-register','--state',state,'--kind','final_content_truth:'+sid,'--artifact',truth)
        run('runtime.py','stage2-text-reconciled','--state',state,'--slide-id',sid)
        canvas=write(sid+'-canvas.json',{'width':1600,'height':900})
        run('runtime.py','stage2-artifact-register','--state',state,'--kind','canvas:'+sid,'--artifact',canvas)
    display=run('runtime.py','stage2-formal-display','--state',state)
    assert [item['slide_id'] for item in display['renders']]==['S001','S002','S003']
    run('runtime.py','stage2-visual-approve','--state',state,'--message','确认，按这组页面进入下一阶段。')
    for sid in ('S001','S002','S003'):
        from stage2_test_helpers import write_stage3_text_plan
        manifest_path,inventory_path=write_stage3_text_plan(tmp_path/(sid+'-manifest.json'),tmp_path/(sid+'-inventory.json'),sid,'OFFLINE '+sid)
        run('canva_bridge.py','register-text-plan','--state',state,'--slide-id',sid,'--manifest',manifest_path,'--inventory',inventory_path)
        run('runtime.py','stage2-artifact-register','--state',state,'--kind','font_fallback:'+sid,'--artifact',write(sid+'-font_fallback.json',{}))
        render=tmp_path/(sid+'.png')
        run('canva_bridge.py','register-text-clean','--state',state,'--slide-id',sid,'--artifact',render)
    assert run('runtime.py','stage2-handoff-status','--state',state)['status']=='COMPLETE'
    for sid in ('S001','S002','S003'):
        attempt=run('canva_bridge.py','begin-attempt','--state',state,'--slide-id',sid)
        result=write(sid+'-result.json',{**attempt,'design_id':'OFFLINE-'+sid,'assessment':'ACCEPT','evidence':'OFFLINE external response fixture'})
        run('canva_bridge.py','accept-attempt','--state',state,'--slide-id',sid,'--attempt-id',attempt['attempt_id'],'--result',result)
        graphics=tmp_path/(sid+'-graphics.pptx');prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(graphics)
        evidence=write(sid+'-download.json',{'attempt_id':attempt['attempt_id'],'slide_id':sid,'design_id':'OFFLINE-'+sid,
             'pptx_sha256':bridge.sha(graphics),'edit_url':'https://www.canva.com/design/OFFLINE-'+sid+'/edit',
             'download_entry':'OFFLINE download event fixture','completion_evidence':'OFFLINE complete fixture'})
        run('canva_bridge.py','bind-download','--state',state,'--slide-id',sid,'--pptx',graphics,'--evidence',evidence)
        restored=tmp_path/(sid+'-restored.pptx');prs.slides[0].shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text='OFFLINE '+sid;prs.save(restored)
        assert Presentation(restored).slides[0].shapes[0].text=='OFFLINE '+sid
        # Finalized Text Manifest was validated and registered before destructive cleanup.
        assert runtime.stage2_artifact_current(state,'finalized_manifest:'+sid)
        assert runtime.stage2_artifact_current(state,'font_fallback:'+sid)
        checks={k:'PASS' for k in runtime.GATE_NAMES}
        gates=write(sid+'-review.json',{'attempt_id':attempt['attempt_id'],'slide_id':sid,'design_id':'OFFLINE-'+sid,
            'download_sha256':bridge.sha(graphics),'restored_pptx_sha256':bridge.sha(restored),
            'text_restore_fingerprint':runtime.expected_lineage_from_state(runtime.load_runtime_state(state),sid)['text_restore_fingerprint'],
            'checks':checks,'evidence':{k:'OFFLINE review of synthetic native-text fixture' for k in checks}})
        run('canva_bridge.py','seal-page','--state',state,'--slide-id',sid,'--pptx',restored,'--gates',gates)
    assert run('canva_bridge.py','status','--state',state)['status']=='PREPARE_DECK'
    cp,prepared=_rc9_assembly_cli('prepare','--state',state,'--work-dir',tmp_path/'work')
    assert cp.returncode==0 and prepared['status']=='AWAITING_DECK_VALIDATION'
    cp,_=_rc9_assembly_cli('publish','--state',state,'--validation-report',prepared['validation_report'],'--output',tmp_path/'outputs/deck.pptx')
    assert cp.returncode==1 and not (tmp_path/'outputs/deck.pptx').exists()
    _complete_deck_review(Path(prepared['validation_report']))
    cp,published=_rc9_assembly_cli('publish','--state',state,'--validation-report',prepared['validation_report'],'--output',tmp_path/'outputs/deck.pptx')
    assert cp.returncode==0 and published['status']=='PASS'
    deck=Presentation(published['path']);assert [slide.shapes[0].text for slide in deck.slides]==['OFFLINE S001','OFFLINE S002','OFFLINE S003']
    assert run('canva_bridge.py','status','--state',state)['status']=='COMPLETE'


def test_rc9_hash_matching_historical_invalid_files_do_not_resume_as_valid(tmp_path):
    import copy,pytest,canva_bridge as bridge,runtime,assemble_deck
    for kind in ('graphics','restored'):
        folder=tmp_path/kind;folder.mkdir();state,_,pages=_ready_codex_deck(folder)
        current=runtime.load_runtime_state(state);graphics,restored,_=pages['S001']
        path=graphics if kind=='graphics' else restored;path.write_bytes(b'OFFLINE invalid file')
        if kind=='graphics':
            current['canva_bridge']['downloads']['S001']['sha256']=bridge.sha(path)
            current['stage2_run']['artifacts']['graphics_first_pptx:S001']['sha256']=bridge.sha(path)
        else:
            current['slides']['S001']['validated_single_page']['validated_pptx_sha256']=bridge.sha(path)
            current['canva_bridge']['validated']['S001']['seal']=copy.deepcopy(current['slides']['S001']['validated_single_page'])
        runtime._save_runtime_state(state,current);before=state.read_bytes()
        with pytest.raises(ValueError,match='GRAPHICS_FIRST_PPTX_INVALID'):assemble_deck.current_inputs(runtime.load_runtime_state(state))
        assert bridge.status(state)['status'] not in ('COMPLETE','PREPARE_DECK')
        assert state.read_bytes()==before


def test_rc9_finalizer_missing_python_dependency_does_not_move_source(tmp_path):
    import subprocess,shutil,os,pytest
    valid=_rc9_package_cases(tmp_path)['valid'];before=valid.read_bytes()
    shadow=tmp_path/'shadow';shadow.mkdir();(shadow/'pptx.py').write_text("raise ImportError('OFFLINE unavailable pptx dependency')")
    env=os.environ.copy();env['PYTHONPATH']=str(shadow)
    shell=shutil.which('pwsh') or shutil.which('powershell')
    if not shell: pytest.skip('PowerShell not available')
    script=Path(__file__).resolve().parents[1]/'scripts/canva_pptx_finalize.ps1'
    cp=subprocess.run([shell,'-NoProfile','-File',str(script),'-DownloadedFile',str(valid),'-OutputDirectory',str(tmp_path/'out'),
                       '-TargetFileName','S001.pptx','-StableSamples','1','-PythonExecutable',sys.executable],capture_output=True,text=True,env=env)
    assert cp.returncode!=0 and 'PPTX_VALIDATION_DEPENDENCY_MISSING' in cp.stderr
    assert 'GRAPHICS_FIRST_PPTX_INVALID' not in cp.stderr
    assert valid.read_bytes()==before and not (tmp_path/'out/S001.pptx').exists()


# rc10: use actual bridge/assembly CLI; only external/native edges are fixtures.
def test_rc10_changed_canvas_blocks_old_clean_and_new_registration_resumes(tmp_path):
    import json,canva_bridge as bridge,runtime
    state,_,_=_ready_codex_deck(tmp_path)
    saved=runtime.load_runtime_state(state);old= saved['canva_bridge']['active']['S001'].copy()
    canvas=tmp_path/'new-canvas.json';canvas.write_text(json.dumps({'width':1920,'height':1080}))
    runtime.register_stage2_artifact(state,'canvas:S001',canvas)
    before=state.read_bytes()
    cp,data=_rc9_cli('canva_bridge.py','begin-attempt','--state',state,'--slide-id','S001','--rebuild')
    assert cp.returncode==1 and cp.stderr=='' and data['status']=='BLOCKED'
    assert state.read_bytes()==before and bridge.status(state)['status']=='PREPARE_TEXT'
    # A human has rechecked the formal clean for the new source; no remote call here.
    clean=runtime.load_runtime_state(state)['slides']['S001']['text_clean']
    cp,data=_rc9_cli('canva_bridge.py','register-text-clean','--state',state,'--slide-id','S001','--artifact',clean)
    assert cp.returncode==0 and cp.stderr=='' and data['status']=='TEXT_CLEAN_REGISTERED'
    assert runtime.load_runtime_state(state)['canva_bridge']['active']['S001']==old
    cp,data=_rc9_cli('canva_bridge.py','begin-attempt','--state',state,'--slide-id','S001','--rebuild')
    assert cp.returncode==0 and data['status']=='PENDING' and data['attempt_id']!=old['attempt_id']
    assert data['source_fingerprint']!=old['source_fingerprint']


def test_rc10_same_canvas_semantics_and_legacy_local_binding_keep_design(tmp_path):
    import json,canva_bridge as bridge,runtime
    state,_,_=_ready_codex_deck(tmp_path)
    canvas=tmp_path/'same-canvas.json';canvas.write_text(json.dumps({'height':900,'width':1600},indent=2))
    runtime.register_stage2_artifact(state,'canvas:S001',canvas)
    assert bridge.status(state)['status']=='PREPARE_DECK'
    saved=runtime.load_runtime_state(state);design=saved['canva_bridge']['active']['S001'].copy()
    saved['canva_bridge'].pop('text_preparation');runtime._save_runtime_state(state,saved)
    assert bridge.status(state)['status']=='PREPARE_TEXT'
    before=state.read_bytes()
    cp,data=_rc9_cli('canva_bridge.py','begin-attempt','--state',state,'--slide-id','S001','--rebuild')
    assert cp.returncode==1 and data['status']=='BLOCKED' and state.read_bytes()==before
    for sid in saved['stage2_run']['page_order']:
        bridge.register_text_clean(state,sid,saved['slides'][sid]['text_clean'])
    current=runtime.load_runtime_state(state)
    assert current['canva_bridge']['active']['S001']==design
    assert 'validated_deck' not in current and bridge.status(state)['status']=='PREPARE_DECK'


def test_rc10_pending_binding_repair_preserves_attempt_and_budget(tmp_path):
    import pytest,canva_bridge as bridge,runtime
    from test_preflight_scope import _bridge_fixture
    state=_bridge_fixture(tmp_path,n=1);attempt=bridge.begin_attempt(state,'S001')
    current=runtime.load_runtime_state(state);current['canva_bridge']['active']['S001']['retry_state']='consumed'
    current['canva_bridge'].pop('text_preparation');runtime._save_runtime_state(state,current)
    before=state.read_bytes();assert bridge.status(state)['status']=='BLOCKED';assert state.read_bytes()==before
    bridge.register_text_clean(state,'S001',current['slides']['S001']['text_clean'])
    saved=runtime.load_runtime_state(state)
    assert saved['canva_bridge']['active']==current['canva_bridge']['active']
    assert bridge.status(state)['status']=='WAIT_RECONSTRUCTION'
    before=state.read_bytes()
    with pytest.raises(ValueError):bridge.begin_attempt(state,'S001',True)
    assert state.read_bytes()==before


def test_rc10_register_failure_is_atomic_and_current_page_only(tmp_path,monkeypatch):
    import pytest,canva_bridge as bridge,runtime
    from test_preflight_scope import _bridge_fixture
    state=_bridge_fixture(tmp_path);saved=runtime.load_runtime_state(state)
    # Global handoff may be incomplete; current-page registration is still valid.
    saved['stage2_run']['text_reconciliation'].pop('S003');runtime._save_runtime_state(state,saved)
    result=bridge.register_text_clean(state,'S001',saved['slides']['S001']['text_clean'])
    assert result['status']=='TEXT_CLEAN_REGISTERED'
    current=runtime.load_runtime_state(state)
    assert len(current['canva_bridge']['text_preparation_history']['S001'])==1
    before=state.read_bytes()
    with pytest.raises(ValueError):bridge.register_text_clean(state,'S003',saved['slides']['S003']['text_clean'])
    assert state.read_bytes()==before
    original=bridge.os.replace
    def fail(src,dst):
        if Path(dst).resolve()==state.resolve():raise OSError('OFFLINE commit failure')
        return original(src,dst)
    with monkeypatch.context() as patch:
        patch.setattr(bridge.os,'replace',fail)
        with pytest.raises(OSError):bridge.register_text_clean(state,'S001',saved['slides']['S001']['text_clean'])
    assert state.read_bytes()==before and not state.with_name(state.name+'.bridge-lock').exists()
    assert not runtime.load_runtime_state(state).get('validated_deck')


def test_rc10_new_run_archives_binding_and_resume_keeps_it(tmp_path):
    import json,canva_bridge as bridge,runtime
    from test_preflight_scope import _bridge_fixture
    state=_bridge_fixture(tmp_path,n=1);before=state.read_bytes();bridge.start_run(state)
    assert state.read_bytes()==before
    old=runtime.load_runtime_state(state)['canva_bridge']['text_preparation']
    outline=tmp_path/'replacement.json'
    outline.write_text(json.dumps({'slides':[{k:k+' replacement' for k in runtime.OUTLINE_FIELDS} for i in range(2)]}),encoding='utf-8')
    runtime.record_stage1_draft(state,outline);runtime.present_stage1_review(state);runtime.handle_stage1_reply(state,'同意，按这个继续。', decision='approve')
    bridge.start_run(state);saved=runtime.load_runtime_state(state)
    assert 'text_preparation' not in saved['canva_bridge']
    assert saved['canva_bridge']['run_history'][-1]['bridge_records']['text_preparation']==old


import pytest as _rc10_pytest
@_rc10_pytest.mark.parametrize('invalid',[[],None,'invalid',{'checks':[]},{'evidence':[]}])
def test_rc10_invalid_review_cli_is_controlled_and_atomic(tmp_path,monkeypatch,invalid):
    import json,assemble_deck
    state,_,_=_ready_codex_deck(tmp_path);_local_native_only(monkeypatch)
    prepared=assemble_deck.prepare(state,tmp_path/'work');review=Path(prepared['validation_report'])
    _complete_deck_review(review)
    if isinstance(invalid,dict):
        value=json.loads(review.read_text());value.update(invalid)
    else:value=invalid
    review.write_text(json.dumps(value));before=state.read_bytes();output=tmp_path/'outputs/deck.pptx'
    cp,data=_rc9_cli('assemble_deck.py','publish','--state',state,'--validation-report',review,'--output',output)
    assert cp.returncode==1 and cp.stderr=='' and data['status']=='FAIL'
    assert data['code']=='DECK_LEVEL_VALIDATION_FAILED' and data['details']
    assert state.read_bytes()==before and not output.exists()



def test_rc10_registration_write_failure_keeps_design_and_no_new_pass(tmp_path,monkeypatch):
    import pytest,canva_bridge as bridge,runtime
    from test_preflight_scope import _bridge_fixture
    state=_bridge_fixture(tmp_path,n=1);bridge.begin_attempt(state,'S001')
    saved=runtime.load_runtime_state(state);before=state.read_bytes()
    original=runtime._save_runtime_state
    def fail(path,value):
        if 'text_preparation_history' in value.get('canva_bridge',{}):raise OSError('OFFLINE binding write failure')
        return original(path,value)
    with monkeypatch.context() as patch:
        patch.setattr(runtime,'_save_runtime_state',fail)
        with pytest.raises(OSError):bridge.register_text_clean(state,'S001',saved['slides']['S001']['text_clean'])
    assert state.read_bytes()==before and bridge.status(state)['status']=='WAIT_RECONSTRUCTION'
    assert not state.with_name(state.name+'.bridge-lock').exists()
    assert not runtime.load_runtime_state(state).get('validated_deck')


def test_rc10_changed_canvas_invalidates_download_page_and_deck_review(tmp_path,monkeypatch):
    import json,pytest,canva_bridge as bridge,runtime,assemble_deck
    state,_,pages=_ready_codex_deck(tmp_path);_local_native_only(monkeypatch)
    prepared=assemble_deck.prepare(state,tmp_path/'work');review=Path(prepared['validation_report']);_complete_deck_review(review)
    canvas=tmp_path/'changed-canvas.json';canvas.write_text(json.dumps({'width':1920,'height':1080}))
    runtime.register_stage2_artifact(state,'canvas:S001',canvas)
    current=runtime.load_runtime_state(state);before=state.read_bytes()
    with pytest.raises(ValueError):bridge.download_current(current,'S001')
    with pytest.raises(ValueError):bridge.page_provenance(current,'S001')
    cp,data=_rc9_cli('assemble_deck.py','publish','--state',state,'--validation-report',review,'--output',tmp_path/'outputs/deck.pptx')
    assert cp.returncode==1 and cp.stderr=='' and data['status']=='FAIL'
    assert state.read_bytes()==before and not (tmp_path/'outputs/deck.pptx').exists()
    assert all(path.exists() for path in pages['S001']) and Path(prepared['path']).exists()


def test_rc17_bound_download_blocks_magic_layers_rebuild_during_restore_text(tmp_path):
    import json, pytest, canva_bridge as bridge, runtime
    from test_preflight_scope import _bridge_fixture
    from pptx import Presentation
    state = _bridge_fixture(tmp_path, n=1)
    active = bridge.begin_attempt(state, 'S001')
    result = {**active, 'design_id': 'design-A', 'assessment': 'ACCEPT', 'evidence': 'synthetic accepted reconstruction'}
    result_path = tmp_path / 'accept.json'; result_path.write_text(json.dumps(result))
    bridge.accept_attempt(state, 'S001', active['attempt_id'], result_path)
    seed = tmp_path / 'seed.pptx'; deck = Presentation(); deck.slides.add_slide(deck.slide_layouts[6]); deck.save(seed)
    graphics, _, _ = _download_and_review(state, 'S001', tmp_path, seed)
    assert bridge.status(state)['status'] == 'RESTORE_TEXT'
    before = state.read_bytes()
    saved = runtime.load_runtime_state(state)
    attempt_id = saved['canva_bridge']['active']['S001']['attempt_id']
    download_sha = saved['canva_bridge']['downloads']['S001']['sha256']
    with pytest.raises(ValueError, match='Graphics-first PPTX already bound'):
        bridge.begin_attempt(state, 'S001', True)
    assert state.read_bytes() == before
    saved = runtime.load_runtime_state(state)
    assert bridge.status(state)['status'] == 'RESTORE_TEXT'
    assert saved['canva_bridge']['active']['S001']['attempt_id'] == attempt_id
    assert saved['canva_bridge']['downloads']['S001']['sha256'] == download_sha == bridge.sha(graphics)
