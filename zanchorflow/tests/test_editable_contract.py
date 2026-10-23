"""Offline contract and local algorithm tests; no model-behavior claims."""
import sys,json,hashlib,copy
from pathlib import Path
import pytest,yaml
from PIL import Image,ImageDraw
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
from editable_split import region_components
from editable_formula import build_latex_document
from editable_build import build,validate
from test_editable_build import plan

def test_ninth_protocol_registered_without_changing_historical_versions():
 manifest=yaml.safe_load((ROOT/'manifest.yaml').read_text(encoding='utf-8'));items=manifest['canonical_protocols']
 assert len(items)==9 and manifest['skill_version']=='1.0' and manifest['protocol_baseline']=='RC17'
 assert items[-1]['path']=='references/09_Editable_Path重建协议_v1.0.md'
 for item in items:assert hashlib.sha256((ROOT/item['path']).read_bytes()).hexdigest()==item['sha256']

def test_whole_object_regions_preserve_disconnected_faint_detail(tmp_path):
 image=Image.new('RGBA',(120,90));draw=ImageDraw.Draw(image);draw.rectangle([20,25,65,65],fill=(20,70,90,255));draw.ellipse([45,8,49,12],fill=(20,70,90,7))
 path=tmp_path/'regions.json';path.write_text(json.dumps({'regions':[{'name':'icon','box':[5,3,80,78]}]}))
 result=region_components(image,path)
 assert len(result)==1 and result[0]['box'][1]==8

def test_regions_reject_unowned_or_boundary_foreground(tmp_path):
 image=Image.new('RGBA',(100,100));ImageDraw.Draw(image).rectangle([10,10,30,30],fill='red');path=tmp_path/'regions.json';path.write_text(json.dumps({'regions':[{'name':'icon','box':[10,10,40,40]}]}))
 with pytest.raises(SystemExit,match='boundary'):region_components(image,path)

def test_formula_display_wrapper_is_valid_in_standalone_lr_box():
 document=build_latex_document(r'\mathbf{x}=\frac{a}{b}')
 assert r'$\displaystyle' in document and r'\mathbf{x}' in document and r'\[' not in document

def test_final_manifest_rebuild_preserves_object_structure(tmp_path):
 p=plan();one=build(p,tmp_path/'one.pptx');resolved=json.loads(Path(one['resolved_manifest']).read_text(encoding='utf-8'));two=build(resolved['plan'],tmp_path/'two.pptx')
 assert one['objects']==two['objects'] and one['occurrences']==two['occurrences']
 assert resolved['pptx_sha256']==one['sha256']

def test_rich_text_mismatch_is_rejected_before_build(tmp_path):
 p=plan();p['objects'][0]['paragraphs']=[{'runs':[{'text':'wrong'}]}]
 with pytest.raises(ValueError,match='formatted'):build(p,tmp_path/'bad.pptx')


def test_chroma_preserves_existing_faint_foreground_alpha():
 from editable_chroma import _apply_alpha_to_image
 im=Image.new('RGBA',(3,1));im.putpixel((0,0),(0,255,0,255));im.putpixel((1,0),(20,70,90,7));im.putpixel((2,0),(20,70,90,128))
 _apply_alpha_to_image(im,key=(0,255,0),tolerance=10,spill_cleanup=False,soft_matte=False,transparent_threshold=10,opaque_threshold=80)
 assert [im.getpixel((i,0))[3] for i in range(3)]==[0,7,128]


def test_source_corner_radius_is_explicit_and_native(tmp_path):
 from pptx import Presentation
 p=plan();p['objects']=[{'id':'card','type':'shape','preset':'roundRect','box_px':[10,10,700,200],'source_corner_radius_px':8,'fill':'FFFFFF','stroke':'000000'}]
 result=build(p,tmp_path/'corner.pptx');shape=Presentation(result['path']).slides[0].shapes[0]
 assert abs(shape.adjustments[0]-0.04)<1e-6
 p['objects'][0]['source_corner_radius_px']=101
 with pytest.raises(ValueError,match='corner'):validate(p)


def test_table_border_children_precede_fill_per_drawingml(tmp_path):
 from pptx import Presentation
 from pptx.oxml.ns import qn
 p=plan();p['objects'][1]['style']={'fill':'FFFFFF','borders':{'left':{'color':'203040','width':1}}};result=build(p,tmp_path/'borders.pptx');table=next(s.table for s in Presentation(result['path']).slides[0].shapes if s.has_table);tags=[child.tag for child in table.cell(0,0)._tc.tcPr]
 assert tags.index(qn('a:lnL'))<tags.index(qn('a:solidFill'))


def test_one_logical_curve_cannot_be_declared_as_many_fragments():
 p=plan();p['objects']=[{'id':'a','type':'path','semantic_line_id':'one-curve','box_px':[10,10,100,100],'path_px':[{'op':'moveTo','points':[[10,10]]},{'op':'lnTo','points':[[110,110]]}]},{'id':'b','type':'path','semantic_line_id':'one-curve','box_px':[110,110,100,100],'path_px':[{'op':'moveTo','points':[[110,110]]},{'op':'lnTo','points':[[210,210]]}]}]
 with pytest.raises(ValueError,match='logical line'):validate(p)
