import hashlib, json, sys
from pathlib import Path
import pytest
from pptx import Presentation
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from runtime import (resume_stage, OUTLINE_FIELDS, record_stage1_draft,
                     present_stage1_review, handle_stage1_reply,
                     start_stage2_run, register_stage2_artifact)
from merge_pptx import validate_merge_inputs, merge
from stage2_test_helpers import write_png, write_truth, complete_stage2_page


def case(tmp_path):
    def put(name, value):
        p=tmp_path/name
        if isinstance(value, bytes): p.write_bytes(value)
        else: p.write_text(json.dumps(value),encoding='utf-8')
        return str(p)
    pptx=tmp_path/'S01.pptx'
    prs=Presentation(); prs.slides.add_slide(prs.slide_layouts[6]); prs.save(pptx)
    state={'deck_order':['S01'], 'slides':{'S01':{
        'approved_render':write_png(tmp_path/'render.png', b'render-v1'),
        'final_content_truth':write_truth(tmp_path/'truth.json', 'S01', 'v1'),
        'canvas':put('canvas.json',{'width':10,'height':7.5}),
        'text_clean':put('clean.png',b'clean-v1'),
        'finalized_manifest':put('manifest.json',{'text':'v1'}),
        'font_fallback':put('fonts.json',{}),
    }}}
    state_path=tmp_path/'runtime-state.json'; state_path.write_text(json.dumps(state),encoding='utf-8')
    outline_path=tmp_path/'outline.json'
    outline_path.write_text(json.dumps({'slides':[{field:field for field in OUTLINE_FIELDS}]}),encoding='utf-8')
    record_stage1_draft(state_path,outline_path)
    present_stage1_review(state_path)
    handle_stage1_reply(state_path,'同意，按这个继续。', decision='approve')
    start_stage2_run(state_path)
    for kind in ('anchor','style_dna'):
        asset = tmp_path / kind
        asset.write_bytes(kind.encode())
        register_stage2_artifact(state_path, kind, asset)
    complete_stage2_page(__import__('runtime'), state_path, 'S01', state['slides']['S01']['approved_render'], state['slides']['S01']['final_content_truth'])
    return state_path,pptx,state


def seal_case(state_path,pptx):
    from runtime import seal_validated_single_page
    return seal_validated_single_page(state_path,'S01',pptx,{k:'PASS' for k in ('content_truth','semantic_fidelity','functional_editability','visual_fidelity')})


def assert_merge_rejects(state_path,record,tmp_path,reason):
    from runtime import validate_current_artifact
    state=json.loads(state_path.read_text(encoding='utf-8'))
    assert reason in validate_current_artifact(record,state,'S01')
    result=merge([record],['S01'],tmp_path/'merged.pptx',state_path)
    assert result['code']=='DECK_MERGE_INTEGRITY_FAILED'
    assert reason in result['details']
    assert not (tmp_path/'merged.pptx').exists()


def test_stale_source_fingerprint(tmp_path):
    state_path,pptx,state=case(tmp_path); record=seal_case(state_path,pptx)
    Path(state['slides']['S01']['approved_render']).write_bytes(b'render-v2')
    assert_merge_rejects(state_path,record,tmp_path,'stale_source_fingerprint')


def test_stale_text_clean_fingerprint(tmp_path):
    state_path,pptx,state=case(tmp_path); record=seal_case(state_path,pptx)
    Path(state['slides']['S01']['text_clean']).write_bytes(b'clean-v2')
    assert_merge_rejects(state_path,record,tmp_path,'stale_text_clean_fingerprint')


def test_stale_text_restore_fingerprint(tmp_path):
    state_path,pptx,state=case(tmp_path); record=seal_case(state_path,pptx)
    Path(state['slides']['S01']['font_fallback']).write_text(json.dumps({'font':'other'}),encoding='utf-8')
    assert_merge_rejects(state_path,record,tmp_path,'stale_text_restore_fingerprint')


def test_validated_pptx_bytes_changed(tmp_path):
    state_path,pptx,state=case(tmp_path); record=seal_case(state_path,pptx)
    pptx.write_bytes(pptx.read_bytes()+b'changed')
    assert_merge_rejects(state_path,record,tmp_path,'validated_pptx_hash_mismatch')


def test_resume_rejects_stale_or_unsealed_flag(tmp_path):
    state_path,pptx,state=case(tmp_path); record=seal_case(state_path,pptx)
    assert resume_stage({'slide_id':'S01','validated_single_page_pptx':True},state_path)=='deck_merge'
    pptx.write_bytes(pptx.read_bytes()+b'changed')
    assert resume_stage({'slide_id':'S01','validated_single_page_pptx':True},state_path)!='deck_merge'
    assert resume_stage({'validated_single_page_pptx':True})!='deck_merge'


def test_seal_requires_four_gates_and_is_frozen(tmp_path):
    from runtime import seal_validated_single_page
    state_path,pptx,state=case(tmp_path)
    with pytest.raises(ValueError): seal_validated_single_page(state_path,'S01',pptx,{'content_truth':'PASS'})
    record=seal_case(state_path,pptx)
    assert record['validated_pptx_sha256']==hashlib.sha256(pptx.read_bytes()).hexdigest()
    with pytest.raises(ValueError): seal_case(state_path,pptx)

def test_new_gate_pass_can_supersede_stale_seal(tmp_path):
    state_path,pptx,state=case(tmp_path)
    old=seal_case(state_path,pptx)
    Path(state['slides']['S01']['approved_render']).write_bytes(b'render-v2')
    new=seal_case(state_path,pptx)
    assert new['source_fingerprint'] != old['source_fingerprint']
    assert_merge_rejects(state_path,old,tmp_path,'stale_source_fingerprint')

def test_stale_job_claim_cannot_override_trusted_state_or_start_com(tmp_path,monkeypatch):
    import merge_pptx
    state_path,pptx,state=case(tmp_path);record=seal_case(state_path,pptx)
    Path(state['slides']['S01']['approved_render']).write_bytes(b'new approved render')
    record['expected_lineage']={k:record[k] for k in ('source_fingerprint','text_clean_fingerprint','text_restore_fingerprint')}
    def forbidden(*args,**kwargs): raise AssertionError('PowerPoint COM import must not run')
    monkeypatch.setattr(merge_pptx,'_native_merge',forbidden)
    result=merge_pptx.merge([record],['S01'],tmp_path/'out.pptx',state_path)
    assert result['code']=='DECK_MERGE_INTEGRITY_FAILED'
    assert 'stale_source_fingerprint' in result['details']
