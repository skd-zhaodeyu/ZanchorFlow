import hashlib
import json
import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import runtime
from stage2_test_helpers import write_png, write_truth, complete_stage2_page


GATES = {name: "PASS" for name in runtime.GATE_NAMES}
FLAGS = {
    "slide_id": "S01",
    "validated_single_page_pptx": True,
    "merged_pptx": True,
    "validated_deck_pptx": True,
}


def _write(path, data):
    path.write_bytes(data)
    return str(path)


def _case(tmp_path):
    slides = {}
    for slide_id in ("S01", "S02", "S03"):
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        pptx = tmp_path / (slide_id + ".pptx")
        prs.save(pptx)
        slides[slide_id] = {
            "approved_render": write_png(tmp_path / (slide_id + "-render.png"), slide_id.encode()),
            "final_content_truth": write_truth(tmp_path / (slide_id + "-truth.json"), slide_id),
            "canvas": _write(tmp_path / (slide_id + "-canvas.json"), b'{"width":10}'),
            "text_clean": _write(tmp_path / (slide_id + "-clean.png"), b"clean"),
            "finalized_manifest": _write(tmp_path / (slide_id + "-manifest.json"), b'{"text":"x"}'),
            "font_fallback": _write(tmp_path / (slide_id + "-fonts.json"), b"{}"),
        }
    state_path = tmp_path / "state.json"
    state_path.write_text(json.dumps({"deck_order": ["S01", "S02", "S03"], "slides": slides}), encoding="utf-8")
    outline = {'slides': [{field: field + '-' + slide_id for field in runtime.OUTLINE_FIELDS}
                          for slide_id in slides]}
    outline_path = tmp_path / 'outline.json'
    outline_path.write_text(json.dumps(outline), encoding='utf-8')
    runtime.record_stage1_draft(state_path, outline_path)
    runtime.present_stage1_review(state_path)
    runtime.handle_stage1_reply(state_path, '同意，按这个继续。', decision='approve')
    runtime.start_stage2_run(state_path)
    for kind in ('anchor','style_dna'):
        asset = tmp_path / kind
        asset.write_bytes(kind.encode())
        runtime.register_stage2_artifact(state_path, kind, asset)
    for slide_id, slide in slides.items():
        complete_stage2_page(runtime, state_path, slide_id, slide['approved_render'], slide['final_content_truth'])
    records = {}
    for slide_id in slides:
        records[slide_id] = runtime.seal_validated_single_page(
            state_path, slide_id, tmp_path / (slide_id + ".pptx"), GATES
        )
    merged = tmp_path / "merged.pptx"
    prs = Presentation()
    for _ in range(3):
        prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(merged)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    merged_hash = hashlib.sha256(merged.read_bytes()).hexdigest()
    state["merged_deck"] = {
        "deck_order": ["S01", "S02", "S03"],
        "slide_inputs": [
            {"slide_id": sid, "validated_pptx_sha256": records[sid]["validated_pptx_sha256"]}
            for sid in ("S01", "S02", "S03")
        ],
        "merged_pptx_path": str(merged),
        "merged_pptx_sha256": merged_hash,
    }
    state["validated_deck"] = {"status": "PASS", "merged_pptx_sha256": merged_hash}
    state_path.write_text(json.dumps(state), encoding="utf-8")
    return state_path, merged


def _state(path):
    return json.loads(path.read_text(encoding="utf-8"))


def _save(path, state):
    path.write_text(json.dumps(state), encoding="utf-8")


def test_d1_deck_order_change_remerges_without_revalidating_pages(tmp_path):
    state_path, _ = _case(tmp_path)
    state = _state(state_path)
    state["deck_order"] = ["S01", "S03", "S02"]
    _save(state_path, state)
    assert runtime.resume_stage(FLAGS, state_path) == "deck_merge"


def test_d2_new_single_page_seal_invalidates_old_merged_deck(tmp_path):
    state_path, _ = _case(tmp_path)
    state = _state(state_path)
    slide = state["slides"]["S02"]
    path = Path(slide["validated_single_page"]["path"])
    path.write_bytes(path.read_bytes() + b"new validated bytes")
    runtime.seal_validated_single_page(state_path, "S02", path, GATES)
    assert runtime.resume_stage(FLAGS, state_path) == "deck_merge"


def test_d3_merged_bytes_change_invalidates_old_deck(tmp_path):
    state_path, merged = _case(tmp_path)
    merged.write_bytes(merged.read_bytes() + b"changed")
    assert runtime.resume_stage(FLAGS, state_path) == "deck_merge"


def test_d4_remerge_invalidates_prior_deck_validation(tmp_path):
    state_path, merged = _case(tmp_path)
    state = _state(state_path)
    merged.write_bytes(merged.read_bytes() + b"new merge")
    new_hash = hashlib.sha256(merged.read_bytes()).hexdigest()
    state["merged_deck"]["merged_pptx_sha256"] = new_hash
    state["validated_deck"]["merged_pptx_sha256"] = "old-merged-hash"
    _save(state_path, state)
    assert runtime.resume_stage(FLAGS, state_path) == "deck_validation"


def test_d5_current_sealed_deck_completes_without_boolean_authority(tmp_path):
    state_path, _ = _case(tmp_path)
    flags = {**FLAGS, "merged_pptx": False, "validated_deck_pptx": False}
    assert runtime.resume_stage(flags, state_path) == "complete"


def test_deck_resume_uses_trusted_state_without_single_slide_hint(tmp_path):
    state_path, _ = _case(tmp_path)
    assert runtime.resume_stage({}, state_path) == "complete"




def _fake_native_merge(inputs, order, target):
    prs = Presentation()
    for _ in order:
        prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(target)


def test_merge_job_order_must_match_trusted_state(tmp_path, monkeypatch):
    import merge_pptx
    state_path, merged = _case(tmp_path)
    state = _state(state_path)
    records = [state["slides"][sid]["validated_single_page"] for sid in ("S01", "S02", "S03")]
    def forbidden(*args, **kwargs):
        raise AssertionError("COM must not start")
    monkeypatch.setattr(merge_pptx, "_native_merge", forbidden)
    result = merge_pptx.merge(records, ["S01", "S03", "S02"], merged, state_path)
    assert result["status"] == "FAIL"
    assert "deck_order_mismatch" in result["details"]


def test_successful_remerge_clears_old_validation_even_if_bytes_repeat(tmp_path, monkeypatch):
    import merge_pptx
    state_path, merged = _case(tmp_path)
    old = _state(state_path)
    records = [old["slides"][sid]["validated_single_page"] for sid in old["deck_order"]]
    monkeypatch.setattr(merge_pptx, "_native_merge", _fake_native_merge)
    result = merge_pptx.merge(records, old["deck_order"], merged, state_path)
    assert result["status"] == "PASS", result
    state = _state(state_path)
    assert state["merged_deck"]["deck_order"] == old["deck_order"]
    assert state.get("validated_deck") is None
    assert runtime.resume_stage(FLAGS, state_path) == "deck_validation"
    runtime.seal_validated_deck(state_path, "PASS")
    assert runtime.resume_stage(FLAGS, state_path) == "complete"


def test_failed_final_seal_cannot_leave_old_validation_trusted(tmp_path, monkeypatch):
    import merge_pptx
    state_path, merged = _case(tmp_path)
    old = _state(state_path)
    records = [old["slides"][sid]["validated_single_page"] for sid in old["deck_order"]]
    monkeypatch.setattr(merge_pptx, "_native_merge", _fake_native_merge)
    def fail_commit(*args, **kwargs):
        raise OSError("simulated final state commit failure")
    monkeypatch.setattr(merge_pptx, "seal_merged_deck", fail_commit, raising=False)
    result = merge_pptx.merge(records, old["deck_order"], merged, state_path)
    assert result["status"] == "FAIL"
    state = _state(state_path)
    assert state.get("validated_deck") is None
    assert runtime.resume_stage(FLAGS, state_path) == "deck_merge"

