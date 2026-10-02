import sys, zipfile, json, importlib.util
import pytest
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from merge_pptx import merge, inspect_ooxml_dependencies, _shape_signature, GEOMETRY_TOLERANCE_EMU
from runtime import seal_validated_single_page
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pptx.dml.color import RGBColor
from PIL import Image

pytestmark = pytest.mark.skipif(
    importlib.util.find_spec('win32com') is None,
    reason='requires Windows PowerPoint COM / pywin32')


def make_page(path, label, layout, img):
    prs = Presentation()
    master = prs.slide_masters[0]
    master.background.fill.solid()
    master.background.fill.fore_color.rgb = RGBColor(180, 30, 30) if label.startswith('S1') else RGBColor(30, 30, 180)
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text = label
    slide.shapes.add_picture(str(img), Inches(1), Inches(2), width=Inches(1))
    data = CategoryChartData()
    data.categories = ['A','B']
    data.add_series(label, (2,3))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(3), Inches(2), Inches(4), Inches(3), data)
    prs.save(path)


def test_native_merge_preserves_independent_resources(tmp_path):
    img = tmp_path/'image.png'
    Image.new('RGB', (20,20), (80,120,180)).save(img)
    a, b, out = (tmp_path/n for n in ('z-first.pptx','a-second.pptx','merged.pptx'))
    make_page(a, 'S1 editable', 1, img)
    make_page(b, 'S2 editable', 3, img)
    for p in (a,b):
        with zipfile.ZipFile(p) as z:
            assert 'ppt/media/image1.png' in z.namelist()
            assert 'ppt/charts/chart1.xml' in z.namelist()
            assert 'ppt/embeddings/Microsoft_Excel_Sheet1.xlsx' in z.namelist()
    state = {'deck_order':['S2','S1'], 'slides':{}}
    for slide_id, path in [('S1',a),('S2',b)]:
        source_files = {}
        for key, value in [
            ('approved_render',b'render-'+slide_id.encode()),
            ('final_content_truth',{'text':slide_id}),
            ('canvas',{'width':10,'height':7.5}),
            ('text_clean',b'clean-'+slide_id.encode()),
            ('finalized_manifest',{'text':slide_id}),
            ('font_fallback',{}),
        ]:
            item = tmp_path/f'{slide_id}-{key}'
            if isinstance(value,bytes): item.write_bytes(value)
            else: item.write_text(json.dumps(value),encoding='utf-8')
            source_files[key]=str(item)
        state['slides'][slide_id]=source_files
    state_path=tmp_path/'runtime-state.json'
    state_path.write_text(json.dumps(state),encoding='utf-8')
    gates={k:'PASS' for k in ('content_truth','semantic_fidelity','functional_editability','visual_fidelity')}
    records=[seal_validated_single_page(state_path,slide_id,path,gates) for slide_id,path in [('S1',a),('S2',b)]]
    result = merge(records, ['S2','S1'], out, state_path)
    assert result['status'] == 'PASS', result
    prs = Presentation(out)
    assert len(prs.slides) == 2
    assert [next(s.text for s in sl.shapes if s.has_text_frame and s.text) for sl in prs.slides] == ['S2 editable','S1 editable']
    assert all(sum(s.has_chart for s in sl.shapes) == 1 for sl in prs.slides)
    assert [sl.slide_layout.slide_master.background.fill.fore_color.rgb for sl in prs.slides] == [RGBColor(30,30,180), RGBColor(180,30,30)]
    assert not inspect_ooxml_dependencies(out)
    source_signatures=[_shape_signature(Presentation(path))[0] for path in (b,a)]
    output_signatures=_shape_signature(prs)
    deltas=[abs(before-after) for src,dst in zip(source_signatures,output_signatures)
            for source_shape,output_shape in zip(src,dst)
            for before,after in zip(source_shape['geometry'],output_shape['geometry'])]
    assert len(deltas)==44
    assert max(deltas)==0
    assert GEOMETRY_TOLERANCE_EMU==0
    with zipfile.ZipFile(out) as z:
        assert len([n for n in z.namelist() if n.startswith('ppt/charts/chart') and n.endswith('.xml')]) == 2
        assert len([n for n in z.namelist() if n.startswith('ppt/embeddings/') and n.endswith('.xlsx')]) == 2
