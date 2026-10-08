"""New transport-policy regressions. Fixtures are synthetic, not live Canva claims."""
import copy
import io
import json
import shutil
import sys
from pathlib import Path
import urllib.error
import zipfile
import pytest
from pptx import Presentation
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import download_transport as dt
import download_evidence as de
import host_acquisition as h
import canva_bridge as b
import runtime
from test_acquisition_request import setup
from test_file_first_download import make_ppt

URL='https://export-download.canva.com/opaque/fixture-design/1/export.pptx?secret=must-not-persist'

def fixture(tmp_path,title=None):
    state,identity,lookup=setup(tmp_path)
    lookup['observation']='synthetic transport test'
    if title is not None:lookup['response']['design']['title']=title
    lp=tmp_path/'lookup.json';lp.write_text(json.dumps(lookup))
    b.register_design_observation(state,'S001',lp)
    record=tmp_path/'record.json'
    h.prepare(record,state,'S001',lp,tmp_path/'out','S001.pptx')
    directory=tmp_path/'browser';directory.mkdir()
    url='https://www.canva.com/design/fixture-design/edit'
    snap=h.snapshot(record,state,directory,url)
    h.lock(record)
    return state,identity,record,directory,url,snap

def response(monkeypatch,body,length=True,status=200,error=None):
    class Response(io.BytesIO):
        url=URL
        def __init__(self):
            super().__init__(body);self.status=status
            self.headers={'Content-Length':str(len(body))} if length else {}
    class Opener:
        def open(self,url,timeout):
            assert timeout==30
            if error:raise error
            return Response()
    monkeypatch.setattr(dt.urllib.request,'build_opener',lambda *_:Opener())

def pptbytes(tmp_path,pages=1):
    p=tmp_path/'source.pptx';prs=Presentation()
    for _ in range(pages):prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(p);return p.read_bytes()

def finalize(monkeypatch):
    def move(source,out,name,python):
        target=Path(out)/name;target.parent.mkdir(exist_ok=True)
        shutil.move(source,target)
        return {'ValidPptx':True,'SlideCount':1,'Path':str(target),'SHA256':de.sha(target)}
    monkeypatch.setattr(h,'_finalize',move)

@pytest.mark.parametrize('title',[None,'未命名的设计','generic'])
def test_direct_binds_without_name_event_history_or_preview(tmp_path,monkeypatch,title):
    state,identity,record,directory,url,snap=fixture(tmp_path,title)
    response(monkeypatch,pptbytes(tmp_path))
    result=h.direct(record,state,url,URL,tmp_path/'work')
    assert result['status']=='FILE_COMPLETE'
    receipt=de.sealed(de.read(result['file_evidence'])['transport_receipt'])
    assert receipt['get_count']==1 and 'must-not-persist' not in json.dumps(receipt)
    assert receipt['completion_seconds']<1  # no post-stream stability sleep
    finalize(monkeypatch)
    bound=h.finish(record,state,None,url,False,file_evidence=result['file_evidence'])
    assert bound['status']=='DOWNLOAD_BOUND'
    assert h.finish(record,state,None,url,False)['reused']
    current=b.download_current(runtime.load_runtime_state(state),'S001')
    assert current['identity']==identity and current['require_transport_receipt']
    assert not Path(h._read(record)['lease_path']).exists()

@pytest.mark.parametrize('pages',[1,2])
@pytest.mark.parametrize('length',[True,False])
def test_generic_downloader_accepts_expected_pages(tmp_path,monkeypatch,pages,length):
    response(monkeypatch,pptbytes(tmp_path,pages),length)
    result=dt.download(URL,'fixture-design',tmp_path/'work',pages)
    assert result['result']=='PASS' and result['slide_count']==pages
    assert result['sha256']==de.sha(result['path'])

def operation(tmp_path,identity,record,url,path=None,exclusive=None):
    capture={'source':'canva_host.download','identity':identity,'acquisition_id':h._read(record)['acquisition_id'],
             'operation_id':'synthetic-offline-click','observed_url':url,'action':'normal'}
    if path:capture['path']=str(path)
    if exclusive:
        capture.update(exclusive_directory=str(exclusive),directory_owned_by_operation=True)
    op=tmp_path/'actual-operation.json';op.write_text(json.dumps(capture));return op

def test_browser_no_event_path_capture_binds_generic_file(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    file=make_ppt(directory/'未命名的设计.pptx')
    op=operation(tmp_path,identity,record,url,file)
    monkeypatch.setattr(de.time,'sleep',lambda _:None)
    result=h.observe(record,state,snap['snapshot']['path'],1,False,operation_capture=op)
    assert result['status']=='FILE_COMPLETE'
    finalize(monkeypatch)
    assert h.finish(record,state,None,url,False,file_evidence=result['file_evidence'])['status']=='DOWNLOAD_BOUND'

def test_shared_directory_one_file_is_not_source_proof(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    file=make_ppt(directory/'generic.pptx');monkeypatch.setattr(de.time,'sleep',lambda _:None)
    result=h.observe(record,state,snap['snapshot']['path'],1,False)
    assert result['status']=='FILE_SOURCE_PENDING' and file.exists()
    assert not h._read(record).get('file_proofs')

def test_exclusive_directory_capture_can_bind_without_event(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    make_ppt(directory/'generic.pptx')
    op=operation(tmp_path,identity,record,url,exclusive=directory)
    monkeypatch.setattr(de.time,'sleep',lambda _:None)
    assert h.observe(record,state,snap['snapshot']['path'],1,False,operation_capture=op)['status']=='FILE_COMPLETE'

def test_exclusive_directory_multiple_files_pending(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    make_ppt(directory/'one.pptx');make_ppt(directory/'two.pptx')
    op=operation(tmp_path,identity,record,url,exclusive=directory)
    monkeypatch.setattr(de.time,'sleep',lambda _:None)
    assert h.observe(record,state,snap['snapshot']['path'],1,False,operation_capture=op)['status']=='FILE_SOURCE_PENDING'

def test_direct_failure_then_browser_recovery_same_identity(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    response(monkeypatch,b'',error=urllib.error.HTTPError(URL,403,'expired',{},None))
    failed=h.direct(record,state,url,URL,tmp_path/'work')
    assert failed['status']=='DIRECT_UNAVAILABLE'
    recovery=h.snapshot(record,state,directory,url,'recovery',URL)
    file=make_ppt(directory/'generic.pptx')
    op=operation(tmp_path,identity,record,url,file)
    data=de.read(op);data['action']='recovery';data['export_source']=de.recovery_url(URL,identity['design_id']);op.write_text(json.dumps(data))
    monkeypatch.setattr(de.time,'sleep',lambda _:None)
    result=h.observe(record,state,recovery['snapshot']['path'],1,False,URL,operation_capture=op)
    assert result['status']=='FILE_COMPLETE'
    finalize(monkeypatch)
    assert h.finish(record,state,None,url,False,file_evidence=result['file_evidence'])['status']=='DOWNLOAD_BOUND'

@pytest.mark.parametrize('field',de.IDENTITY_KEYS)
def test_identity_mismatch_rejected(tmp_path,monkeypatch,field):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    response(monkeypatch,pptbytes(tmp_path));result=h.direct(record,state,url,URL,tmp_path/'work')
    proof=de.read(result['file_evidence']);proof['identity'][field]='wrong'
    with pytest.raises(ValueError):de.validate_transport_proof(proof,identity)

def test_receipt_removal_cannot_downgrade(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    response(monkeypatch,pptbytes(tmp_path));result=h.direct(record,state,url,URL,tmp_path/'work')
    proof=de.read(result['file_evidence']);proof.pop('transport_receipt')
    with pytest.raises(ValueError,match='RECEIPT_REQUIRED'):b.validate_completion_evidence(runtime.load_runtime_state(state),identity,proof,proof['selected']['path'])

def test_receipt_tamper_and_file_tamper_rejected(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    response(monkeypatch,pptbytes(tmp_path));result=h.direct(record,state,url,URL,tmp_path/'work')
    proof=de.read(result['file_evidence']);rp=Path(proof['transport_receipt']['path'])
    original=rp.read_bytes();rp.write_text('{}')
    with pytest.raises(ValueError,match='EVIDENCE_CHANGED'):de.validate_transport_proof(proof,identity)
    rp.write_bytes(original);Path(proof['selected']['path']).write_bytes(b'wrong')
    with pytest.raises((ValueError,zipfile.BadZipFile)):de.validate_transport_proof(proof,identity)

@pytest.mark.parametrize('url',[
 'http://export-download.canva.com/opaque/fixture-design/1/a.pptx',
 'https://export-download.canva.com.evil.test/opaque/fixture-design/1/a.pptx',
 'https://export-download.canva.com/opaque/wrong/1/a.pptx'])
def test_bad_redirect_or_initial_url_rejected(url):
    with pytest.raises(ValueError):dt.validate_url(url,'fixture-design')
    with pytest.raises(ValueError):dt.Redirect('fixture-design',[]).redirect_request(None,None,302,'',{},url)

def test_valid_redirect_revalidated():
    import urllib.request
    events=[];url='https://export-download.canva.com/opaque/fixture-design/2/new.pptx'
    req=urllib.request.Request(URL)
    assert dt.Redirect('fixture-design',events).redirect_request(req,None,302,'',{},url).full_url==url
    assert events[0]['source']['path'].endswith('new.pptx')

@pytest.mark.parametrize('kind',['zip','pages','length'])
def test_invalid_transfer_never_finalized(tmp_path,monkeypatch,kind):
    body=b'bad' if kind=='zip' else pptbytes(tmp_path,2 if kind=='pages' else 1)
    response(monkeypatch,body)
    if kind=='length':
        class Bad(io.BytesIO):
            status=200;url=URL;headers={'Content-Length':str(len(body)+1)}
        monkeypatch.setattr(dt.urllib.request,'build_opener',lambda *_:type('Opener',(),{'open':lambda *a,**k:Bad(body)})())
    result=dt.download(URL,'fixture-design',tmp_path/'work')
    assert result['result']=='FAIL' and not list((tmp_path/'work').glob('*.pptx'))

def test_retries_and_recovery_snapshots_have_no_hard_budget(tmp_path):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    for n in range(12):
        assert h.reserve_retry(record,state,'S001','ui')['retry_number']==n+1
        h.snapshot(record,state,directory,url,'recovery',URL)
    assert Path(str(record)+'.intent').exists()

def test_wrong_page_count_rejected_by_formal_route(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    response(monkeypatch,pptbytes(tmp_path,2))
    assert h.direct(record,state,url,URL,tmp_path/'work')['status']=='DIRECT_UNAVAILABLE'

def test_new_ready_defaults_to_required_receipt(tmp_path):
    record=tmp_path/'record.json'
    assert h.ready(record,{'status':'WAIT_DOWNLOAD','identity':{},'edit_url':'unused'})['require_transport_receipt']

def test_completed_binding_policy_cannot_be_removed(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    response(monkeypatch,pptbytes(tmp_path));result=h.direct(record,state,url,URL,tmp_path/'work')
    finalize(monkeypatch);h.finish(record,state,None,url,False,file_evidence=result['file_evidence'])
    data=runtime.load_runtime_state(state);data['canva_bridge'].pop('transport_requirements')
    with pytest.raises(ValueError,match='requirement missing'):b.download_current(data,'S001')

def test_all_owned_leases_released_including_after_bind_interruption(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    response(monkeypatch,pptbytes(tmp_path));result=h.direct(record,state,url,URL,tmp_path/'work')
    finalize(monkeypatch)
    h.finish(record,state,None,url,False,file_evidence=result['file_evidence'])
    data=h._read(record)
    assert len(set(data['lease_paths']))==2
    assert all(not Path(p).exists() for p in data['lease_paths'])
    for lease in data['lease_paths']:de.write_once(lease,{'acquisition_id':data['acquisition_id']})
    h.finish(record,state,None,url,False)
    assert all(not Path(p).exists() for p in data['lease_paths'])

def test_three_actual_same_failures_recommend_diagnosis_without_refusal(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    response(monkeypatch,b'',error=TimeoutError())
    for n in range(4):
        result=h.direct(record,state,url,URL,tmp_path/'work')
        assert result['status']=='DIRECT_UNAVAILABLE'
        assert result['diagnose_before_repeat'] is (n>=2)
    assert len(h._read(record)['actions'])==4

def test_browser_wrong_operation_action_rejected(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    file=make_ppt(directory/'generic.pptx')
    op=operation(tmp_path,identity,record,url,file);data=de.read(op)
    data['action']='recovery';op.write_text(json.dumps(data))
    monkeypatch.setattr(de.time,'sleep',lambda _:None)
    with pytest.raises(ValueError,match='operation capture'):h.observe(record,state,snap['snapshot']['path'],1,False,operation_capture=op)

def test_complete_pending_proof_reused_without_another_get(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    response(monkeypatch,pptbytes(tmp_path));first=h.direct(record,state,url,URL,tmp_path/'work')
    def denied(*a,**k):pytest.fail('duplicate network transfer')
    monkeypatch.setattr(dt,'download',denied)
    second=h.direct(record,state,url,URL,tmp_path/'work')
    assert second['reused'] and second['file_evidence']==first['file_evidence']

def test_text_plan_invalidation_clears_transport_policy(tmp_path):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    data=runtime.load_runtime_state(state)
    assert 'S001' in data['canva_bridge']['transport_requirements']
    b._invalidate_after_text_plan_change(data,'S001')
    assert 'S001' not in data['canva_bridge'].get('transport_requirements',{})

def test_explicit_termination_releases_owned_directory_and_requires_new_record(tmp_path):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    original=h._read(record)
    assert h.terminate(record,state)['status']=='TERMINATED'
    assert Path(str(record)+'.intent').exists()
    assert all(not Path(p).exists() for p in original['lease_paths'])
    with pytest.raises(ValueError,match='terminated'):h.direct(record,state,url,URL,tmp_path/'work')
    new=tmp_path/'new-record.json'
    h.prepare(new,state,'S001',original['request']['lookup_path'],tmp_path/'out','S001.pptx')
    h.snapshot(new,state,directory,url)
    assert runtime.load_runtime_state(state)['canva_bridge']['transport_requirements']['S001']['acquisition_id']==h._read(new)['acquisition_id']

def test_browser_unreadable_href_can_still_recover(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    recovery=h.snapshot(record,state,directory,url,'recovery')
    file=make_ppt(directory/'generic.pptx');op=operation(tmp_path,identity,record,url,file)
    data=de.read(op);data['action']='recovery';op.write_text(json.dumps(data))
    monkeypatch.setattr(de.time,'sleep',lambda _:None)
    assert h.observe(record,state,recovery['snapshot']['path'],1,False,operation_capture=op)['status']=='FILE_COMPLETE'

def test_terminal_export_retry_allowed_under_same_identity_lock(tmp_path,monkeypatch):
    state,identity,record,directory,url,snap=fixture(tmp_path)
    h.fail(record,'EXPORT_FAILED')
    retry=h.snapshot(record,state,directory,url,'export_retry')
    file=make_ppt(directory/'generic.pptx');op=operation(tmp_path,identity,record,url,file)
    monkeypatch.setattr(de.time,'sleep',lambda _:None)
    result=h.observe(record,state,retry['snapshot']['path'],1,False,operation_capture=op)
    assert result['status']=='FILE_COMPLETE' and Path(str(record)+'.intent').exists()
