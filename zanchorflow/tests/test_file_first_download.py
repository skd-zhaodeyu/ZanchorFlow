"""File-first behavior and source identity regressions; synthetic offline fixture only."""
import copy
import json
import shutil
import sys
from pathlib import Path
import pytest
from pptx import Presentation
from PIL import Image
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import canva_bridge as b
import host_acquisition as h
import download_evidence as de
import runtime
from test_acquisition_request import setup
from test_preflight_scope import _bridge_fixture

def fixture(tmp_path):
    state,identity,lookup=setup(tmp_path)
    lookup['observation']='synthetic trusted lookup; no live design call'
    title=de.title_for('S001',identity['attempt_id'],identity['text_clean_fingerprint'])
    lookup['response']['design']['title']=title
    lp=tmp_path/'lookup.json';lp.write_text(json.dumps(lookup))
    b.register_design_observation(state,'S001',lp)
    downloads=tmp_path/'browser';downloads.mkdir()
    record=tmp_path/'record.json'
    h.ready(record,{'status':'WAIT_DOWNLOAD','identity':identity,'edit_url':lookup['response']['design']['urls']['edit_url']})
    legacy=h._read(record);legacy.pop('require_transport_receipt');h._save(record,legacy)
    h.prepare(record,state,'S001',lp,tmp_path/'out','S001.pptx')
    url='https://www.canva.com/design/fixture-design/edit'
    snap=h.snapshot(record,state,downloads,url,lease_dir=tmp_path/'leases')
    h.lock(record)
    return state,identity,record,downloads,title,url,snap

def make_ppt(path,text=''):
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    if text:slide.shapes.add_textbox(1,1,100,100).text=text
    prs.save(path);return path

def capture(tmp_path,monkeypatch,recovery=False):
    state,identity,record,directory,title,url,snap=fixture(tmp_path)
    if recovery:
        snap=h.snapshot(record,state,directory,url,'recovery','https://export-download.canva.com/opaque/fixture-design/1/export.pptx',tmp_path/'leases')
    source=make_ppt(directory/(title+'.pptx'))
    monkeypatch.setattr(de.time,'sleep',lambda _:None)
    observed=h.observe(record,state,snap['snapshot']['path'],wait_seconds=1,event_confirmed=False)
    assert observed['status']=='FILE_COMPLETE'
    return state,identity,record,source,url,observed

def checked(tmp_path,monkeypatch,recovery=False):
    state,identity,record,source,url,observed=capture(tmp_path,monkeypatch,recovery)
    state_data=runtime.load_runtime_state(state);input_file=b.resolve(state_data,state_data['slides']['S001']['text_clean'])
    # Contract fixture only: synthetic receipt/review is labelled, no rendering assertion.
    png=tmp_path/'synthetic-preview.png';Image.new('RGB',(16,9),'red').save(png)
    preview={'schema_version':2,'identity':identity,'source_sha256':de.sha(source),
             'preview_path':str(png),'preview_sha256':de.sha(png),'fixture_only':True}
    pref=de.write_once(tmp_path/'synthetic-preview.json',preview)
    review={'identity':identity,'file_sha256':de.sha(source),'input_sha256':identity['text_clean_fingerprint'],
            'input_render':{'path':str(input_file),'sha256':de.sha(input_file)},'preview_receipt':pref,
            'assessment':'MATCH','evidence':'synthetic offline comparison contract; not a live visual PASS'}
    rp=tmp_path/'review.json';rp.write_text(json.dumps(review))
    final=h.register_page_check(record,state,observed['file_evidence'],rp)
    return state,identity,record,source,url,final

@pytest.mark.parametrize('recovery',[False,True])
def test_no_event_no_history_file_proof_finishes_and_binds(tmp_path,monkeypatch,recovery):
    state,identity,record,source,url,final=checked(tmp_path,monkeypatch,recovery)
    def finalize(source,out,name,python):
        target=Path(out)/name;target.parent.mkdir();shutil.move(source,target)
        return {'ValidPptx':True,'SlideCount':1,'Path':str(target),'SHA256':de.sha(target)}
    monkeypatch.setattr(h,'_finalize',finalize)
    result=h.finish(record,state,None,url,False,file_evidence=final['file_evidence'])
    assert result['status']=='DOWNLOAD_BOUND'
    current=runtime.load_runtime_state(state)
    bound=b.download_current(current,'S001')
    assert bound['identity']==identity
    evidence=json.loads(bound['evidence']['completion_evidence'])
    assert evidence['download_event_confirmed'] is False
    assert b.status(state)['status']=='RESTORE_TEXT'
    assert h.finish(record,state,None,url,False)['reused'] is True
    assert not Path(h._read(record)['lease_path']).exists()

@pytest.mark.parametrize('field',['attempt_id','slide_id','run_id','source_fingerprint','text_clean_fingerprint','approved_outline_fingerprint','design_id'])
def test_each_identity_field_mismatch_rejected(tmp_path,monkeypatch,field):
    state,identity,record,source,url,final=checked(tmp_path,monkeypatch)
    proof=de.read(final['file_evidence']);proof['identity'][field]='wrong'
    with pytest.raises(ValueError,match='IDENTITY_MISMATCH'):de.validate_file_proof(proof,identity)

def test_similar_single_page_from_other_design_name_not_selected(tmp_path,monkeypatch):
    state,identity,record,directory,title,url,snap=fixture(tmp_path)
    make_ppt(directory/'zf-S002-other-title.pptx')
    result=h.observe(record,state,snap['snapshot']['path'],wait_seconds=0)
    assert result['status']=='NO_FILE_STARTED'
    assert not h._read(record).get('file_proofs')

@pytest.mark.parametrize('href',[
 'https://export-download.canva.com/opaque/other-design/1/export.pptx',
 'https://export-download.canva.com/opaque/fixture-design-extra/1/export.pptx',
 'https://export-download.canva.com.evil.test/opaque/fixture-design/1/export.pptx',
 'http://export-download.canva.com/opaque/fixture-design/1/export.pptx'])
def test_wrong_recovery_design_or_host_rejected(tmp_path,href):
    state,identity,record,directory,title,url,snap=fixture(tmp_path)
    with pytest.raises(ValueError):h.snapshot(record,state,directory,url,'recovery',href)

def test_unchanged_old_exact_title_not_a_new_download(tmp_path,monkeypatch):
    state,identity,record,directory,title,url,snap=fixture(tmp_path)
    source=make_ppt(directory/(title+'.pptx'))
    recovery=h.snapshot(record,state,directory,url,'recovery','https://export-download.canva.com/opaque/fixture-design/1/export.pptx')
    assert h.observe(record,state,recovery['snapshot']['path'],wait_seconds=0)['status']=='NO_FILE_STARTED'

def test_overwritten_file_with_new_bytes_detected(tmp_path,monkeypatch):
    state,identity,record,directory,title,url,snap=fixture(tmp_path)
    source=make_ppt(directory/(title+'.pptx'),'before')
    recovery=h.snapshot(record,state,directory,url,'recovery','https://export-download.canva.com/opaque/fixture-design/1/export.pptx')
    make_ppt(source,'after');monkeypatch.setattr(de.time,'sleep',lambda _:None)
    result=h.observe(record,state,recovery['snapshot']['path'],wait_seconds=1)
    assert result['status']=='FILE_COMPLETE' and result['selected']['sha256']==de.sha(source)

def test_different_content_candidates_do_not_use_visual_similarity(tmp_path,monkeypatch):
    state,identity,record,directory,title,url,snap=fixture(tmp_path)
    make_ppt(directory/(title+'.pptx'),'A');make_ppt(directory/(title+' (1).pptx'),'B')
    monkeypatch.setattr(de.time,'sleep',lambda _:None)
    result=h.observe(record,state,snap['snapshot']['path'],wait_seconds=1)
    assert result['status']=='SOURCE_AMBIGUOUS' and len(result['candidates'])==2

def test_byte_identical_duplicates_are_both_logged(tmp_path,monkeypatch):
    state,identity,record,directory,title,url,snap=fixture(tmp_path)
    source=make_ppt(directory/(title+'.pptx'));shutil.copyfile(source,directory/(title+' (1).pptx'))
    monkeypatch.setattr(de.time,'sleep',lambda _:None)
    result=h.observe(record,state,snap['snapshot']['path'],wait_seconds=1)
    assert result['status']=='FILE_COMPLETE' and len(result['duplicates'])==1

def test_bound_evidence_tamper_invalidates_download_current(tmp_path,monkeypatch):
    state,identity,record,source,url,final=checked(tmp_path,monkeypatch)
    h.acquire_files(record,state,final['file_evidence'])
    data=h._read(record);payload=data['download_evidence']
    envelope={k:identity[k] for k in ('attempt_id','slide_id','design_id')}
    envelope.update(edit_url=url,download_entry='synthetic browser',completion_evidence=json.dumps(payload),pptx_sha256=de.sha(source))
    ep=tmp_path/'binding.json';ep.write_text(json.dumps(envelope));b.bind_download(state,'S001',source,ep)
    snap=Path(payload['snapshot']['path']);snap.write_text('{}')
    with pytest.raises(ValueError,match='EVIDENCE_CHANGED'):b.download_current(runtime.load_runtime_state(state),'S001')

def test_missing_page_review_retains_actual_file(tmp_path,monkeypatch):
    state,identity,record,source,url,observed=capture(tmp_path,monkeypatch)
    with pytest.raises(ValueError,match='PAGE_CHECK_REQUIRED'):h.acquire_files(record,state,observed['file_evidence'])
    assert source.exists() and Path(str(record)+'.intent').exists()

def test_directory_lease_conflict_does_not_unlock_original(tmp_path):
    directory=tmp_path/'browser';directory.mkdir();leases=tmp_path/'leases'
    path=de.acquire_lease(directory,leases,'first')
    with pytest.raises(ValueError,match='IN_USE'):de.acquire_lease(directory,leases,'second')
    with pytest.raises(ValueError):de.release_lease(path,'second')
    assert Path(path).exists();de.release_lease(path,'first')

def test_upload_registration_captures_actual_current_image_before_return_id(tmp_path):
    state=_bridge_fixture(tmp_path,n=1);pending=b.begin_attempt(state,'S001')
    data=runtime.load_runtime_state(state);image=b.resolve(data,data['slides']['S001']['text_clean'])
    result=b.register_upload(state,'S001',image)
    assert result['image_sha256']==pending['text_clean_fingerprint']
    assert pending['attempt_id'] in result['requested_title']
    assert 'design_id' not in result['identity']
    assert b.register_upload(state,'S001',image)==result

def test_upload_registration_rejects_another_similar_image_path(tmp_path):
    state=_bridge_fixture(tmp_path,n=1);b.begin_attempt(state,'S001')
    data=runtime.load_runtime_state(state);image=b.resolve(data,data['slides']['S001']['text_clean'])
    other=tmp_path/'other.png';shutil.copyfile(image,other)
    with pytest.raises(ValueError,match='actual upload'):b.register_upload(state,'S001',other)

def test_design_observation_never_rewrites_registered_attempt(tmp_path):
    state,identity,record,directory,title,url,snap=fixture(tmp_path)
    data=runtime.load_runtime_state(state);p=b.resolve(data,data['slides']['S001']['reconstruction_attempt'])
    before=p.read_bytes();lookup=h._read(record)['request']['lookup_path']
    b.register_design_observation(state,'S001',lookup)
    assert p.read_bytes()==before and b.active_identity(runtime.load_runtime_state(state),'S001')==identity


def test_maintenance_does_not_modify_unrelated_existing_functions():
    import ast,base64
    from download_scope import data
    root=Path(__file__).resolve().parents[1]
    allowed={'canva_bridge.py':{'start_run','_invalidate_after_text_plan_change','begin_attempt',
                               'download_current','bind_download','accept_attempt','main'},
             'host_acquisition.py':{'_history','acquire','acquire_wait','finish','main','ready','reserve_retry','_strategy_record'}}
    new_allowed={'canva_bridge.py':{'register_upload','register_upload_result','register_design_observation','validate_completion_evidence','register_transport_requirement'},
                 'host_acquisition.py':{'snapshot','observe','acquire_files','register_page_check','_register_transport_proof','direct','terminate'}}
    for name,changed in allowed.items():
        entry=data()['zanchorflow/scripts/'+name]
        baseline=ast.parse(base64.b64decode(entry['baseline_b64']).decode())
        current=ast.parse((root/'scripts'/name).read_text(encoding='utf-8'))
        old={n.name:ast.dump(n) for n in baseline.body if isinstance(n,ast.FunctionDef)}
        now={n.name:ast.dump(n) for n in current.body if isinstance(n,ast.FunctionDef)}
        assert set(now)-set(old)==new_allowed[name]
        assert not set(old)-set(now)
        for function in set(old)-changed:assert old[function]==now[function],function


def test_scope_detects_unapproved_candidate_mutations():
    from download_scope import data,restore_repo_bytes,new_runtime,verify_new_runtime
    root=Path(__file__).resolve().parents[1]
    for rel in data():
        with pytest.raises(AssertionError):restore_repo_bytes(rel,(root.parent/rel).read_bytes()+b'\n# unapproved\n')
    for rel in new_runtime():
        with pytest.raises(AssertionError):verify_new_runtime(rel,(root/rel).read_bytes()+b'\n# unapproved\n')


@pytest.mark.parametrize('operation',['normal','recovery'])
def test_normalized_filename_uses_current_design_export_not_title(tmp_path,monkeypatch,operation):
    from urllib.parse import quote
    state,identity,record,directory,title,url,snap=fixture(tmp_path)
    header="attachment; filename*=UTF-8''"+quote('exported-page.pptx')
    href='https://export-download.canva.com/opaque/fixture-design/1/export.pptx?response-content-disposition='+quote(header)
    if operation=='recovery':snap=h.snapshot(record,state,directory,url,'recovery',href)
    make_ppt(directory/'exported-page.pptx')
    monkeypatch.setattr(de.time,'sleep',lambda _:None)
    result=h.observe(record,state,snap['snapshot']['path'],1,False,href)
    assert result['status']=='FILE_COMPLETE'
    proof=de.read(result['file_evidence']);assert de.validate_file_proof(proof,identity)['identity']==identity
    assert proof['export_observation']['filename']=='exported-page.pptx'
    assert 'response-content-disposition' not in json.dumps(proof)

def test_unknown_completion_schema_does_not_bypass_validation(tmp_path):
    state,identity,_=setup(tmp_path)
    with pytest.raises(ValueError,match='unsupported completion'):b.validate_completion_evidence(runtime.load_runtime_state(state),identity,{'schema_version':3},tmp_path/'unused')

def test_two_current_pages_cannot_claim_one_design_id(tmp_path):
    from test_preflight_scope import _accepted_attempt
    state=_bridge_fixture(tmp_path,n=2)
    _accepted_attempt(state,'S001',tmp_path,design='duplicate-design')
    identity=_accepted_attempt(state,'S002',tmp_path,design='duplicate-design')
    lookup={'schema_version':1,'source':'canva_connector.get_design','identity':identity,'observation':'synthetic',
            'response':{'design':{'id':'duplicate-design','title':de.title_for('S002',identity['attempt_id'],identity['text_clean_fingerprint']),
                                 'page_count':1,'urls':{'edit_url':'https://www.canva.com/design/duplicate-design/edit'}}}}
    path=tmp_path/'lookup.json';path.write_text(json.dumps(lookup))
    with pytest.raises(ValueError,match='multiple current pages'):b.register_design_observation(state,'S002',path)


def test_generic_file_with_only_visual_claim_is_not_identity_proof(tmp_path):
    identity={'attempt_id':'a'*32}
    source=make_ppt(tmp_path/'generic.pptx')
    proof={'selected':{'path':str(source)},'download_event_confirmed':False}
    with pytest.raises(ValueError,match='SOURCE_ASSOCIATION_PENDING'):de.source_association(proof,identity,{'actual_title':'unique-design'},source)

def test_normalized_file_with_current_unique_core_title_associates(tmp_path):
    identity={'attempt_id':'a'*32};source=tmp_path/'generic.pptx'
    prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.core_properties.title='unique-design';prs.save(source)
    proof={'selected':{'path':str(source)},'download_event_confirmed':False}
    assert de.source_association(proof,identity,{'actual_title':'unique-design'},source)=='unique_design_core_title'


def test_missing_api_title_uses_captured_same_design_host_observation(tmp_path):
    state,identity,lookup=setup(tmp_path)
    lookup['observation']='synthetic API response without optional title'
    lp=tmp_path/'lookup.json';lp.write_text(json.dumps(lookup))
    title=de.title_for('S001',identity['attempt_id'],identity['text_clean_fingerprint'])
    url='https://www.canva.com/design/fixture-design/edit'
    capture=de.write_once(tmp_path/'host-capture.json',{'observed_url':url,'actual_title':title,'fixture_only':True})
    hp=tmp_path/'host-title.json';hp.write_text(json.dumps({'source':'canva_host.readonly','identity':identity,
        'observed_url':url,'actual_title':title,'evidence':'synthetic actual name-control observation','capture':capture}))
    assert b.register_design_observation(state,'S001',lp,hp)['actual_title']==title
    data=runtime.load_runtime_state(state);obs=de.sealed(data['canva_bridge']['design_observations']['S001'])
    assert de.sealed(obs['host_title'])['identity']==identity
    wrong=de.read(hp);wrong['observed_url']='https://www.canva.com/design/other-design/edit';hp.write_text(json.dumps(wrong))
    with pytest.raises(ValueError):b.register_design_observation(state,'S001',lp,hp)


def test_same_design_url_token_change_can_resume_without_reexport(tmp_path,monkeypatch):
    state,identity,record,source,url,final=checked(tmp_path,monkeypatch)
    def finalize(source,out,name,python):
        target=Path(out)/name;target.parent.mkdir();shutil.move(source,target)
        return {'ValidPptx':True,'SlideCount':1,'Path':str(target),'SHA256':de.sha(target)}
    monkeypatch.setattr(h,'_finalize',finalize)
    result=h.finish(record,state,None,url+'?ui-token=rotated',False,file_evidence=final['file_evidence'])
    assert result['status']=='DOWNLOAD_BOUND'
    assert b.download_current(runtime.load_runtime_state(state),'S001')['identity']==identity


def test_observe_cli_defaults_to_ten_seconds_not_legacy_finish_wait(tmp_path,monkeypatch,capsys):
    calls=[]
    def observed(*args):calls.append(args[3]);return {'status':'NO_FILE_STARTED'}
    monkeypatch.setattr(h,'observe',observed)
    monkeypatch.setattr(sys,'argv',['host_acquisition.py','observe','--record',str(tmp_path/'record.json'),
        '--state',str(tmp_path/'state.json'),'--snapshot',str(tmp_path/'snapshot.json')])
    assert h.main()==0 and calls==[10]


def test_acceptance_checks_actual_magic_return_independently_of_proposed_id(tmp_path):
    state=_bridge_fixture(tmp_path,n=1);active=b.begin_attempt(state,'S001')
    data=runtime.load_runtime_state(state);image=b.resolve(data,data['slides']['S001']['text_clean'])
    upload=b.register_upload(state,'S001',image)
    raw={'tool':'canva_image_to_design','call_args':{'image_file':str(image),'title':upload['requested_title']},
         'result':{'content':[{'type':'text','text':json.dumps({'design':{'id':'actual-returned-design'}})}]}}
    rp=tmp_path/'raw-tool.json';rp.write_text(json.dumps(raw))
    assert b.register_upload_result(state,'S001',rp)['design_id']=='actual-returned-design'
    proposed={**active,'design_id':'wrong-but-similar-design','assessment':'ACCEPT','evidence':'synthetic assessment'}
    ap=tmp_path/'accept.json';ap.write_text(json.dumps(proposed))
    with pytest.raises(ValueError,match='actual Magic return'):b.accept_attempt(state,'S001',active['attempt_id'],ap)
    proposed['design_id']='actual-returned-design';ap.write_text(json.dumps(proposed))
    assert b.accept_attempt(state,'S001',active['attempt_id'],ap)['design_id']=='actual-returned-design'

@pytest.mark.parametrize('result',[{}, {'design_id':'a','design':{'id':'b'}}, {'isError':True,'design_id':'a'}])
def test_ambiguous_or_error_magic_result_is_not_guessed(result):
    with pytest.raises(ValueError):de.returned_design_id(result)


def test_preview_keeps_existing_user_presentations_open(tmp_path,monkeypatch):
    import types,download_preview
    state,identity,_=setup(tmp_path);source=make_ppt(tmp_path/'source.pptx');before=de.sha(source)
    calls=[]
    class Slide:
        def Export(self,path,format):Image.new('RGB',(16,9),'red').save(path)
    doc=types.SimpleNamespace(Slides=types.SimpleNamespace(Count=1,Item=lambda _:Slide()),Close=lambda:calls.append('own_doc_closed'))
    presentations=types.SimpleNamespace(Count=1,Open=lambda *_:doc)
    app=types.SimpleNamespace(Presentations=presentations,Quit=lambda:calls.append('app_quit'))
    client=types.ModuleType('win32com.client');client.DispatchEx=lambda _:app
    package=types.ModuleType('win32com');package.client=client
    monkeypatch.setitem(sys.modules,'win32com',package);monkeypatch.setitem(sys.modules,'win32com.client',client)
    result=download_preview.render(source,tmp_path/'preview',state,'S001','powerpoint')
    assert calls==['own_doc_closed'] and result['owned_application'] is False
    assert de.sha(source)==before and Path(result['preview_path']).is_file()
    assert package.__gen_path__==str(tmp_path/'preview'/'com-cache')
