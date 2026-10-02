import copy
import json
import sys
from pathlib import Path
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import canva_bridge as b


def canvas(tw=1000000,th=1000000):
    return {'source_width':1000,'source_height':1000,
            'target_slide_width_emu':tw,'target_slide_height_emu':th}


def graphic(shapes):
    shape=shapes.add_shape(MSO_SHAPE.RECTANGLE,100000,100000,300000,200000)
    shape.fill.solid(); shape.line.fill.background()
    style=shape._element.find(qn('p:style'))
    if style is not None: shape._element.remove(style)
    return shape


def source(tmp_path,height=2000000,group=False,text=False):
    p=tmp_path/'download.pptx';prs=Presentation();prs.slide_width=2000000;prs.slide_height=height
    sl=prs.slides.add_slide(prs.slide_layouts[6])
    shapes=sl.shapes.add_group_shape().shapes if group else sl.shapes
    g=graphic(shapes)
    t=None
    if text:
        t=sl.shapes.add_textbox(10000,10000,200000,100000);t.text='ordinary source label'
    prs.save(p)
    return p,g.shape_id,t.shape_id if t else None


def registration(p):
    keys=('same_framing','no_crop','no_expansion','no_rotation','no_relayout')
    return {'download_sha256':b.sha(p),'text_clean_fingerprint':'synthetic-clean-fingerprint',
            'checks':{k:'PASS' for k in keys},'evidence':{k:'synthetic controlled registration fixture' for k in keys}}


def restored(tmp_path,p,can,reg=None,cleanup=None):
    prs,report=b.prepare_canvas_mapping(p,can,reg or registration(p),cleanup)
    t=prs.slides[0].shapes.add_textbox(10000,10000,200000,100000);t.text='restored native label'
    out=tmp_path/'restored.pptx';prs.save(out)
    report['restored_pptx_sha256']=b.sha(out)
    return out,report


def test_original_five_pixel_path_does_not_move_graphics(tmp_path):
    p,_,_=source(tmp_path,height=2010000)
    before=Presentation(p).slides[0].shapes[0]
    prs,report=b.prepare_canvas_mapping(p,canvas(2000000,2000000))
    assert report['plan']['mode']=='unchanged_objects'
    assert etree.tostring(before._element)==etree.tostring(prs.slides[0].shapes[0]._element)
    assert [prs.slide_width,prs.slide_height]==[2000000,2000000]


@pytest.mark.parametrize('height,pass_expected',[(2009998,True),(2010000,True),(2010002,False)])
def test_adaptation_five_source_pixel_boundary(tmp_path,height,pass_expected):
    p,_,_=source(tmp_path,height=height)
    if not pass_expected:
        with pytest.raises(ValueError,match='DECK_CANVAS_INVARIANT_FAILED'):
            b.prepare_canvas_mapping(p,canvas(),registration(p))
    else:
        out,report=restored(tmp_path,p,canvas())
        assert report['plan']['mode']=='uniform_fit'
        b.verify_canvas_mapping(p,out,canvas(),report,'synthetic-clean-fingerprint')


def test_group_child_transform_is_not_scaled_twice(tmp_path):
    p,_,_=source(tmp_path,group=True)
    original=Presentation(p).slides[0].shapes[0].shapes[0]._element
    prs,report=b.prepare_canvas_mapping(p,canvas(),registration(p))
    assert etree.tostring(original)==etree.tostring(prs.slides[0].shapes[0].shapes[0]._element)
    out=tmp_path/'group.pptx';prs.save(out);report['restored_pptx_sha256']=b.sha(out)
    b.verify_canvas_mapping(p,out,canvas(),report)


def test_text_bbox_and_glyph_height_use_same_letterboxed_rectangle(tmp_path):
    p,_,_=source(tmp_path,height=2010000)
    _,report=b.prepare_canvas_mapping(p,canvas(),registration(p))
    geometry=b.map_native_text_geometry({'visual_bbox_normalized':[0,0,1,1],'font_size_normalized':0.1},report['plan'])
    rect=report['plan']['content_rect_emu']
    assert geometry['bbox_emu']==[round(v) for v in rect]
    assert geometry['visible_glyph_height_emu']==pytest.approx(rect[3]*0.1)
    assert geometry['visible_glyph_height_pt']==pytest.approx(rect[3]*0.1/12700)


@pytest.mark.parametrize('field,value',[('source_width',0),('source_height',float('nan')),('target_slide_width_emu',-1),('target_slide_height_emu',1.5)])
def test_invalid_dimensions_fail_closed(tmp_path,field,value):
    p,_,_=source(tmp_path); can=canvas();can[field]=value
    with pytest.raises(ValueError,match='CANVAS_MAPPING_UNRESOLVED'):
        b.prepare_canvas_mapping(p,can,registration(p))


@pytest.mark.parametrize('bad',['absent','crop','unbound'])
def test_framing_cannot_be_inferred_from_aspect(tmp_path,bad):
    p,_,_=source(tmp_path);reg=registration(p)
    if bad=='absent':reg=None
    if bad=='crop':reg['checks']['no_crop']='FAIL'
    if bad=='unbound':reg['download_sha256']='wrong'
    with pytest.raises(ValueError,match='CANVAS_MAPPING_UNRESOLVED'):
        b.prepare_canvas_mapping(p,canvas(),reg)


def test_safe_neutrality_cleanup_and_forged_graphic_cleanup(tmp_path):
    p,g,t=source(tmp_path,text=True)
    clean={'download_sha256':b.sha(p),'removed_text_shape_ids':[t],'evidence':'synthetic ordinary-text-only neutrality decision'}
    out,report=restored(tmp_path,p,canvas(),cleanup=clean)
    assert report['retained_shape_ids']==[g]
    b.verify_canvas_mapping(p,out,canvas(),report)
    clean['removed_text_shape_ids']=[g]
    with pytest.raises(ValueError,match='cannot remove non-text graphics'):
        b.prepare_canvas_mapping(p,canvas(),registration(p),clean)


@pytest.mark.parametrize('mutation',['remove','reorder','path','native_geometry','target','new_graphic','report'])
def test_preservation_tampering_is_rejected(tmp_path,mutation):
    p,_,_=source(tmp_path);before=b.sha(p)
    out,report=restored(tmp_path,p,canvas())
    prs=Presentation(out);sl=prs.slides[0]
    if mutation=='remove':sl.shapes._spTree.remove(sl.shapes[0]._element)
    elif mutation=='reorder':sl.shapes._spTree.remove(sl.shapes[0]._element);sl.shapes._spTree.append(Presentation(p).slides[0].shapes[0]._element)
    elif mutation=='path':sl.shapes[0]._element.find('.//'+qn('a:prstGeom')).set('prst','ellipse')
    elif mutation=='native_geometry':sl.shapes[0].left+=1
    elif mutation=='target':prs.slide_width+=1
    elif mutation=='new_graphic':graphic(sl.shapes)
    elif mutation=='report':report['plan']['scale_ratio']=[1,3]
    prs.save(out);report['restored_pptx_sha256']=b.sha(out)
    with pytest.raises(ValueError,match='CANVAS_MAPPING_UNRESOLVED|DECK_CANVAS_INVARIANT_FAILED'):
        b.verify_canvas_mapping(p,out,canvas(),report)
    assert b.sha(p)==before


@pytest.mark.parametrize('unsupported',['style','line','table','text'])
def test_whole_page_preflight_rejects_unsupported_objects(tmp_path,unsupported):
    p,_,_=source(tmp_path);prs=Presentation(p);sl=prs.slides[0]
    if unsupported=='style':
        sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,100000,100000)
    elif unsupported=='line':sl.shapes[0].line.fill.solid()
    elif unsupported=='table':sl.shapes.add_table(1,1,0,0,100000,100000)
    else:sl.shapes.add_textbox(0,0,100000,100000).text='unresolved ordinary text'
    prs.save(p);before=b.sha(p)
    with pytest.raises(ValueError,match='CANVAS_MAPPING_UNRESOLVED|GRAPHICS_FIRST_TEXT_CONTAMINATION'):
        b.prepare_canvas_mapping(p,canvas(),registration(p))
    assert b.sha(p)==before


def sealed_fixture(tmp_path,ordinary_text=False,original_tolerance=False):
    import runtime as r
    from test_preflight_scope import _bridge_fixture, _accepted_attempt
    state=_bridge_fixture(tmp_path,n=1)
    can={'source_width':16,'source_height':9,'target_slide_width_emu':6096000,'target_slide_height_emu':3429000}
    cp=tmp_path/'new-canvas.json';cp.write_text(json.dumps(can))
    r.register_stage2_artifact(state,'canvas:S001',cp)
    r.mark_stage2_text_reconciled(state,'S001')
    r.present_stage2_formal_display(state);r.seal_stage2_visual_approval(state,'synthetic fixture explicit visual approval')
    b.register_text_plan(state,'S001',tmp_path/'S001-manifest.json',tmp_path/'S001-inventory.json')
    b.register_text_clean(state,'S001',tmp_path/'S001-text_clean.png')
    identity=_accepted_attempt(state,'S001',tmp_path)
    p=tmp_path/'download.pptx';prs=Presentation();prs.slide_width=can['target_slide_width_emu'] if original_tolerance else 15925800;prs.slide_height=can['target_slide_height_emu'] if original_tolerance else 8953500
    sl=prs.slides.add_slide(prs.slide_layouts[6]);graphic(sl.shapes)
    text_id=None
    if ordinary_text:
        tx=sl.shapes.add_textbox(10000,10000,200000,100000);tx.text='A';text_id=tx.shape_id
    prs.save(p)
    ev={**{k:identity[k] for k in ('attempt_id','slide_id','design_id')},
        'edit_url':'https://www.canva.com/design/design-A/edit','download_entry':'synthetic PowerPoint entry',
        'completion_evidence':'synthetic exact Host/History completed-file fixture','pptx_sha256':b.sha(p)}
    ep=tmp_path/'download-evidence.json';ep.write_text(json.dumps(ev));b.bind_download(state,'S001',p,ep)
    reg=registration(p);reg['text_clean_fingerprint']=identity['text_clean_fingerprint']
    cleanup={'download_sha256':b.sha(p),'removed_text_shape_ids':[text_id],'evidence':'synthetic safe ordinary-text cleanup'} if ordinary_text else None
    out,report=restored(tmp_path,p,can,reg,cleanup)
    st=r.load_runtime_state(state)
    review={**{k:identity[k] for k in ('attempt_id','slide_id','design_id')},
        'download_sha256':b.sha(p),'restored_pptx_sha256':b.sha(out),
        'text_restore_fingerprint':r.expected_lineage_from_state(st,'S001')['text_restore_fingerprint'],
        'checks':{k:'PASS' for k in r.GATE_NAMES},'evidence':{k:'synthetic offline gate fixture' for k in r.GATE_NAMES},
        'canvas_mapping':report}
    rp=tmp_path/'page-review.json';rp.write_text(json.dumps(review))
    return state,p,out,can,review,rp


def test_formal_seal_requires_mapping_and_preserves_state_on_failure(tmp_path):
    state,p,out,can,review,rp=sealed_fixture(tmp_path)
    original=state.read_bytes();review.pop('canvas_mapping');rp.write_text(json.dumps(review))
    with pytest.raises(ValueError,match='uniform-fit review required'):
        b.seal_page(state,'S001',out,rp)
    assert state.read_bytes()==original and not Path(str(state)+'.bridge-lock').exists()


def test_formal_seal_stores_mapping_review_and_detects_changed_proof(tmp_path):
    import runtime as r
    state,p,out,can,review,rp=sealed_fixture(tmp_path)
    assert b.seal_page(state,'S001',out,rp)['status']=='PAGE_SEALED'
    st=r.load_runtime_state(state);record=b.page_provenance(st,'S001')
    assert record['canvas_mapping_review']['sha256']==b.sha(rp)
    assert b.sha(p)==review['download_sha256']
    rp.write_text(rp.read_text()+' ')
    with pytest.raises(ValueError,match='sealed mapping review changed'):
        b.page_provenance(r.load_runtime_state(state),'S001')


def test_unknown_extension_is_rejected_before_any_transform(tmp_path):
    p,_,_=source(tmp_path);prs=Presentation(p)
    from pptx.oxml.xmlchemy import OxmlElement
    props=prs.slides[0].shapes[0]._element.find(qn('p:spPr'))
    props.append(OxmlElement('a:extLst'));prs.save(p);before=b.sha(p)
    with pytest.raises(ValueError,match='extensions'):
        b.prepare_canvas_mapping(p,canvas(),registration(p))
    assert b.sha(p)==before


def test_picture_media_preserved_and_modified_media_rejected(tmp_path):
    import zipfile
    from PIL import Image
    p,_,_=source(tmp_path);image=tmp_path/'image.png';Image.new('RGB',(8,8),'red').save(image)
    prs=Presentation(p);prs.slides[0].shapes.add_picture(str(image),0,0,width=200000);prs.save(p)
    out,report=restored(tmp_path,p,canvas());b.verify_canvas_mapping(p,out,canvas(),report)
    altered=tmp_path/'altered.pptx'
    with zipfile.ZipFile(out) as old,zipfile.ZipFile(altered,'w') as new:
        for name in old.namelist():
            data=old.read(name)
            if name.startswith('ppt/media/'):data+=b'changed'
            new.writestr(name,data)
    report['restored_pptx_sha256']=b.sha(altered)
    with pytest.raises(ValueError,match='media or slide relationships changed'):
        b.verify_canvas_mapping(p,altered,canvas(),report)


def test_cleanup_cannot_hide_text_extension_objects(tmp_path):
    from pptx.oxml.xmlchemy import OxmlElement
    p,g,t=source(tmp_path,text=True);prs=Presentation(p)
    prs.slides[0].shapes[-1]._element.append(OxmlElement('p:extLst'));prs.save(p)
    cleanup={'download_sha256':b.sha(p),'removed_text_shape_ids':[t],'evidence':'synthetic text-only claim'}
    with pytest.raises(ValueError,match='cannot remove non-text graphics'):
        b.prepare_canvas_mapping(p,canvas(),registration(p),cleanup)


@pytest.mark.parametrize('invalid',['0','invalid'])
def test_invalid_group_child_extent_is_rejected(tmp_path,invalid):
    p,_,_=source(tmp_path,group=True);prs=Presentation(p)
    prs.slides[0].shapes[0]._element.find('.//'+qn('a:chExt')).set('cx',invalid);prs.save(p)
    with pytest.raises(ValueError,match='group child'):
        b.prepare_canvas_mapping(p,canvas(),registration(p))


@pytest.mark.parametrize('actual_width',[2000000,3000000])
def test_coordinate_rounding_has_at_most_one_emu_error(tmp_path,actual_width):
    from fractions import Fraction
    p,_,_=source(tmp_path);prs=Presentation(p)
    prs.slide_width=prs.slide_height=actual_width
    shape=prs.slides[0].shapes[0]
    shape.left=100001;shape.top=-100003;shape.width=300001;shape.height=200003
    original=[shape.left,shape.top,shape.width,shape.height];prs.save(p)
    mapped,report=b.prepare_canvas_mapping(p,canvas(),registration(p))
    result=mapped.slides[0].shapes[0]
    written=[result.left,result.top,result.width,result.height]
    scale=Fraction(*report['plan']['scale_ratio'])
    for old,new in zip(original,written):
        ideal=old*scale
        assert abs(new-ideal)<=1
        assert new==round(ideal)
    if actual_width==2000000:
        assert written[:2]==[50000,-50002]  # Both signed half-EMU cases.


@pytest.mark.parametrize('target_height,accepted',[(1004999,True),(1005000,True),(1005001,False)])
def test_target_to_source_aspect_is_checked_independently(tmp_path,target_height,accepted):
    p,_,_=source(tmp_path);can=canvas(th=target_height)
    if accepted:
        out,report=restored(tmp_path,p,can)
        b.verify_canvas_mapping(p,out,can,report)
    else:
        with pytest.raises(ValueError,match='DECK_CANVAS_INVARIANT_FAILED'):
            b.prepare_canvas_mapping(p,can,registration(p))
