import sys,copy,zipfile
from pathlib import Path
import pytest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from editable_build import build
from pptx import Presentation

def plan():
 return {'schema_version':1,'source_size_px':[1000,600],'slide_size_emu':[10000000,6000000],'objects':[{'id':'title','type':'text','box_px':[40,20,700,80],'text':'中文 Title','occurrence_ids':['t1'],'font_size':28},{'id':'table','type':'table','box_px':[40,150,700,300],'column_widths':[2,1],'row_heights':[1,2],'cells':[[{'text':'Header','col_span':2,'occurrence_ids':['t2']},{}],[{'text':'Alpha','occurrence_ids':['t3']},{'text':'42','occurrence_ids':['t4']}]]}]}

def test_native_table_and_text_mapping(tmp_path):
 p=plan();before=copy.deepcopy(p);result=build(p,tmp_path/'page.pptx');assert p==before
 prs=Presentation(result['path']);tb=next(s.table for s in prs.slides[0].shapes if s.has_table)
 assert tb.cell(0,0).is_merge_origin and tb.cell(0,1).is_spanned
 assert tb.cell(1,1).text=='42' and set(result['occurrences'])=={'t1','t2','t3','t4'}
 assert result['occurrences']['t4']['cell']==[1,1]

@pytest.mark.parametrize('fault',['duplicate','missing_geometry','bad_merge','bool_coordinate'])
def test_invalid_plan_rejected_before_output(tmp_path,fault):
 p=plan()
 if fault=='duplicate':p['objects'][1]['cells'][1][1]['occurrence_ids']=['t1']
 if fault=='missing_geometry':del p['objects'][0]['box_px']
 if fault=='bad_merge':p['objects'][1]['cells'][0][1]={'text':'hidden'}
 if fault=='bool_coordinate':p['objects'][0]['box_px'][0]=True
 out=tmp_path/'bad.pptx'
 with pytest.raises(ValueError):build(p,out)
 assert not out.exists()

def test_one_native_curve_with_dash_and_arrow(tmp_path):
 p=plan();p['objects']=[{'id':'curve','type':'path','box_px':[100,100,400,200],'fill':'none','dash':'dash','end_arrow':'triangle','path_px':[{'op':'moveTo','points':[[100,300]]},{'op':'cubicBezTo','points':[[200,100],[400,100],[500,300]]}]}]
 result=build(p,tmp_path/'curve.pptx')
 with zipfile.ZipFile(result['path']) as z:
  xml=z.read('ppt/slides/slide1.xml')
  assert xml.count(b'<p:sp>')==1 and b'cubicBezTo' in xml and b'prstDash' in xml and b'tailEnd' in xml
