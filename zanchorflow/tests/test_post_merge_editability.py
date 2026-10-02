"""Windows/PowerPoint integration: real edits, save, close and reopen."""
import json, sys, zipfile
from pathlib import Path
import pytest
from openpyxl import Workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from merge_pptx import merge
from runtime import seal_validated_single_page


def _make_merged(tmp_path):
    workbook=tmp_path/'independent.xlsx'
    wb=Workbook();ws=wb.active;ws['A1']='separate workbook';ws['B2']=41;wb.save(workbook)
    pptx=tmp_path/'source.pptx'
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    text=slide.shapes.add_textbox(Inches(0.5),Inches(0.5),Inches(3),Inches(0.8));text.text='Native text'
    data=CategoryChartData();data.categories=['A','B'];data.add_series('Business', (2,3))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1),Inches(2),Inches(4),Inches(3),data)
    slide.shapes.add_ole_object(str(workbook),'Excel.Sheet.12',Inches(6),Inches(2),Inches(2),Inches(2))
    prs.save(pptx)
    def write(name,value):
        p=tmp_path/name
        if isinstance(value,bytes):p.write_bytes(value)
        else:p.write_text(json.dumps(value),encoding='utf-8')
        return str(p)
    state={'deck_order':['S01'], 'slides':{'S01':{
        'approved_render':write('render.png',b'approved'),
        'final_content_truth':write('truth.json',{'text':'Native text'}),
        'canvas':write('canvas.json',{'width':10,'height':7.5}),
        'text_clean':write('clean.png',b'clean'),
        'finalized_manifest':write('manifest.json',{'text':'Native text'}),
        'font_fallback':write('fonts.json',{}),
    }}}
    state_path=tmp_path/'runtime-state.json';state_path.write_text(json.dumps(state),encoding='utf-8')
    record=seal_validated_single_page(state_path,'S01',pptx,{k:'PASS' for k in ('content_truth','semantic_fidelity','functional_editability','visual_fidelity')})
    merged=tmp_path/'merged.pptx';result=merge([record],['S01'],merged,state_path)
    assert result['status']=='PASS',result
    with zipfile.ZipFile(merged) as z:
        embeddings=[n for n in z.namelist() if n.startswith('ppt/embeddings/')]
        assert len(embeddings)==2, embeddings
        assert any(n.endswith('.xlsx') for n in embeddings)
        assert any(n.endswith('.bin') for n in embeddings)
    return merged


@pytest.mark.integration
def test_powerpoint_post_merge_editability(tmp_path):
    win32=pytest.importorskip('win32com.client')
    merged=_make_merged(tmp_path)
    app=win32.DispatchEx('PowerPoint.Application'); app.Visible=True; prs=None
    try:
        prs=app.Presentations.Open(str(merged),0,0,-1)
        slide=prs.Slides(1)
        text=next(s for s in slide.Shapes if s.HasTextFrame and s.TextFrame.HasText and s.TextFrame.TextRange.Text=='Native text')
        chart=next(s.Chart for s in slide.Shapes if s.HasChart)
        ole=next(s for s in slide.Shapes if s.Type==7)
        text.TextFrame.TextRange.Text='Text changed after merge'
        chart.ChartData.Activate()
        chart_wb=chart.ChartData.Workbook
        chart_wb.Worksheets(1).Cells(2,2).Value=9
        chart_wb.Close(True)
        app.ActiveWindow.ViewType=1
        ole.OLEFormat.Activate()
        independent=ole.OLEFormat.Object
        independent.Worksheets(1).Cells(2,2).Value=73
        independent.Save()
        prs.Save();prs.Close();prs=None
        prs=app.Presentations.Open(str(merged),0,0,-1)
        slide=prs.Slides(1)
        assert any(s.HasTextFrame and s.TextFrame.HasText and s.TextFrame.TextRange.Text=='Text changed after merge' for s in slide.Shapes)
        chart=next(s.Chart for s in slide.Shapes if s.HasChart)
        chart.ChartData.Activate();chart_wb=chart.ChartData.Workbook
        assert chart_wb.Worksheets(1).Cells(2,2).Value==9
        assert list(chart.SeriesCollection(1).Values)[0]==9
        chart_wb.Close(False)
        ole=next(s for s in slide.Shapes if s.Type==7)
        app.ActiveWindow.ViewType=1
        app.ActiveWindow.ViewType=1
        ole.OLEFormat.Activate();independent=ole.OLEFormat.Object
        assert independent.Worksheets(1).Cells(2,2).Value==73
    finally:
        if prs is not None:prs.Close()
        app.Quit()
