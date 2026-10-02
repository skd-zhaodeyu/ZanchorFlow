import sys, zipfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from runtime import active_attempt_matches, validate_manifest_hashes, text_clean_fingerprint, text_restore_fingerprint, resume_stage
from merge_pptx import inspect_ooxml_dependencies
from pptx import Presentation

def test_canonical_hashes_and_stage_resume():
    root = Path(__file__).resolve().parents[1]
    assert not validate_manifest_hashes(root)
    assert resume_stage({'outline': True, 'anchor': True, 'style_dna': True}) == 'AWAITING_STAGE1_APPROVAL'
    assert resume_stage({'validated_single_page_pptx': True}) != 'deck_merge'

def test_lineage_rejects_stale_attempt():
    a = {'slide_id':'S1','source_fingerprint':'v1','text_clean_fingerprint':'clean1'}
    assert active_attempt_matches(a,'S1','v1','clean1')
    assert not active_attempt_matches(a,'S1','v2','clean1')
    assert not active_attempt_matches(a,'S1','v1','clean2')
    assert text_clean_fingerprint(b'a') != text_clean_fingerprint(b'b')
    assert text_restore_fingerprint({'x':1},{'font':'A'}) != text_restore_fingerprint({'x':1},{'font':'B'})

def test_inspector_detects_missing_relationship_target(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    good = tmp_path/'good.pptx'
    bad = tmp_path/'bad.pptx'
    prs.save(good)
    with zipfile.ZipFile(good) as src, zipfile.ZipFile(bad,'w') as dst:
        for item in src.infolist():
            if item.filename != 'ppt/slideLayouts/slideLayout7.xml':
                dst.writestr(item, src.read(item.filename))
    assert inspect_ooxml_dependencies(bad)
