"""Positive canvas continuation and source-bound formal wiring; no remote calls."""
import json,sys,copy
from pathlib import Path
import pytest
from PIL import Image
from pptx import Presentation
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
sys.path.insert(0,str(Path(__file__).resolve().parent))
import runtime, reconstruction_router as router, layer_package as lp, layer_bridge as lb
from test_reconstruction_backend_router import ready_state
from test_layer_bridge import image_state,write_plan,approve_ready_plan,lb_sha
from test_layer_revision import pattern,fixture
from layer_build_pptx import build_page_pptx

def write_canvas(state,sid,size,target):
    path=state.parent/(sid+'-canvas.json')
    data={'source_width':size[0],'source_height':size[1]}
    if target is not None:
        data.update(target_slide_width_emu=target[0],target_slide_height_emu=target[1])
    path.write_text(json.dumps(data),encoding='utf-8')
    runtime.register_stage2_artifact(state,'canvas:'+sid,path)
    return path

def explicit_state(tmp_path,sizes=((941,1672),(940,1672)),target=(5143500,9144000)):
    state=ready_state(tmp_path)
    for sid,size in zip(('S001','S002'),sizes):
        clean=tmp_path/(sid+'-clean-full.png')
        Image.new('RGB',size,'white').save(clean)
        runtime.register_stage2_artifact(state,'text_clean:'+sid,clean)
        write_canvas(state,sid,size,target)
    return state

def test_940_941_differences_continue_to_one_approved_canvas(tmp_path):
    state=explicit_state(tmp_path)
    rec=router.choose_backend(state,'image_layer')
    assert rec['target_slide_emu']==[5143500,9144000]
    before=state.read_bytes()
    assert router.choose_backend(state,'image_layer')==rec
    assert state.read_bytes()==before

@pytest.mark.parametrize('scale',[1,2])
def test_proportional_resolution_change_continues(tmp_path,scale):
    state=explicit_state(tmp_path,((900*scale,1600*scale),)*2)
    assert router.choose_backend(state,'image_layer')['target_slide_emu']==[5143500,9144000]

@pytest.mark.parametrize('width,accepted',[(904,True),(905,True),(906,False)])
def test_five_source_pixel_boundary_is_preserved(tmp_path,width,accepted):
    state=explicit_state(tmp_path,((900,900),)*2,(width*10000,9000000))
    if accepted: assert router.choose_backend(state,'image_layer')['target_slide_emu']==[width*10000,9000000]
    else:
        with pytest.raises(ValueError,match='five source pixels'):router.choose_backend(state,'image_layer')

@pytest.mark.parametrize('kind',['missing','half','boolean','zero','oversize','float','conflict','dimensions','stale'])
def test_invalid_approval_is_not_silently_used(tmp_path,kind):
    state=explicit_state(tmp_path)
    p=state.parent/'S002-canvas.json';c=json.loads(p.read_text())
    if kind=='missing':c={'source_width':940,'source_height':1672}
    elif kind=='half':c.pop('target_slide_width_emu')
    elif kind=='boolean':c['target_slide_width_emu']=True
    elif kind=='zero':c['target_slide_width_emu']=0
    elif kind=='oversize':c['target_slide_width_emu']=999999999
    elif kind=='float':c['target_slide_width_emu']=5143500.0
    elif kind=='conflict':c['target_slide_width_emu']*=2;c['target_slide_height_emu']*=2
    elif kind=='dimensions':c['source_width']=900
    elif kind=='stale':c['source_width']=939
    p.write_text(json.dumps(c),encoding='utf-8')
    if kind!='stale':runtime.register_stage2_artifact(state,'canvas:S002',p)
    before=state.read_bytes()
    with pytest.raises(ValueError):router.choose_backend(state,'image_layer')
    assert state.read_bytes()==before

def test_no_explicit_target_keeps_original_fallback(tmp_path):
    state=explicit_state(tmp_path,((1600,900),)*2,None)
    assert router.choose_backend(state,'image_layer')['target_slide_emu']==[12192000,6858000]

def test_no_explicit_target_keeps_mixed_ratio_rejection(tmp_path):
    state=explicit_state(tmp_path,((941,1672),(940,1672)),None)
    with pytest.raises(ValueError,match='inconsistent source aspect'):router.choose_backend(state,'image_layer')

def test_reselection_does_not_silently_retarget(tmp_path):
    state=explicit_state(tmp_path);router.choose_backend(state,'image_layer')
    for sid,size in zip(('S001','S002'),((941,1672),(940,1672))):write_canvas(state,sid,size,(10287000,18288000))
    before=state.read_bytes()
    with pytest.raises(ValueError,match='selected target differs'):router.choose_backend(state,'image_layer')
    assert state.read_bytes()==before
    assert router.selected_backend(state)=='image_layer'

def current_formal_fixture(tmp_path):
    source,plan,result=fixture(tmp_path)
    payload=json.loads(plan.read_text())
    payload.update(schema_version=1,max_foregrounds=6,deferred=[])
    for target in payload['targets']:target.update(label='A',edit_action='move',selection_reason='fixture')
    plan.write_text(json.dumps(payload),encoding='utf-8')
    state=ready_state(tmp_path)
    for sid in ('S001','S002'):
        runtime.register_stage2_artifact(state,'text_clean:'+sid,source)
        write_canvas(state,sid,(512,384),(9144000,6858000))
    router.choose_backend(state,'image_layer')
    overlay=tmp_path/'overlay.png';Image.open(source).save(overlay)
    lb.register_plan(state,'S001',plan,overlay);approve_ready_plan(state,'S001')
    s=runtime.load_runtime_state(state);p=s['layer_bridge']['plans']['S001']
    s['layer_bridge'].setdefault('results',{})['S001']={'status':'LAYER_RESULT_BOUND','slide_id':'S001',
        'run_id':p['run_id'],'plan_sha256':p['plan_sha256'],'text_clean_sha256':p['text_clean_sha256'],
        'task_id':'one','path':str(result),'sha256':lb_sha(result),'usage':{}}
    runtime._save_runtime_state(state,s)
    return state,source,plan,result

def test_formal_cli_builds_approved_canvas_and_correct_object_placement(tmp_path):
    state,source,plan,result=current_formal_fixture(tmp_path)
    bundle=tmp_path/'formal-bundle'
    assert lp.main(['--state',str(state),'--slide-id','S001','--source-image',str(source),
        '--plan',str(plan),'--result',str(result),'--output',str(bundle)])==0
    manifest=json.loads((bundle/'manifest.json').read_text())
    assert manifest['ppt']=={'slide_width_emu':9144000,'slide_height_emu':6858000,'target_canvas_source':'explicit_target'}
    page=tmp_path/'page.pptx';build_page_pptx(bundle,page)
    prs=Presentation(page);shape=prs.slides[0].shapes[1]
    assert [prs.slide_width,prs.slide_height]==[9144000,6858000]
    assert [shape.left,shape.top,shape.width,shape.height]==[
        round(9144000*70/512),round(6858000*65/384),round(9144000*90/512),round(6858000*65/384)]
    assert lb.bind_bundle(state,'S001',bundle)['status']=='LAYER_BUNDLE_READY'

@pytest.mark.parametrize('bad',['source','plan','result','page','stale'])
def test_formal_cli_identity_rejection_writes_no_bundle(tmp_path,bad):
    state,source,plan,result=current_formal_fixture(tmp_path)
    if bad in ('source','plan','result'):
        original={'source':source,'plan':plan,'result':result}[bad]
        copy_path=tmp_path/('other-'+original.name);copy_path.write_bytes(original.read_bytes())
        if bad=='source':source=copy_path
        if bad=='plan':plan=copy_path
        if bad=='result':result=copy_path
    if bad=='stale':(tmp_path/'S001-canvas.json').write_text('{}')
    bundle=tmp_path/'not-created'
    assert lp.main(['--state',str(state),'--slide-id','S999' if bad=='page' else 'S001',
        '--source-image',str(source),'--plan',str(plan),'--result',str(result),'--output',str(bundle)])==1
    assert not bundle.exists()

def test_cli_pair_is_required(tmp_path):
    with pytest.raises(SystemExit):
        lp.main(['--state','state','--source-image','s','--plan','p','--result','r','--output','o'])

def test_direct_bundle_binding_checks_current_canvas(tmp_path):
    state,source,plan,result=current_formal_fixture(tmp_path)
    bundle=tmp_path/'bundle';lp.package_layers(source,plan,result,bundle,explicit_target_emu=[9144000,6858000])
    for sid in ('S001','S002'):write_canvas(state,sid,(512,384),(18288000,13716000))
    before=state.read_bytes()
    with pytest.raises(ValueError,match='selected target differs'):lb.bind_bundle(state,'S001',bundle)
    assert state.read_bytes()==before

def test_canvas_change_prevents_new_submit_without_touching_budget(tmp_path,monkeypatch):
    state=image_state(tmp_path);plan,overlay=write_plan(tmp_path)
    lb.register_plan(state,'S001',plan,overlay);approve_ready_plan(state,'S001')
    for sid in ('S001','S002'):write_canvas(state,sid,(1600,900),(24384000,13716000))
    monkeypatch.setenv(router.IMAGE_LAYER_API_KEY_ENV,'synthetic-test-key')
    class Provider:
        calls=0
        @staticmethod
        def submit_task(*a):Provider.calls+=1;return {'task_id':'fake'}
    before=state.read_bytes()
    with pytest.raises(ValueError,match='selected target differs'):lb.submit(state,'S001',tmp_path/'work',Provider)
    assert Provider.calls==0 and state.read_bytes()==before

def test_existing_pending_task_still_queries_after_canvas_change(tmp_path,monkeypatch):
    state=image_state(tmp_path);plan,overlay=write_plan(tmp_path)
    lb.register_plan(state,'S001',plan,overlay);approve_ready_plan(state,'S001')
    monkeypatch.setenv(router.IMAGE_LAYER_API_KEY_ENV,'synthetic-test-key')
    class Provider:
        calls=0
        @staticmethod
        def submit_task(*a):Provider.calls+=1;return {'task_id':'fake'}
        @staticmethod
        def query_task(task,key):return {'task_id':task,'normalized_status':'RUNNING'}
    assert lb.submit(state,'S001',tmp_path/'work',Provider)['status']=='LAYER_REQUEST_PENDING'
    for sid in ('S001','S002'):write_canvas(state,sid,(1600,900),(24384000,13716000))
    before=state.read_bytes()
    assert lb.query(state,'S001',tmp_path/'work',Provider)['status']=='LAYER_REQUEST_PENDING'
    assert Provider.calls==1 and state.read_bytes()==before

@pytest.mark.parametrize('size',[(940,1672),(941,1672),(1800,3200)])
def test_pixel_difference_and_resolution_change_produce_approved_pptx(tmp_path,size):
    source=tmp_path/'source.png';im=pattern(source,size);im.save(tmp_path/'resized.png')
    im.convert('RGBA').save(tmp_path/'bg.png')
    from PIL import ImageDraw
    fg=Image.new('RGBA',size);ImageDraw.Draw(fg).rectangle((70,65,159,129),fill=(180,20,30,230));fg.save(tmp_path/'fg.png')
    plan=tmp_path/'plan.json';plan.write_text(json.dumps({'page_id':'S001','targets':[{'id':'a','bbox':[60,55,170,140]}],'recommended_total_layers':2}))
    result=tmp_path/'result.json';result.write_text(json.dumps({'backend_canvas':list(size),'resized_image':'resized.png',
        'image_boxes':[[60,55,170,140]],'resized_image_boxes':[[60,55,170,140]],'boxes_mapping_index':[0],
        'layers':[{'type':'background','file':'bg.png'},{'type':'foreground','target_id':'a','file':'fg.png'}]}))
    bundle=tmp_path/'bundle'
    lp.package_layers(source,plan,result,bundle,explicit_target_emu=[5143500,9144000])
    page=tmp_path/'page.pptx';build_page_pptx(bundle,page);prs=Presentation(page)
    assert [prs.slide_width,prs.slide_height]==[5143500,9144000]
    shape=prs.slides[0].shapes[1]
    assert [shape.left,shape.top]==[round(5143500*70/size[0]),round(9144000*65/size[1])]

def test_exactly_five_source_pixels_produces_pptx(tmp_path):
    source=tmp_path/'source.png';im=pattern(source,(900,900));im.save(tmp_path/'resized.png');im.convert('RGBA').save(tmp_path/'bg.png')
    from PIL import ImageDraw
    fg=Image.new('RGBA',im.size);ImageDraw.Draw(fg).rectangle((70,65,159,129),fill=(180,20,30,230));fg.save(tmp_path/'fg.png')
    plan=tmp_path/'plan.json';plan.write_text(json.dumps({'page_id':'S001','targets':[{'id':'a','bbox':[60,55,170,140]}],'recommended_total_layers':2}))
    result=tmp_path/'result.json';result.write_text(json.dumps({'backend_canvas':[900,900],'resized_image':'resized.png',
        'image_boxes':[[60,55,170,140]],'resized_image_boxes':[[60,55,170,140]],'boxes_mapping_index':[0],
        'layers':[{'type':'background','file':'bg.png'},{'type':'foreground','target_id':'a','file':'fg.png'}]}))
    bundle=tmp_path/'bundle';lp.package_layers(source,plan,result,bundle,explicit_target_emu=[9050000,9000000])
    page=tmp_path/'page.pptx';build_page_pptx(bundle,page)
    prs=Presentation(page);assert [prs.slide_width,prs.slide_height]==[9050000,9000000]
