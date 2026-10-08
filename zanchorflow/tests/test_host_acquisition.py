import json, sqlite3, sys
from pathlib import Path
import pytest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import host_acquisition as host
import canva_bridge as bridge
import runtime
from test_acquisition_request import setup
from pptx import Presentation

def fixture(tmp_path):
    state,identity,lookup=setup(tmp_path)
    req={'identity':identity,'edit_url':lookup['response']['design']['urls']['edit_url'],'status':'WAIT_DOWNLOAD'}
    record=tmp_path/'record.json';host.ready(record,req)
    # Historical persisted acquisition shape, not a new-policy record.
    legacy=host._read(record);legacy.pop('require_transport_receipt');host._save(record,legacy)
    host.lock(record)
    profile=tmp_path/'profile';profile.mkdir()
    target=tmp_path/'custom-downloads';target.mkdir()
    (profile/'Preferences').write_text(json.dumps({'download':{'default_directory':str(target)}}))
    f=target/'exact.pptx';deck=Presentation();deck.slides.add_slide(deck.slide_layouts[6]);deck.save(f)
    url='https://www.canva.com/design/fixture-design/edit'
    db=sqlite3.connect(profile/'History')
    db.execute('create table downloads (id integer, guid text, target_path text, state integer, received_bytes integer, total_bytes integer, tab_url text, start_time integer)')
    locked=json.loads(record.read_text())['locked_at_chrome']
    db.execute('insert into downloads values (1,?,?,?,?,?,?,?)',('fixture-guid',str(f),1,f.stat().st_size,f.stat().st_size,url,locked+1));db.commit();db.close()
    return state,record,profile,f,url

def test_custom_browser_default_and_exact_record(tmp_path):
    state,record,profile,f,url=fixture(tmp_path);before=state.read_bytes()
    result=host.acquire(record,state,profile,url,True)
    assert result['status']=='ACQUIRED' and result['download_path']==str(f)
    assert result['download_evidence']['browser_default_directory']==str(f.parent)
    assert state.read_bytes()==before and bridge.status(state)['status']=='WAIT_DOWNLOAD'

def test_intent_and_failed_never_allow_second_click(tmp_path):
    state,record,profile,f,url=fixture(tmp_path)
    with pytest.raises(ValueError):host.lock(record)
    host.fail(record,'DOWNLOAD_UNKNOWN')
    with pytest.raises(ValueError):host.lock(record)
    assert bridge.status(state)['status']=='WAIT_DOWNLOAD'

def test_restart_reuses_accepted_identity(tmp_path):
    state,record,profile,f,url=fixture(tmp_path)
    assert json.loads(record.read_text())['identity']==bridge.active_identity(runtime.load_runtime_state(state),'S001')
    result=host.acquire(record,state,profile,url,True)
    assert result['identity']['design_id']=='fixture-design'

@pytest.mark.parametrize('kind',['identity','incomplete','ambiguous','file_changed','url'])
def test_unproven_download_rejected_without_runtime_change(tmp_path,kind):
    state,record,profile,f,url=fixture(tmp_path);before=state.read_bytes();event=True
    if kind=='event':event=False
    elif kind=='identity':
        data=json.loads(record.read_text());data['identity']['attempt_id']='wrong';record.write_text(json.dumps(data))
    elif kind=='url':url='https://www.canva.com.evil.test/design/fixture-design/edit'
    elif kind=='file_changed':f.write_bytes(b'bad')
    else:
        db=sqlite3.connect(profile/'History')
        if kind=='incomplete':db.execute('update downloads set state=0')
        else:db.execute('insert into downloads select 2,guid,target_path,state,received_bytes,total_bytes,tab_url,start_time from downloads')
        db.commit();db.close()
    with pytest.raises(ValueError):host.acquire(record,state,profile,url,event)
    assert state.read_bytes()==before


def _add_recovery_url_chain(profile, design_id='fixture-design'):
    db = sqlite3.connect(profile/'History')
    db.execute("update downloads set tab_url='' where id=1")
    db.execute('create table downloads_url_chains (id integer, chain_index integer, url text)')
    db.execute('insert into downloads_url_chains values (1,0,?)',
               (f'https://www.canva.com/design/{design_id}/export/download?format=pptx',))
    db.commit(); db.close()


def test_recovery_link_download_with_blank_tab_url_uses_strict_url_chain_provenance(tmp_path):
    state,record,profile,f,url=fixture(tmp_path)
    _add_recovery_url_chain(profile)
    result=host.acquire(record,state,profile,url,True)
    assert result['status']=='ACQUIRED'
    evidence=result['download_evidence']
    assert evidence['design_provenance']['method']=='recovery_download_url'
    assert evidence['design_provenance']['design_id']=='fixture-design'
    assert 'fixture-design' in evidence['design_provenance']['evidence_url']
    assert evidence['download_record']['tab_url']==''


def test_blank_tab_url_without_recovery_url_provenance_still_fails_closed(tmp_path):
    state,record,profile,f,url=fixture(tmp_path); before=state.read_bytes()
    db=sqlite3.connect(profile/'History'); db.execute("update downloads set tab_url='' where id=1"); db.commit(); db.close()
    with pytest.raises(ValueError,match='same-design download record pending'):
        host.acquire(record,state,profile,url,True)
    assert state.read_bytes()==before


def test_recovery_url_for_other_design_is_rejected(tmp_path):
    state,record,profile,f,url=fixture(tmp_path); before=state.read_bytes()
    _add_recovery_url_chain(profile,'other-design')
    with pytest.raises(ValueError,match='same-design download record pending'):
        host.acquire(record,state,profile,url,True)
    assert state.read_bytes()==before


def test_recovery_link_canonical_finish_binds_same_acquisition(tmp_path,monkeypatch):
    import shutil
    state,record,profile,f,url=fixture(tmp_path)
    # Add the canonical trusted connector request used by finish.
    _,identity,lookup=setup(tmp_path/'lookup-fixture')
    # Reuse this fixture's current identity while preserving the trusted /d/ URL association.
    lookup['identity']=json.loads(record.read_text())['identity']
    lookup['response']['design']['id']='fixture-design'
    lookup['observation']='synthetic readonly connector evidence'
    lookup_path=tmp_path/'lookup.json'; lookup_path.write_text(json.dumps(lookup))
    final_dir=tmp_path/'graphics-first'; final_dir.mkdir()
    host.prepare(record,state,'S001',lookup_path,final_dir,'S001.pptx')
    _add_recovery_url_chain(profile)
    def fake_finalize(source,target_dir,target_name,python_executable):
        target=Path(target_dir)/target_name
        shutil.copyfile(source,target)
        digest=bridge.sha(target)
        return {'Status':'Success','Path':str(target),'SHA256':digest,'ValidPptx':True,'SlideCount':1}
    monkeypatch.setattr(host,'_finalize',fake_finalize)
    result=host.finish(record,state,profile,url,True,wait_seconds=0)
    saved=json.loads(record.read_text())
    assert result['status']=='DOWNLOAD_BOUND'
    assert saved['status']=='ACQUIRED'
    assert saved['download_evidence']['design_provenance']['method']=='recovery_download_url'
    assert bridge.status(state)['status']=='RESTORE_TEXT'
    current=runtime.load_runtime_state(state)
    assert current['canva_bridge']['downloads']['S001']['identity']['design_id']=='fixture-design'
    assert current['canva_bridge']['downloads']['S001']['sha256']==result['sha256']


def test_normal_tab_url_records_primary_provenance(tmp_path):
    state,record,profile,f,url=fixture(tmp_path)
    result=host.acquire(record,state,profile,url,True)
    assert result['download_evidence']['design_provenance']=={
        'method':'tab_url','design_id':'fixture-design','evidence_url':url}


def test_nonblank_mismatched_tab_url_cannot_fallback_to_matching_recovery_chain(tmp_path):
    state,record,profile,f,url=fixture(tmp_path);before=state.read_bytes()
    db=sqlite3.connect(profile/'History')
    db.execute("update downloads set tab_url='https://www.canva.com/design/other-design/edit' where id=1")
    db.execute('create table downloads_url_chains (id integer, chain_index integer, url text)')
    db.execute('insert into downloads_url_chains values (1,0,?)',
               ('https://www.canva.com/design/fixture-design/export/download?format=pptx',))
    db.commit();db.close()
    with pytest.raises(ValueError,match='same-design download record pending'):
        host.acquire(record,state,profile,url,True)
    assert state.read_bytes()==before


def test_recovery_url_chain_requires_trusted_canva_host_and_exact_design_token(tmp_path):
    state,record,profile,f,url=fixture(tmp_path);before=state.read_bytes()
    db=sqlite3.connect(profile/'History')
    db.execute("update downloads set tab_url='' where id=1")
    db.execute('create table downloads_url_chains (id integer, chain_index integer, url text)')
    db.execute('insert into downloads_url_chains values (1,0,?)',
               ('https://www.canva.com.evil.test/design/fixture-design/export/download',))
    db.execute('insert into downloads_url_chains values (1,1,?)',
               ('https://www.canva.com/design/fixture-design-extra/export/download',))
    db.commit();db.close()
    with pytest.raises(ValueError,match='same-design download record pending'):
        host.acquire(record,state,profile,url,True)
    assert state.read_bytes()==before


def test_recovery_url_chain_ambiguous_multiple_download_records_fail_closed(tmp_path):
    state,record,profile,f,url=fixture(tmp_path);before=state.read_bytes()
    _add_recovery_url_chain(profile)
    db=sqlite3.connect(profile/'History')
    locked=json.loads(record.read_text())['locked_at_chrome']
    second=f.parent/'second.pptx';deck=Presentation();deck.slides.add_slide(deck.slide_layouts[6]);deck.save(second)
    db.execute('insert into downloads values (2,?,?,?,?,?,?,?)',
               ('fixture-guid-2',str(second),1,second.stat().st_size,second.stat().st_size,'',locked+2))
    db.execute('insert into downloads_url_chains values (2,0,?)',
               ('https://www.canva.com/design/fixture-design/export/download?format=pptx',))
    db.commit();db.close()
    with pytest.raises(ValueError,match='exactly one same-design download record required'):
        host.acquire(record,state,profile,url,True)
    assert state.read_bytes()==before


def test_recovery_url_chain_before_intent_lock_is_ignored(tmp_path):
    state,record,profile,f,url=fixture(tmp_path);before=state.read_bytes()
    _add_recovery_url_chain(profile)
    db=sqlite3.connect(profile/'History')
    locked=json.loads(record.read_text())['locked_at_chrome']
    db.execute('update downloads set start_time=? where id=1',(locked-1,))
    db.commit();db.close()
    with pytest.raises(ValueError,match='same-design download record pending'):
        host.acquire(record,state,profile,url,True)
    assert state.read_bytes()==before


def test_valid_optional_record_without_event_is_accepted(tmp_path):
    state,record,profile,f,url=fixture(tmp_path)
    result=host.acquire(record,state,profile,url,False)
    assert result['status']=='ACQUIRED'
    assert result['download_evidence']['download_event_confirmed'] is False
