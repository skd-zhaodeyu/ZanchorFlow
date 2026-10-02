import sys
from pathlib import Path
from pptx import Presentation
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
sys.path.insert(0,str(Path(__file__).resolve().parent))
import layer_package as lp
import layer_build_pptx as bp
from test_layer_bundle import fixture


def test_build_page_uses_manifest_target_canvas_and_canonical_geometry(tmp_path):
    source,plan,result=fixture(tmp_path)
    manifest=lp.package_layers(source,plan,result,tmp_path/'bundle', schema_version=2)
    out=tmp_path/'page.pptx'; built=bp.build_page_pptx(tmp_path/'bundle',out)
    prs=Presentation(out)
    assert len(prs.slides)==1
    assert prs.slide_width==manifest['ppt']['slide_width_emu']
    assert prs.slide_height==manifest['ppt']['slide_height_emu']
    assert len(prs.slides[0].shapes)==manifest['counts']['actual_total_layers']==2
    fg=prs.slides[0].shapes[1]; norm=manifest['layers'][1]['canonical_bbox_normalized']
    assert abs(fg.left-round(prs.slide_width*norm[0]))<=1
    assert abs(fg.top-round(prs.slide_height*norm[1]))<=1
    assert abs(fg.width-round(prs.slide_width*norm[2]))<=1
    assert abs(fg.height-round(prs.slide_height*norm[3]))<=1
    assert built['status']=='PASS'


def test_graphics_first_has_only_picture_shapes_and_z_order(tmp_path):
    source,plan,result=fixture(tmp_path); lp.package_layers(source,plan,result,tmp_path/'bundle', schema_version=2)
    out=tmp_path/'page.pptx'; bp.build_page_pptx(tmp_path/'bundle',out); prs=Presentation(out)
    shapes=list(prs.slides[0].shapes)
    assert [s.name for s in shapes]==['Background','Foreground 01 - a']
    assert all(not getattr(s,'has_text_frame',False) for s in shapes)


def test_local_layer_pipeline_has_cli_entrypoints(tmp_path,capsys):
    import layer_validate_bundle as lv
    source,plan,result=fixture(tmp_path)
    bundle=tmp_path/'bundle'; out=tmp_path/'page.pptx'
    assert lp.main(['--schema-version','2','--source-image',str(source),'--plan',str(plan),'--result',str(result),'--output',str(bundle)])==0
    assert lv.main(['--bundle',str(bundle)])==0
    assert bp.main(['--bundle',str(bundle),'--output',str(out)])==0
    assert out.is_file()
