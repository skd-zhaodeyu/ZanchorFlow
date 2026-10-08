"""Office selection/lifecycle tests; ordinary cases never create real Office."""
import copy,hashlib,json,shutil,sys,types,zipfile
from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import preflight,merge_pptx as m,runtime,assemble_deck as a
ROOT=Path(__file__).resolve().parents[1]


def _probes(monkeypatch,pp=True,wps=True):
    calls=[]
    def probe(host,ok):
        def run(_):calls.append(host);return ok,'synthetic '+host
        return run
    monkeypatch.setattr(preflight,'_probe_powerpoint',probe('powerpoint',pp))
    monkeypatch.setattr(preflight,'_probe_wps',probe('wps',wps))
    original=preflight.importlib.util.find_spec
    monkeypatch.setattr(preflight.importlib.util,'find_spec',lambda n:object() if n=='win32com' else original(n))
    return calls

@pytest.mark.parametrize('choice,pp,wps,selected,calls',[
 ('auto',True,True,'powerpoint',['powerpoint']),
 ('auto',False,True,'wps',['powerpoint','wps']),
 ('auto',False,False,None,['powerpoint','wps']),
 ('powerpoint',False,True,None,['powerpoint']),
 ('wps',True,True,'wps',['wps']),
 ('wps',True,False,None,['wps'])])
def test_scoped_selection(monkeypatch,tmp_path,choice,pp,wps,selected,calls):
    observed=_probes(monkeypatch,pp,wps)
    report=preflight.check(ROOT,tmp_path,'merge',office_host=choice)
    assert observed==calls
    assert report['office_host']==selected
    assert preflight.exit_code(report)==(0 if selected else 1)
    if selected=='wps':
        assert report['office_progid']=='KWPP.Application'
        assert report['powerpoint_com']!='PASS'
        assert report['wps_com']=='PASS' and report['blockers']==[]
    if not selected:assert len(report['blockers'])==len(calls)


def test_probe_exception_can_choose_wps(monkeypatch,tmp_path):
    calls=_probes(monkeypatch)
    def failed(_):raise OSError('synthetic cleanup RPC failure')
    monkeypatch.setattr(preflight,'_probe_powerpoint',failed)
    report=preflight.check(ROOT,tmp_path,'merge')
    assert report['office_host']=='wps' and report['blockers']==[]
    assert 'RPC' in report['office_diagnostics'][0]['detail']
    assert calls==['wps']

@pytest.mark.parametrize('scope',['package','acquisition','image_layer'])
def test_non_merge_scope_never_probes_office(monkeypatch,tmp_path,scope):
    def forbidden(_):raise AssertionError('Office probed in unrelated scope')
    monkeypatch.setattr(preflight,'_probe_powerpoint',forbidden)
    monkeypatch.setattr(preflight,'_probe_wps',forbidden)
    report=preflight.check(ROOT,tmp_path,scope)
    assert report['office_host'] is None
    assert report['powerpoint_com']==report['wps_com']=='NOT_REQUIRED'

@pytest.mark.parametrize('missing',['pptx','lxml','win32com','yaml'])
def test_missing_common_dependency_is_not_application_fallback(monkeypatch,tmp_path,missing):
    calls=_probes(monkeypatch)
    original=preflight.importlib.util.find_spec
    monkeypatch.setattr(preflight.importlib.util,'find_spec',lambda n:None if n==missing else original(n))
    report=preflight.check(ROOT,tmp_path,'merge')
    assert calls==[] and report['office_host'] is None
    assert any(x['code']=='PYTHON_DEPENDENCY_MISSING' and x['detail']==missing for x in report['blockers'])


def _mock_com(monkeypatch,src,pp_unavailable=False,import_error=False,cleanup_error=False,save_error=False,reconnect_error=False):
    client=types.ModuleType('win32com.client'); pkg=types.ModuleType('win32com');pkg.client=client
    apps=[];calls=[];saved={};shared=types.SimpleNamespace(closed=False)
    class Slides:
        Count=0
        def Add(self,*_):self.Count+=1
        def InsertFromFile(self,path,_):
            if import_error:raise OSError('PRIMARY import failure')
            self.Count+=1; self.source=path;return 1
        def __call__(self,_):return types.SimpleNamespace(Design=None)
    class Doc:
        def __init__(self,app,path=None):
            self.app=app;self.Slides=Slides();self.PageSetup=types.SimpleNamespace()
            self.Designs=types.SimpleNamespace(Load=lambda _:object());self.path=path
            if path:self.Slides.Count=saved[path]
        def SaveAs(self,path,_):
            if save_error:raise OSError('PRIMARY save failure')
            saved[path]=self.Slides.Count
            if hasattr(self.Slides,'source'):shutil.copyfile(self.Slides.source,path)
            else:
                p=Presentation();p.slides.add_slide(p.slide_layouts[6]);p.save(path)
        def Save(self):pass
        def Close(self):
            if cleanup_error:raise OSError('CLEANUP document failure')
            if self.app.host=='wps':self.app.dead=True
            self.app.own_closes+=1
    class App:
        Name='Microsoft PowerPoint'
        def __init__(self,host):self.host=host;self.dead=False;self.own_closes=0;self.quits=0
        @property
        def Presentations(self):
            if self.dead:raise OSError('old WPS handle cannot be used')
            return types.SimpleNamespace(Add=lambda _:Doc(self),Open=lambda path,*_:Doc(self,path))
        def Quit(self):
            self.quits+=1
            if cleanup_error:raise OSError('CLEANUP quit failure')
            assert self.host=='powerpoint','WPS Quit would affect shared documents'
    def dispatch(progid):
        calls.append(progid)
        if pp_unavailable and progid=='PowerPoint.Application':raise OSError('synthetic PP unavailable')
        host='powerpoint' if progid=='PowerPoint.Application' else 'wps'
        if reconnect_error and host=='wps' and any(x.host=='wps' for x in apps):raise OSError('synthetic WPS reconnect failure')
        app=App(host);apps.append(app);return app
    client.DispatchEx=dispatch
    monkeypatch.setitem(sys.modules,'win32com',pkg);monkeypatch.setitem(sys.modules,'win32com.client',client)
    return apps,calls,shared


def _one_page(path,text='S1'):
    p=Presentation();sl=p.slides.add_slide(p.slide_layouts[6]);sl.shapes.add_textbox(Inches(.5),Inches(.5),Inches(7),Inches(.7)).text=text;p.save(path)
    return path

@pytest.mark.parametrize('host,pp_failed,expected', [('auto',False,'powerpoint'),('auto',True,'wps'),('wps',False,'wps')])
def test_native_selection_reopen_and_progid_identity(monkeypatch,tmp_path,host,pp_failed,expected):
    src=_one_page(tmp_path/'source.pptx');apps,calls,shared=_mock_com(monkeypatch,src,pp_unavailable=pp_failed)
    info=m._native_merge([{'slide_id':'S1','path':src}],['S1'],tmp_path/'out.pptx',office_host=host)
    assert info['office_host']==expected
    assert info['office_progid']==m.OFFICE_PROGIDS[expected]
    assert not shared.closed
    if expected=='wps':
        assert len([x for x in apps if x.host=='wps'])==2
        assert calls[-2:]==['KWPP.Application']*2 and all(x.quits==0 for x in apps)
    else:assert calls==['PowerPoint.Application'] and apps[0].quits==1

@pytest.mark.parametrize('error',['import','save','reconnect'])
def test_operational_failure_does_not_switch_after_creation(monkeypatch,tmp_path,error):
    src=_one_page(tmp_path/'source.pptx')
    args={error+'_error':True} if error!='import' else {'import_error':True}
    apps,calls,_=_mock_com(monkeypatch,src,**args)
    host='wps' if error=='reconnect' else 'auto'
    with pytest.raises(OSError):m._native_merge([{'slide_id':'S1','path':src}],['S1'],tmp_path/'out.pptx',office_host=host)
    assert calls==(['KWPP.Application']*2 if host=='wps' else ['PowerPoint.Application'])


def test_primary_error_survives_cleanup(monkeypatch,tmp_path):
    src=_one_page(tmp_path/'source.pptx');_mock_com(monkeypatch,src,import_error=True,cleanup_error=True)
    info={}
    with pytest.raises(OSError,match='PRIMARY import'):
        m._native_merge([{'slide_id':'S1','path':src}],['S1'],tmp_path/'out.pptx',diagnostics=info)
    assert len(info['cleanup_errors'])==2


def test_powerpoint_probe_retains_operation_and_cleanup_errors(monkeypatch,tmp_path):
    src=_one_page(tmp_path/'source.pptx');_mock_com(monkeypatch,src,save_error=True,cleanup_error=True)
    ok,detail=preflight._probe_powerpoint(tmp_path)
    assert not ok and all(x in detail for x in ('PRIMARY save','CLEANUP document','CLEANUP quit'))


def test_real_wps_probe_contract_with_com_double(monkeypatch,tmp_path):
    src=_one_page(tmp_path/'source.pptx');apps,calls,_=_mock_com(monkeypatch,src)
    assert preflight._probe_wps(tmp_path)==(True,'')
    assert calls==['KWPP.Application']*2 and all(x.quits==0 for x in apps)


def test_prepare_passes_chosen_host_once_and_retains_audit(monkeypatch,tmp_path):
    from test_merge import _assembly_fixture
    assembly,r,state=_assembly_fixture(tmp_path,monkeypatch);calls=[]
    def check(*args,**kw):calls.append(('probe',kw['office_host']));return {'blockers':[],'office_host':'wps'}
    monkeypatch.setattr(preflight,'check',check)
    original=assembly.merge
    def merge(*args,**kw):
        calls.append(('merge',kw['office_host']));result=original(*args,**kw)
        return {**result,'office_host':'wps','office_progid':'KWPP.Application'}
    monkeypatch.setattr(assembly,'merge',merge)
    result=assembly.prepare(state,tmp_path/'work')
    assert calls==[('probe','auto'),('merge','wps')]
    assert result['office_host']=='wps'
    assert json.loads(Path(result['validation_report']).read_text())['native_merge']['office_progid']=='KWPP.Application'


def test_prepare_failure_preserves_prior_state_and_files(monkeypatch,tmp_path):
    from test_merge import _assembly_fixture
    assembly,r,state=_assembly_fixture(tmp_path,monkeypatch)
    old=state.read_bytes();inputs={p:p.read_bytes() for p in tmp_path.glob('S*') if p.is_file()}
    def failed(*args,**kw):return {'status':'FAIL','details':['synthetic selected host failure'],'office_host':'wps'}
    monkeypatch.setattr(assembly,'merge',failed)
    with pytest.raises(ValueError):assembly.prepare(state,tmp_path/'work')
    assert state.read_bytes()==old and all(p.read_bytes()==v for p,v in inputs.items())
    assert not (tmp_path/'work/deck-review.json').exists()


def test_declared_scope_restores_all_existing_text_files_and_detects_mutation():
    from office_scope import data,restore_repo_bytes
    d=data();repo=ROOT.parent
    for rel in {x['repo_path'] for x in d['patches']}:
        assert hashlib.sha256(restore_repo_bytes(rel,(repo/rel).read_bytes())).hexdigest()==d['baseline'][rel]
    for rel in ('scripts/runtime.py','scripts/merge_pptx.py'):
        with pytest.raises(AssertionError):restore_repo_bytes('zanchorflow/'+rel,(ROOT/rel).read_bytes()+b'\n# unauthorized change\n')
    assert m.GEOMETRY_TOLERANCE_EMU==0


def test_all_other_runtime_files_are_baseline_bytes():
    from office_scope import data,restore_repo_bytes
    d=data()
    for p in (ROOT/'scripts').iterdir():
        if not p.is_file():continue
        rel='zanchorflow/'+p.relative_to(ROOT).as_posix()
        from download_scope import new_runtime, verify_new_runtime
        if p.relative_to(ROOT).as_posix() in new_runtime():
            verify_new_runtime(p.relative_to(ROOT).as_posix(),p.read_bytes());continue
        assert hashlib.sha256(restore_repo_bytes(rel,p.read_bytes())).hexdigest()==d['baseline'][rel]


def _real_inputs(folder):
    from PIL import Image
    folder.mkdir(parents=True,exist_ok=True);state={'deck_order':['S2','S1'],'slides':{}};pages=[]
    for sid,color in [('S1',(180,30,30)),('S2',(30,30,180))]:
        # Separate source folders deliberately produce conflicting media/image1.png names.
        local=folder/sid;local.mkdir();img=local/'picture.png';Image.new('RGB',(32,24),color).save(img)
        p=Presentation();p.slide_width=Inches(10);p.slide_height=Inches(7.5)
        sl=p.slides.add_slide(p.slide_layouts[6]);bg=sl.slide_layout.slide_master.background.fill;bg.solid();bg.fore_color.rgb=RGBColor(*color)
        for i,(text,size) in enumerate([(sid+' 标题 Editable',30),('中文 English body',18),('caption',12)]):
            shape=sl.shapes.add_textbox(Inches(.6),Inches(.6+i*.8),Inches(8.8),Inches(.65))
            shape.text=text
            for run in shape.text_frame.paragraphs[0].runs:run.font.name='Microsoft YaHei';run.font.size=Pt(size)
        sl.shapes.add_shape(1,Inches(1),Inches(4),Inches(2),Inches(1)).fill.solid()
        sl.shapes.add_picture(str(img),Inches(6),Inches(4),Inches(2),Inches(1.5))
        pptx=local/'input.pptx';p.save(pptx);pages.append(pptx)
        values={'approved_render':b'synthetic-render-'+sid.encode(),'text_clean':b'synthetic-clean-'+sid.encode(),
          'final_content_truth':{'title':sid},'canvas':{'width':10,'height':7.5},'finalized_manifest':{'title':sid},'font_fallback':{}}
        state['slides'][sid]={}
        for key,value in values.items():
            target=local/key;target.write_bytes(value if isinstance(value,bytes) else json.dumps(value).encode())
            state['slides'][sid][key]=str(target)
    state_path=folder/'state.json';state_path.write_text(json.dumps(state),encoding='utf-8')
    records=[runtime.seal_validated_single_page(state_path,sid,path,{g:'PASS' for g in runtime.GATE_NAMES}) for sid,path in zip(('S1','S2'),pages)]
    return state_path,records,pages



def test_direct_application_failure_has_both_diagnostics_and_preserves_existing_bytes(monkeypatch,tmp_path):
    state,records,pages=_real_inputs(tmp_path/'inputs')
    before=state.read_bytes();out=tmp_path/'existing.pptx';out.write_bytes(b'existing-preserved')
    source_bytes={p:p.read_bytes() for p in pages};calls=[]
    src=pages[0];_mock_com(monkeypatch,src)
    import win32com.client
    def unavailable(progid):calls.append(progid);raise OSError('synthetic unavailable '+progid)
    monkeypatch.setattr(win32com.client,'DispatchEx',unavailable)
    result=m.merge(records,['S2','S1'],out,state)
    assert result['status']=='FAIL' and result['office_host'] is None
    assert calls==['PowerPoint.Application','KWPP.Application']
    assert len(result['office_attempts'])==2
    assert state.read_bytes()==before and out.read_bytes()==b'existing-preserved'
    assert all(p.read_bytes()==v for p,v in source_bytes.items())

@pytest.mark.integration
@pytest.mark.parametrize('host',['powerpoint','wps'])
def test_actual_merge_reopen_edit_and_negative_files(tmp_path,monkeypatch,host):
    import win32com.client
    state,records,pages=_real_inputs(tmp_path/'inputs');source_bytes={p:p.read_bytes() for p in pages}
    probe=preflight.check(ROOT,tmp_path/'probe','merge',office_host=host)
    assert probe['office_host']==host and probe['blockers']==[],probe
    out=tmp_path/'merged.pptx';result=m.merge(records,['S2','S1'],out,state,office_host=host)
    assert result['status']=='PASS' and result['office_host']==host,result
    prs=Presentation(out);assert len(prs.slides)==2
    assert [sl.shapes[0].text for sl in prs.slides]==['S2 标题 Editable','S1 标题 Editable']
    assert not m.inspect_ooxml_dependencies(out)
    # Actual edit/save/reopen uses a copy; input and accepted merged bytes remain intact.
    edit=tmp_path/'edited.pptx';shutil.copyfile(out,edit);progid=m.OFFICE_PROGIDS[host]
    app=win32com.client.DispatchEx(progid);doc=None
    try:
        doc=app.Presentations.Open(str(edit),0,0,0)
        doc.Slides(1).Shapes(1).TextFrame.TextRange.Text='S2 修改后 Edited'
        doc.Save();doc.Close();doc=None
        if host=='wps':app=None;app=win32com.client.DispatchEx(progid)
        doc=app.Presentations.Open(str(edit),0,0,0)
        assert 'Edited' in doc.Slides(1).Shapes(1).TextFrame.TextRange.Text
        doc.Close();doc=None
    finally:
        if doc is not None:doc.Close()
        if host=='powerpoint' and app is not None:app.Quit()
        app=None
    edited=Presentation(edit);assert edited.slides[0].shapes[0].text=='S2 修改后 Edited'
    # All other objects and the entire other page must preserve exact canonical signatures.
    before=m._shape_signature(prs);after=m._shape_signature(edited)
    assert not m._shape_signature_errors(before[0][1:],after[0][1:],0)
    assert not m._shape_signature_errors(before[1],after[1],0)
    assert not m.inspect_ooxml_dependencies(edit)
    initial_state=state.read_bytes();initial_out=out.read_bytes()
    for fault in ('order','image'):
        bad=tmp_path/(fault+'.pptx');broken=Presentation(out)
        if fault=='order':
            ids=broken.slides._sldIdLst;ids.insert(0,ids[1])
        else:
            picture=next(s for s in broken.slides[0].shapes if s.shape_type==13)
            from PIL import Image
            import io
            payload=io.BytesIO();Image.new('RGB',(32,24),(1,250,1)).save(payload,format='PNG')
            picture._pic.blipFill.blip.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed',broken.slides[0].part.get_or_add_image_part(io.BytesIO(payload.getvalue()))[1])
        broken.save(bad)
        def corrupted(inputs,order,target,**kw):
            shutil.copyfile(bad,target);kw['diagnostics'].update(office_host=host,office_progid=progid)
        with monkeypatch.context() as context:
            context.setattr(m,'_native_merge',corrupted)
            failed=m.merge(records,['S2','S1'],out,state,office_host=host)
        assert failed['status']=='FAIL' and failed['code']=='DECK_MERGE_INTEGRITY_FAILED',failed
        if fault=='image':assert any('picture_sha256_changed' in x for x in failed['details'])
        assert state.read_bytes()==initial_state and out.read_bytes()==initial_out
    assert all(p.read_bytes()==v for p,v in source_bytes.items())
    (tmp_path/'evidence.json').write_text(json.dumps({'host':host,'merge':result,'negative_faults':['order','image'],'edit':'SAVED_REOPENED_OTHER_OBJECTS_UNCHANGED'},ensure_ascii=False,indent=2),encoding='utf-8')

@pytest.mark.integration
@pytest.mark.parametrize('backend',['magic_layer','image_layer'])
def test_wps_formal_assembly_with_current_three_page_seals(tmp_path,backend):
    import canva_bridge as cb,layer_bridge as lb
    from test_title_routes import restored_pages
    from test_merge import _complete_deck_review
    state,pages=restored_pages(tmp_path,backend)
    seal=cb.seal_page if backend=='magic_layer' else lb.seal_page
    for sid,(_,restored,review) in pages.items():seal(state,sid,restored,review)
    result=a.prepare(state,tmp_path/'work',office_host='wps')
    assert result['status']=='AWAITING_DECK_VALIDATION' and result['office_host']=='wps',result
    assert result['preflight']['powerpoint_com']=='NOT_REQUIRED'
    review=Path(result['validation_report']);_complete_deck_review(review)
    final=tmp_path/'outputs/deck.pptx';assert a.publish(state,review,final)['status']=='PASS'
    before=state.read_bytes();assert (lb.status(state) if backend=='image_layer' else cb.status(state))['status']=='COMPLETE'
    assert state.read_bytes()==before and len(Presentation(final).slides)==3
