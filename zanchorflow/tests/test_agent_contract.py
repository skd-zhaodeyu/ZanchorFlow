import base64, copy, json, re, sys
from pathlib import Path
import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import host_acquisition as host
import runtime, canva_bridge as bridge
from test_acquisition_request import setup
DOC=Path(__file__).resolve().parents[1]/'docs/codex-canva-bridge.md'

def example(name,values):
    match=re.search(r'<!-- '+re.escape(name)+r' -->\s*```python\n(.*?)```',DOC.read_text(encoding='utf-8'),re.S)
    assert match is not None, 'missing documented safety example: '+name
    env=dict(values);exec(compile(match.group(1),str(DOC)+'#'+name,'exec'),env)
    return env

def original(tmp_path,text_kind=None):
    file=tmp_path/'source.pptx';p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
    shape=s.shapes.build_freeform(0,0).convert_to_shape()
    body=shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}txBody')
    if body is not None:shape._element.remove(body)
    if text_kind:
        parent=s.shapes.add_group_shape().shapes if text_kind=='group' else s.shapes
        tx=parent.add_textbox(0,0,Inches(1),Inches(1));tx.text='existing text'
        if text_kind=='hidden':tx._element.xpath('./p:nvSpPr/p:cNvPr')[0].set('hidden','1')
    p.save(file);return file

def test_readonly_freeform_has_no_added_text_body(tmp_path):
    file=original(tmp_path);raw=file.read_bytes()
    result=example('readonly-neutrality',{'source_pptx':file})
    assert result['existing_text_nodes']==[]
    assert result['original_shape_xml']==[(s.shape_id,etree.tostring(s._element)) for s in result['source_slide'].shapes]
    assert file.read_bytes()==raw
    assert not result['source_slide'].shapes[0]._element.xpath('./p:txBody')

@pytest.mark.parametrize('kind',['plain','group','hidden'])
def test_existing_text_is_detected_without_mutation(tmp_path,kind):
    file=original(tmp_path,kind);raw=file.read_bytes()
    result=example('readonly-neutrality',{'source_pptx':file})
    assert result['existing_text_nodes']==['existing text']
    assert result['original_shape_xml']==[(s.shape_id,etree.tostring(s._element)) for s in result['source_slide'].shapes]
    assert file.read_bytes()==raw

@pytest.mark.parametrize('change',[False,True])
def test_saved_old_objects_are_compared_to_source(tmp_path,change):
    file=original(tmp_path);env=example('readonly-neutrality',{'source_pptx':file})
    p=env['source_deck'];s=p.slides[0];s.shapes.add_textbox(0,0,100,100).text='restored'
    if change:s.shapes[0].text  # reproduce the actual bad getter, not an ignored empty-node exception
    dest=tmp_path/'restored.pptx';p.save(dest)
    values={'source_pptx':file,'restored_pptx':dest}
    if change:
        with pytest.raises(AssertionError):example('verify-original-objects',values)
    else:example('verify-original-objects',values)

def decision(count,**changes):
    values=dict(intent_count=count,actual_download_count=count,current_status='WAIT_DOWNLOAD',
        identity_matches=True,terminal_failure_confirmed=True,required_permissions_ok=True,
        sources_unambiguous=True,operation_in_progress=False,completed_file_available=False,
        failure_checkpoint='EXPORT_FAILED')
    values.update(changes)
    return example('download-retry-gate',values)

@pytest.mark.parametrize('success',[True,False])
def test_final_download_fifth_attempt_and_no_sixth(tmp_path,success):
    state,identity,lookup=setup(tmp_path)
    req={'status':'WAIT_DOWNLOAD','identity':identity,'edit_url':lookup['response']['design']['urls']['edit_url']}
    paths=[];clicks=[]
    for number in range(1,6):
        if number>1:
            gate=decision(len(paths));assert gate['retry_allowed'] and gate['next_download_number']==number
        record=tmp_path/(identity['attempt_id']+('' if number==1 else '-download-'+str(number).zfill(2))+'.json')
        host.ready(record,req);host.lock(record);paths.append(record);clicks.append(number)
        host.fail(record,'EXPORT_FAILED' if number<5 or not success else 'FIXTURE_COMPLETED')
    assert clicks==[1,2,3,4,5]
    assert decision(len(paths),completed_file_available=success)['retry_allowed'] is (not success)
    for path in paths:
        assert Path(str(path)+'.intent').exists()
        with pytest.raises(ValueError):host.lock(path)
    assert bridge.status(state)['status']=='WAIT_DOWNLOAD'

@pytest.mark.parametrize('changes',[{'operation_in_progress':True},{'operation_in_progress':None},{'completed_file_available':True},
 {'completed_file_available':None},{'identity_matches':False},{'required_permissions_ok':False},
 {'failure_checkpoint':'FINALIZER'},{'current_status':'RESTORE_TEXT'}])
def test_unknown_or_non_export_failure_never_retries(changes):
    assert not decision(1,**changes)['retry_allowed']

def test_restart_keeps_same_record_budget_but_new_export_gets_fresh_budget(tmp_path):
    state,identity,lookup=setup(tmp_path)
    req={'status':'WAIT_DOWNLOAD','identity':identity,'edit_url':lookup['response']['design']['urls']['edit_url']}
    first=tmp_path/'first.json';host.ready(first,req)
    for _ in range(5):host.reserve_retry(first,state,'S001','query');host.reserve_retry(first,state,'S001','ui')
    for kind in ('query','ui'):
        assert host.reserve_retry(first,state,'S001',kind)['retry_number']==6
    host.lock(first);host.fail(first,'EXPORT_FAILED')
    second=tmp_path/'second.json';host.ready(second,req)
    assert host.reserve_retry(second,state,'S001','query')['retry_number']==1
    assert host.reserve_retry(second,state,'S001','ui')['retry_number']==1
    evidence=tmp_path/'controller.json';evidence.write_text(json.dumps({'intent_count':5,'actual_download_count':5}))
    loaded=json.loads(evidence.read_text())
    assert decision(loaded['intent_count'],actual_download_count=loaded['actual_download_count'])['retry_allowed']
    assert Path(str(first)+'.intent').exists()


def test_runtime_contract_is_recovery_first_and_skill_readonly():
    raw=DOC.read_text(encoding='utf-8')
    assert 'No fixed refusal budget' in raw
    assert 'carry-query-ui-budget' not in raw
    assert 'READ-ONLY PRODUCT' in raw
    assert 'IMPLEMENTATION_DEFECT_SUSPECTED' in raw
    assert 'example.com sanity' not in raw

def test_single_runtime_truth_contains_executable_prepare_lock_finish_path():
    raw=DOC.read_text(encoding='utf-8')
    assert 'host_acquisition.py prepare' in raw
    assert 'host_acquisition.py lock' in raw
    assert 'host_acquisition.py finish' in raw
    assert '--event-confirmed' in raw
    assert 'preflight.py --skill-dir' in raw and '--scope acquisition' in raw
    assert 'wait_seconds' in raw and 'actually wait' in raw

def test_normal_acquisition_does_not_require_magic_layers_or_preselected_download_dir():
    raw=DOC.read_text(encoding='utf-8')
    assert 'magic_layers_tool' not in raw
    assert 'dedicated absolute download directory' not in raw



def _stage2_review_cli_case(tmp_path):
    fields = ('page_task', 'title', 'core_expression', 'key_information',
              'semantic_relation', 'acceptance_criteria')
    outline = {'slides': [{field: f'{field}-1' for field in fields}]}
    outline_path = tmp_path / 'outline.json'
    outline_path.write_text(json.dumps(outline, ensure_ascii=False), encoding='utf-8')
    state_path = tmp_path / 'runtime-state.json'
    state_path.write_text('{"slides":{}}', encoding='utf-8')
    runtime.record_stage1_draft(state_path, outline_path)
    runtime.present_stage1_review(state_path)
    runtime.handle_stage1_reply(state_path, '同意，按这个继续。', decision='approve')
    runtime.start_stage2_run(state_path)
    png = tmp_path / 'candidate.png'
    png.write_bytes(base64.b64decode(
        'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl2pFQAAAAASUVORK5CYII='))
    candidate = runtime.register_stage2_candidate(
        state_path, 'S001', png, generation_intent='structural concept')
    return state_path, candidate


def test_stage2_reject_cli_requires_and_persists_replan_evidence(tmp_path, capsys):
    state_path, candidate = _stage2_review_cli_case(tmp_path)
    common = [
        'stage2-candidate-review', '--state', str(state_path),
        '--candidate-id', candidate['candidate_id'], '--verdict', 'REJECT',
        '--failed-hard-gate', 'H1',
        '--evidence', '主要视觉结构编码了错误关系',
        '--correction', '重新规划主要视觉结构',
    ]
    assert runtime.main(common) == 2
    blocked = json.loads(capsys.readouterr().out)
    assert 'STAGE2_REJECT_REPLAN_EVIDENCE_REQUIRED' in blocked['code']

    assert runtime.main(common + [
        '--replan-evidence', '修复核心关系必须重新规划主要视觉结构'
    ]) == 0
    accepted = json.loads(capsys.readouterr().out)
    assert accepted['review']['replan_evidence'] == '修复核心关系必须重新规划主要视觉结构'

SKILL=Path(__file__).resolve().parents[1]/'SKILL.md'

def test_rc18_skill_has_explicit_reconstruction_branch_gate_and_fixed_image_onboarding():
    raw=SKILL.read_text(encoding='utf-8')
    assert 'Magic Layer 分支（推荐）' in raw
    assert 'Image Layer 分支' in raw
    assert '推荐使用 Magic Layer 分支。' in raw
    image = (SKILL.parent / 'references' / '08_Image_Layer_Model可编辑PPT重建协议_v1.1.md').read_text(encoding='utf-8')
    assert 'Image Layer preparation/QA/retry must read 08' in raw
    assert 'https://research.360.cn/workspace/apikeys' in image
    assert '请核实当前账户的免费权益、余额与价格，并在真实提交前确认费用授权。' in image
    assert 'references/08_Image_Layer_Model可编辑PPT重建协议_v1.1.md' in raw
    assert 'all Text-Clean' in raw or '全部 Text-Clean' in raw
    assert 'The seven files listed' not in raw


def test_rc18_skill_keeps_magic_and_image_protocol_loading_separate():
    raw=SKILL.read_text(encoding='utf-8')
    assert 'Magic Layer 分支' in raw and '05 + 06 + 07' in raw
    assert 'Image Layer 分支' in raw and '08 + 07' in raw


def test_image_layer_visual_qa_is_required_before_graphics_first_binding():
    root=Path(__file__).resolve().parents[1]
    entry=(root/'SKILL.md').read_text(encoding='utf-8')
    assert 'Image Layer preparation/QA/retry must read 08' in entry
    raw=(root/'references'/'08_Image_Layer_Model可编辑PPT重建协议_v1.1.md').read_text(encoding='utf-8')
    assert 'register-visual-qa' in raw
    assert 'layer_isolation' in raw
    assert 'background_repair' in raw
    assert 'recomposite_fidelity' in raw
    assert 'LAYER_VISUAL_QA_REQUIRED' in raw
