"""Offline checks exercise existing cleanup and state contracts; no remote retry proof."""
import sys
from pathlib import Path
import pytest
import yaml
from pptx import Presentation
from lxml import etree
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import canva_bridge as bridge
import runtime
from test_canvas_adaptation import source,canvas,registration
from test_agent_contract import example

ROOT=Path(__file__).resolve().parents[1]


def test_author_and_baseline_metadata_are_separate():
    front=yaml.safe_load((ROOT/'SKILL.md').read_text(encoding='utf-8').split('---',2)[1])
    manifest=yaml.safe_load((ROOT/'manifest.yaml').read_text(encoding='utf-8'))
    assert front['name']==manifest['skill_name']=='zanchorflow'
    assert front['metadata']['author']=='David-Z'
    assert manifest['skill_version']=='0.9.0-rc19-candidate'
    assert len(manifest['canonical_protocols'])==8
    assert manifest['non_runtime_exceptions']==['docs/codex-canva-bridge.md']
    assert manifest['progressive_loading']['stage3_magic_layer']==[6]
    assert manifest['progressive_loading']['stage3_image_layer']==[8]
    assert not (ROOT/'agents/openai.yaml').exists()


def plan(glyph='𝒜',role='native_text'):
    element={'id':'M001','truth_ref':'S001:T001','truth_source':'stage2_final_content_truth',
      'content':glyph,'text_role':'decorative_text','representation_role':role,
      'treatment':'remove_and_restore','visual_bbox_normalized':[.1,.1,.2,.1]}
    item={'region_id':'R001','observed_text':glyph,'detection_source':'full_page_vision',
      'treatment':'remove_and_restore','linked_manifest_ids':['M001'],
      'destructive_cleanup_risk':'low','visual_bbox_normalized':[.1,.1,.2,.1]}
    truth={'slide_id':'S001','meaningful_text':[{'text':glyph}],'exact_facts':[]}
    return {'slide_id':'S001','text_elements':[element]}, {'slide_id':'S001','items':[item]}, truth


def test_special_noncritical_letter_can_remain_native_without_protection_claim():
    manifest,inventory,truth=plan()
    result=runtime.validate_stage3_text_plan(manifest,inventory,truth,'S001')
    assert result['status']=='PASS' and result['preserved_graphic_typography']==[]


@pytest.mark.parametrize('role',['graphic_typography','approved_graphic_asset'])
def test_registered_graphic_role_cannot_silently_use_native_treatment(role):
    manifest,inventory,truth=plan(role=role)
    with pytest.raises(ValueError,match='GRAPHIC_TYPOGRAPHY_TREATMENT_INVALID|PRESERVED_GRAPHIC_CONTENT_INVALID'):
        runtime.validate_stage3_text_plan(manifest,inventory,truth,'S001')


@pytest.mark.parametrize('uniform',[False,True])
@pytest.mark.parametrize('restore',[False,True])
def test_same_source_letter_cleanup_and_optional_native_restore(tmp_path,uniform,restore):
    src,graphic_id,text_id=source(tmp_path,text=True)
    original_hash=bridge.sha(src)
    prs=Presentation(src);prs.slides[0].shapes[-1].text='𝒜';prs.save(src)
    original_hash=bridge.sha(src)
    target=canvas() if uniform else canvas(2000000,2000000)
    cleanup={'download_sha256':original_hash,'removed_text_shape_ids':[text_id],
             'evidence':'synthetic noncritical ordinary textbox; current unprotected content decision'}
    deck,report=bridge.prepare_canvas_mapping(src,target,registration(src),cleanup)
    assert report['plan']['mode']==('uniform_fit' if uniform else 'unchanged_objects')
    assert report['retained_shape_ids']==[graphic_id]
    if restore:
        element={'visual_bbox_normalized':[.1,.1,.2,.1],'font_size_normalized':.05}
        geom=bridge.map_native_text_geometry(element,report['plan'])
        deck.slides[0].shapes.add_textbox(*geom['bbox_emu']).text='𝒜'
    out=tmp_path/'restored.pptx';deck.save(out)
    report['restored_pptx_sha256']=bridge.sha(out)
    bridge.verify_canvas_mapping(src,out,target,report,'synthetic-clean-fingerprint')
    values={'source_pptx':src,'restored_pptx':out,'approved_canvas':target,
            'canvas_mapping_review':report,'current_text_clean_fingerprint':'synthetic-clean-fingerprint'}
    example('verify-original-objects',values)
    actual=Presentation(out)
    assert [s.shape_id for s in list(actual.slides[0].shapes)[:1]]==[graphic_id]
    assert sum(len(s._element.xpath('.//a:t')) for s in actual.slides[0].shapes)==int(restore)
    assert bridge.sha(src)==original_hash


@pytest.mark.parametrize('uniform',[False,True])
@pytest.mark.parametrize('bad',['source','evidence','mixed','unknown_id'])
def test_cleanup_rejection_does_not_change_download(tmp_path,uniform,bad):
    src,graphic_id,text_id=source(tmp_path,text=True)
    if bad=='mixed':
        prs=Presentation(src);prs.slides[0].shapes[0].text='A';prs.save(src)
    before=bridge.sha(src)
    cleanup={'download_sha256':before,'removed_text_shape_ids':[text_id],'evidence':'synthetic proof'}
    if bad=='source':cleanup['download_sha256']='unrelated-download'
    elif bad=='evidence':cleanup['evidence']=''
    elif bad=='mixed':cleanup['removed_text_shape_ids']=[graphic_id]
    else:cleanup['removed_text_shape_ids']=[999]
    target=canvas() if uniform else canvas(2000000,2000000)
    with pytest.raises(ValueError):bridge.prepare_canvas_mapping(src,target,registration(src),cleanup)
    assert bridge.sha(src)==before


@pytest.mark.parametrize('uniform',[False,True])
@pytest.mark.parametrize('tamper',['graphic_loss','graphic_change','cleanup_report'])
def test_cleanup_report_cannot_redefine_graphic_baseline(tmp_path,uniform,tamper):
    src,graphic_id,text_id=source(tmp_path,text=True)
    target=canvas() if uniform else canvas(2000000,2000000)
    cleanup={'download_sha256':bridge.sha(src),'removed_text_shape_ids':[text_id],'evidence':'synthetic text-only cleanup'}
    deck,report=bridge.prepare_canvas_mapping(src,target,registration(src),cleanup)
    if tamper=='graphic_loss':deck.slides[0].shapes._spTree.remove(deck.slides[0].shapes[0]._element)
    elif tamper=='graphic_change':deck.slides[0].shapes[0].left+=1
    else:report['cleanup']['removed_text_shape_ids']=[graphic_id]
    out=tmp_path/'tampered.pptx';deck.save(out);report['restored_pptx_sha256']=bridge.sha(out)
    with pytest.raises(ValueError):bridge.verify_canvas_mapping(src,out,target,report)


def test_no_report_does_not_excuse_deleted_original_objects(tmp_path):
    src,graphic_id,text_id=source(tmp_path,text=True)
    cleanup={'download_sha256':bridge.sha(src),'removed_text_shape_ids':[text_id],'evidence':'synthetic local cleanup'}
    deck,report=bridge.prepare_canvas_mapping(src,canvas(2000000,2000000),cleanup=cleanup)
    out=tmp_path/'restored.pptx';deck.save(out)
    with pytest.raises(AssertionError):example('verify-original-objects',{'source_pptx':src,'restored_pptx':out})


@pytest.mark.parametrize('original_tolerance',[False,True])
def test_formal_seal_records_legal_cleanup_in_both_canvas_paths(tmp_path,original_tolerance):
    import json
    from test_canvas_adaptation import sealed_fixture
    state,src,out,target,review,review_path=sealed_fixture(tmp_path,ordinary_text=True,original_tolerance=original_tolerance)
    original=bridge.sha(src)
    current=runtime.load_runtime_state(state)
    manifest=json.loads(Path(current['slides']['S001']['finalized_manifest']).read_text(encoding='utf-8'))
    contents=[e['content'] for e in manifest['text_elements'] if e['treatment']=='remove_and_restore']
    deck=Presentation(out);deck.slides[0].shapes[-1].text='\n'.join(contents);deck.save(out)
    review['restored_pptx_sha256']=bridge.sha(out)
    review['canvas_mapping']['restored_pptx_sha256']=bridge.sha(out)
    review_path.write_text(json.dumps(review),encoding='utf-8')
    assert review['canvas_mapping']['cleanup']['removed_text_shape_ids']
    assert review['canvas_mapping']['plan']['mode']==('unchanged_objects' if original_tolerance else 'uniform_fit')
    assert bridge.seal_page(state,'S001',out,review_path)['status']=='PAGE_SEALED'
    record=bridge.page_provenance(runtime.load_runtime_state(state),'S001')
    assert record['canvas_mapping_review']['sha256']==bridge.sha(review_path)
    assert bridge.sha(src)==original
    review_path.write_text(review_path.read_text(encoding='utf-8')+' ',encoding='utf-8')
    with pytest.raises(ValueError,match='sealed mapping review changed'):
        bridge.page_provenance(runtime.load_runtime_state(state),'S001')


def test_all_runtime_scripts_remain_exactly_the_delivered_baseline():
    import hashlib
    from refinement_scope import PINNED_RUNTIME_FILES
    assert {str(p.relative_to(ROOT)).replace('\\','/') for p in (ROOT/'scripts').glob('*') if p.is_file()}==set(PINNED_RUNTIME_FILES)
    for rel,digest in PINNED_RUNTIME_FILES.items():
        if rel in {'scripts/reconstruction_router.py','scripts/layer_package.py','scripts/layer_bridge.py'}:
            from publish_scope import assert_approved_runtime
            assert_approved_runtime(rel,ROOT/rel,digest)
        else:
            assert hashlib.sha256((ROOT/rel).read_bytes()).hexdigest()==digest,rel


def test_every_migrated_runtime_block_preserves_original_payload():
    from refinement_scope import MIGRATED_BLOCKS,inherited_doc
    for rel in MIGRATED_BLOCKS:
        inherited_doc(rel,(ROOT/rel).read_text(encoding='utf-8'))
