import json
import pytest
from test_acquisition_request import setup
import canva_bridge as bridge
import runtime
from pptx import Presentation

def evidence(identity):
    return {**identity,'edit_url':'https://www.canva.com/d/fixture-token','download_entry':'fixture download record','completion_evidence':'fixture completed record'}

def register(state,lookup,tmp_path):
    lookup={**lookup,'observation':'controller captured readonly get_design'}
    p=tmp_path/'lookup.json';p.write_text(json.dumps(lookup),encoding='utf-8')
    return bridge.register_design_url(state,'S001',p)

def test_trusted_alias_and_existing_bound_file_reusable(tmp_path):
    state,identity,lookup=setup(tmp_path);register(state,lookup,tmp_path)
    ev=evidence(identity)
    f=tmp_path/'download.pptx';deck=Presentation();deck.slides.add_slide(deck.slide_layouts[6]);deck.save(f)
    ev['pptx_sha256']=bridge.sha(f);ep=tmp_path/'evidence.json';ep.write_text(json.dumps(ev))
    assert bridge.bind_download(state,'S001',f,ep)['status']=='DOWNLOAD_BOUND'
    current=bridge.download_current(runtime.load_runtime_state(state),'S001')
    assert current['identity']==identity and current['sha256']==bridge.sha(f)
    f.write_bytes(b'changed')
    with pytest.raises(ValueError):bridge.download_current(runtime.load_runtime_state(state),'S001')

def test_alias_cannot_be_self_reported(tmp_path):
    state,identity,lookup=setup(tmp_path)
    with pytest.raises(ValueError):bridge._check_download_url(evidence(identity),identity,runtime.load_runtime_state(state))

@pytest.mark.parametrize('url',['https://evil.test/d/token','https://www.canva.com.evil.test/d/token','http://www.canva.com/d/token','https://user:pass@www.canva.com/d/token','https://www.canva.com:444/d/token','https://www.canva.com/design/other/edit'])
def test_alias_security_rejects(tmp_path,url):
    state,identity,lookup=setup(tmp_path)
    lookup['response']['design']['urls']['edit_url']=url
    before=state.read_bytes()
    with pytest.raises(ValueError):register(state,lookup,tmp_path)
    assert state.read_bytes()==before

@pytest.mark.parametrize('key',['attempt_id','design_id','slide_id'])
def test_lookup_identity_rejects(tmp_path,key):
    state,identity,lookup=setup(tmp_path);lookup['identity'][key]='wrong'
    before=state.read_bytes()
    with pytest.raises(ValueError):register(state,lookup,tmp_path)
    assert state.read_bytes()==before

def test_lookup_hash_change_rejects(tmp_path):
    state,identity,lookup=setup(tmp_path);register(state,lookup,tmp_path)
    (tmp_path/'lookup.json').write_text('{}')
    with pytest.raises(ValueError):bridge._check_download_url(evidence(identity),identity,runtime.load_runtime_state(state))
