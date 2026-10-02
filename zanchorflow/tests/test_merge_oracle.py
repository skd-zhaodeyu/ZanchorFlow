import sys
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from merge_pptx import _shape_signature


def base():
    p=Presentation(); s=p.slides.add_slide(p.slide_layouts[6])
    a=s.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(1)); a.text='A'
    b=s.shapes.add_textbox(Inches(4),Inches(1),Inches(2),Inches(1)); b.text='B'
    return p,s,a,b


def test_geometry_change_detected():
    p,s,a,b=base(); before=_shape_signature(p)
    a.left += Inches(0.25)
    assert _shape_signature(p)!=before


def test_z_order_change_detected():
    p,s,a,b=base(); before=_shape_signature(p)
    a._element.addprevious(b._element)
    assert _shape_signature(p)!=before


def test_shape_type_change_detected():
    p,s,a,b=base(); before=_shape_signature(p)
    parent=a._element.getparent(); parent.remove(a._element)
    replacement=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1),Inches(1),Inches(2),Inches(1))
    replacement.text='A'; b._element.addprevious(replacement._element)
    assert _shape_signature(p)!=before


def test_picture_bytes_change_detected(tmp_path):
    first,second=tmp_path/'a.png',tmp_path/'b.png'
    Image.new('RGB',(20,20),(255,0,0)).save(first)
    Image.new('RGB',(20,20),(0,0,255)).save(second)
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_picture(str(first),Inches(1),Inches(1),width=Inches(1))
    q=Presentation();t=q.slides.add_slide(q.slide_layouts[6]);t.shapes.add_picture(str(second),Inches(1),Inches(1),width=Inches(1))
    assert _shape_signature(p)!=_shape_signature(q)
