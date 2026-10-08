"""Synthetic external results, actual title registration/lineage/seal/assembly APIs.
Visual assessments here are fixtures, not claims about live generation or font-fit quality.
"""
import copy, json, io, zipfile
from pathlib import Path
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
import runtime as r
import canva_bridge as cb
import title_policy as t
import layer_bridge as lb
import reconstruction_router as router
import assemble_deck as assembly
from test_title_policy import setup, choose, payload, plan
from test_preflight_scope import _accepted_attempt
from test_image_layer_assembly import prepare_image_page
from test_merge import _complete_deck_review


def build_visual_set(tmp_path):
    state=setup(tmp_path);choose(state)
    # Current host route for the synthetic Magic boundary only; this never calls Canva.
    saved=r.load_runtime_state(state);saved['canva_bridge']['route']='codex-canva';r._save_runtime_state(state,saved)
    for kind in ('anchor','style_dna'):
        p=tmp_path/(kind+'.json');p.write_text('{}',encoding='utf-8');r.register_stage2_artifact(state,kind,p)
    for sid in ('S001','S002','S003'):
        p=tmp_path/(sid+'-visual.png');Image.new('RGB',(1600,900),'white').save(p)
        ctx=t.context(state,sid)
        rec=r.register_stage2_candidate(state,sid,p,'OFFLINE navigation-heading page fixture',title_contract_ref=ctx.get('contract_ref'))
        r.review_stage2_candidate(state,rec['candidate_id'],'PASS')
        if sid=='S002':t.bind(state,sid,payload(state,sid))
    ref=t.context(state,'S003')['contract_ref']
    manifests={}
    for sid in ('S001','S002','S003'):
        m,inv=plan(state,tmp_path,sid,ref,'OFFLINE heading '+sid)
        if sid=='S001':
            data=json.loads(m.read_text(encoding='utf-8'));data.pop('title_contract_ref');data.pop('title_element_ids');m.write_text(json.dumps(data),encoding='utf-8')
        manifests[sid]=(m,inv)
        for kind,obj in [('canvas',{'width':1600,'height':900}),('font_fallback',{'Arial':'Arial'})]:
            p=tmp_path/(sid+'-'+kind+'.json');p.write_text(json.dumps(obj),encoding='utf-8');r.register_stage2_artifact(state,kind+':'+sid,p)
    assert r.stage2_visual_status(state)['status']=='COMPLETE'
    r.present_stage2_formal_display(state);r.seal_stage2_visual_approval(state,'OFFLINE fixture final pages approved')
    return state,manifests


def add_native_title(path, text):
    prs=Presentation(path);slide=prs.slides[0]
    box=slide.shapes.add_textbox(int(prs.slide_width*.1),int(prs.slide_height*.1),int(prs.slide_width*.4),int(prs.slide_height*.1))
    box.text=text
    for paragraph in box.text_frame.paragraphs:
        for run in paragraph.runs:run.font.name='Arial';run.font.size=Pt(24)
    prs.save(path)
    assert [s.text for s in Presentation(path).slides[0].shapes if s.has_text_frame and s.text]==[text]


def restored_pages(tmp_path,backend):
    state,manifests=build_visual_set(tmp_path)
    for sid,(manifest,inv) in manifests.items():
        clean=tmp_path/(sid+'-clean.png');Image.new('RGB',(1600,900),'white').save(clean)
        if backend=='magic_layer':
            cb.register_text_plan(state,sid,manifest,inv);cb.register_text_clean(state,sid,clean)
        else:
            r.register_stage2_artifact(state,'finalized_manifest:'+sid,manifest)
            r.register_stage2_artifact(state,'removal_inventory:'+sid,inv)
            r.register_stage2_artifact(state,'text_clean:'+sid,clean)
    # Font fallback is frozen after Manifest registration, as in the original workflow.
    for sid in manifests:
        p=tmp_path/(sid+'-font_fallback.json');p.write_text(json.dumps({'Arial':'Arial'}),encoding='utf-8')
        r.register_stage2_artifact(state,'font_fallback:'+sid,p)
    router.choose_backend(state,backend)
    pages={}
    for sid,(manifest,inv) in manifests.items():
        if backend=='image_layer':
            graphics,restored,review=prepare_image_page(state,tmp_path,sid,json.loads(manifest.read_text(encoding='utf-8')))
        else:
            identity=_accepted_attempt(state,sid,tmp_path,design='OFFLINE-'+sid)
            graphics=tmp_path/(sid+'-graphics.pptx');prs=Presentation();prs.slide_width=12192000;prs.slide_height=6858000;prs.slides.add_slide(prs.slide_layouts[6]);prs.save(graphics)
            evidence={**identity,'pptx_sha256':cb.sha(graphics),'edit_url':'https://www.canva.com/design/OFFLINE-'+sid+'/edit',
                      'download_entry':'OFFLINE synthetic event fixture','completion_evidence':'OFFLINE complete fixture'}
            ep=tmp_path/(sid+'-download.json');ep.write_text(json.dumps(evidence),encoding='utf-8');cb.bind_download(state,sid,graphics,ep)
            restored=tmp_path/(sid+'-restored.pptx');restored.write_bytes(graphics.read_bytes())
            review=tmp_path/(sid+'-review.json')
            record={**identity,'download_sha256':cb.sha(graphics),'checks':{k:'PASS' for k in r.GATE_NAMES},
                    'evidence':{k:'OFFLINE fixture observation, not live visual evidence' for k in r.GATE_NAMES}}
            review.write_text(json.dumps(record),encoding='utf-8')
        add_native_title(restored,'OFFLINE heading '+sid)
        record=json.loads(review.read_text(encoding='utf-8'));record.update(restored_pptx_sha256=cb.sha(restored),
            text_restore_fingerprint=r.expected_lineage_from_state(r.load_runtime_state(state),sid)['text_restore_fingerprint'])
        review.write_text(json.dumps(record),encoding='utf-8')
        pages[sid]=(graphics,restored,review)
    return state,pages


def offline_native(monkeypatch):
    """Only replace native import for no-Office CI; preserve text/picture relationships."""
    import preflight,merge_pptx
    original=preflight.importlib.util.find_spec
    monkeypatch.setattr(preflight.importlib.util,'find_spec',lambda n:object() if n=='win32com' else original(n))
    monkeypatch.setattr(preflight,'_probe_powerpoint',lambda _: (True,'OFFLINE native capability fixture'))
    def native(inputs,order,target, *, office_host='auto', diagnostics=None):
        by_id={x['slide_id']:x for x in inputs};prs=Presentation()
        first=Presentation(by_id[order[0]]['path']);prs.slide_width=first.slide_width;prs.slide_height=first.slide_height
        for sid in order:
            src=Presentation(by_id[sid]['path']).slides[0];dst=prs.slides.add_slide(prs.slide_layouts[6])
            for shape in src.shapes:
                clone=copy.deepcopy(shape.element)
                for element in clone.iter():
                    for attr,val in list(element.attrib.items()):
                        if attr in ('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed','{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link'):
                            rel=src.part.rels[val]
                            image_part,new_rid=dst.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
                            element.set(attr,new_rid)
                dst.shapes._spTree.insert_element_before(clone,'p:extLst')
        prs.save(target)
    monkeypatch.setattr(merge_pptx,'_native_merge',native)


def complete_flow(tmp_path,backend):
    state,pages=restored_pages(tmp_path,backend)
    sealer=cb.seal_page if backend=='magic_layer' else lb.seal_page
    for sid,(_,restored,review) in pages.items():sealer(state,sid,restored,review)
    result=assembly.prepare(state,tmp_path/'work');assert result['status']=='AWAITING_DECK_VALIDATION',result
    review=Path(result['validation_report']);_complete_deck_review(review)
    output=tmp_path/'outputs'/'deck.pptx';assert assembly.publish(state,review,output)['status']=='PASS'
    with zipfile.ZipFile(output) as z:assert len(z.namelist())==len(set(z.namelist()))
    prs=Presentation(output)
    assert [next(s.text for s in slide.shapes if s.has_text_frame and s.text) for slide in prs.slides]==['OFFLINE heading S001','OFFLINE heading S002','OFFLINE heading S003']
    fonts=[[(run.font.name,run.font.size) for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs for run in p.runs] for slide in prs.slides]
    assert fonts[1]==fonts[2]
    assert router.deck_provenance(r.load_runtime_state(state),['S001','S002','S003'])
    before=state.read_bytes();assert t.context(state,'S003')['contract_ref']==t.context(state,'S002')['contract_ref'];assert state.read_bytes()==before
    return state,output


@pytest.mark.parametrize('backend',['magic_layer','image_layer'])
def test_three_page_title_restoration_seal_and_assembly_offline(tmp_path,monkeypatch,backend):
    offline_native(monkeypatch);complete_flow(tmp_path,backend)


@pytest.mark.integration
@pytest.mark.parametrize('backend',['magic_layer','image_layer'])
def test_three_page_title_restoration_real_office_merge(tmp_path,backend):
    complete_flow(tmp_path,backend)


@pytest.mark.parametrize('backend',['magic_layer','image_layer'])
def test_current_title_binding_checked_before_seal_or_source_use(tmp_path,backend):
    state,pages=restored_pages(tmp_path,backend)
    saved=r.load_runtime_state(state);saved['stage2_run']['title_policy']['page_bindings']['S003']['contract_ref']='unknown-title-ref';r._save_runtime_state(state,saved)
    before=state.read_bytes();files={str(p):p.read_bytes() for triple in pages.values() for p in triple}
    sealer=cb.seal_page if backend=='magic_layer' else lb.seal_page
    with pytest.raises(ValueError,match='TITLE_MANIFEST_REFERENCE_MISMATCH|TITLE_CONTRACT_REFERENCE_UNKNOWN'):sealer(state,'S003',pages['S003'][1],pages['S003'][2])
    assert state.read_bytes()==before
    assert all(Path(p).read_bytes()==data for p,data in files.items())


def test_magic_text_plan_currentness_checks_title_and_preserves_attempt_budget(tmp_path):
    state,pages=restored_pages(tmp_path,'magic_layer');saved=r.load_runtime_state(state)
    old=copy.deepcopy(saved['canva_bridge']);saved['stage2_run']['title_policy']['page_bindings']['S003']['contract_ref']='unknown';r._save_runtime_state(state,saved)
    before=state.read_bytes()
    with pytest.raises(ValueError,match='TITLE_'):cb._check_text_plan(r.load_runtime_state(state),'S003')
    assert r.load_runtime_state(state)['canva_bridge']==old
    assert state.read_bytes()==before
